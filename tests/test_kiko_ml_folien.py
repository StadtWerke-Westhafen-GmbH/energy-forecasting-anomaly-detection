"""Kikos reduzierter ML-Teil in den Folienlayouts des SWW-Designsystems (für Google Slides)."""

from __future__ import annotations

import re
import runpy
import zipfile
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]

# Ergebnisse aus dem Testjahr 2025 zeigt Patrick; sie stehen nicht auf Kikos Hauptfolien
PATRICK_NUMBERS = ["9.188", "15,9", "3.725", "0,901", "114 ", "9,5", "6,2 %", "7.256", "10 von 10", "107 Zähler"]


@pytest.fixture(scope="module")
def mod():
    return runpy.run_path(str(ROOT / "scripts" / "build_kiko_ml_folien.py"))


@pytest.fixture(scope="module")
def deck(mod, tmp_path_factory):
    from pptx import Presentation

    path = mod["build"](tmp_path_factory.mktemp("ml") / "ml.pptx")
    return path, Presentation(path)


def px(emu) -> float:
    return emu / 9525


def all_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from all_shapes(sh.shapes)


def texts(slide) -> str:
    out = []
    for sh in all_shapes(slide.shapes):
        if sh.has_text_frame:
            out.append(sh.text_frame.text)
        elif getattr(sh, "has_table", False) and sh.has_table:
            out.extend(c.text for r in sh.table.rows for c in r.cells)
    return "\n".join(out)


def nums(prs) -> list[str]:
    return [next((sh.text_frame.text for sh in s.shapes if sh.name == "Seitenzahl"), "") for s in prs.slides]


def main_slides(prs):
    return [s for s, n in zip(prs.slides, nums(prs)) if "/" in n]


def test_neun_hauptfolien_und_vier_backups(deck):
    _, prs = deck
    assert nums(prs) == [f"{i} / 9" for i in range(1, 10)] + [f"B{i}" for i in range(1, 5)]


def test_reihenfolge_der_hauptfolien(deck):
    _, prs = deck
    t = [texts(s) for s in main_slides(prs)]
    needles = ["Was in diesem Teil beantwortet wird", "ML Canvas", "Vergangenheit erklärt Zukunft",
               "Vollaststunden", "Warum RMSE", "Random Forest", "99 von 100", "ZL-00147", "Klassifikation"]
    for i, needle in enumerate(needles):
        assert needle in t[i], (i + 1, needle)


def test_keine_testjahr_ergebnisse_auf_hauptfolien(deck):
    _, prs = deck
    for n, s in enumerate(main_slides(prs), 1):
        t = texts(s) + " "
        for value in PATRICK_NUMBERS:
            assert value not in t, (n, value)


def test_rmse_folie_erklaert_die_wahl(deck):
    _, prs = deck
    t = texts(main_slides(prs)[4])
    for needle in ["10 · 10 · 10 · 10", "0 · 0 · 0 · 40", "10 kWh", "20 kWh", "√", "MAE", "Spotmarkt", "kWh",
                   "Notebook 13"]:
        assert needle in t, needle
    bars = [sh for sh in main_slides(prs)[4].shapes if sh.name.startswith("Balken ")]
    assert len(bars) == 4  # natives Diagramm: MAE und RMSE je Szenario


def test_canvas_beschreibt_ml_felder_und_verweist_die_anderen(deck):
    _, prs = deck
    s = main_slides(prs)[1]
    bodies = {sh.name[len("Feld "):]: sh.text_frame.text for sh in s.shapes if sh.name.startswith("Feld ")}
    questions = {sh.name[len("Frage "):]: sh.text_frame.text for sh in s.shapes if sh.name.startswith("Frage ")}
    for nr in ["03", "04", "05", "06", "07", "09"]:
        assert len(bodies[nr]) >= 50, nr
        assert questions[nr].endswith("?"), nr  # jede Karte sagt, welche Frage das Feld beantwortet
    t = texts(s)
    for other in ["Mehrwert", "Datenquellen", "Impact", "Monitoring", "Iana", "Patrick"]:
        assert other in t, other
    assert not any(nr in bodies for nr in ["01", "02", "08", "10"])


def test_schwelle_erklaert_herleitung_und_uebergibt_die_wahl(deck):
    """Kiko erklärt, wie die Grenze entsteht; welche Grenze im Betrieb passt, zeigt Patrick."""
    _, prs = deck
    s = main_slides(prs)[6]
    t = texts(s)
    for needle in ["99 von 100", "144,4", "1.397", "Nov–Dez 2024", "eingefroren", "Faktor", "Startwert",
                   "Patrick"]:
        assert needle in t, needle
    assert "Falsch-positiv" not in t
    assert [sh for sh in s.shapes if sh.name == "Diagramm Kalibrierung"]
    notes = s.notes_slide.notes_text_frame.text
    assert "Falsch-positiv" in notes and "Falsch-negativ" in notes  # Abwägung als Erklärung in den Notizen


def test_prueffall_zeigt_bewertung_lesbar(deck):
    """Die Bewertungsfragen stehen als Folientext, nicht als unlesbar kleiner Screenshot."""
    _, prs = deck
    s = main_slides(prs)[7]
    assert not [sh for sh in s.shapes if sh.name.startswith("Screenshot")]
    t = texts(s)
    for needle in ["Messwert gültig?", "Abweichung erklärt?", "Ursache", "Ja", "Nein", "Label", "+2,30"]:
        assert needle in t, needle


def test_ausblick_klassifikation_und_uebergabe(deck):
    _, prs = deck
    t = texts(main_slides(prs)[8])
    for needle in ["Klassifikation", "Wahrscheinlichkeit", "%", "automatisiert", "Patrick", "Retraining",
                   "Das Modell ersetzt keine Abrechnungsentscheidung"]:
        assert needle in t, needle


def test_designsprache_der_vorlage(deck):
    """Label in Versalien mit Namen, Titel als Aussage, Icon-Scheiben nur auf der Einleitung."""
    _, prs = deck
    for n, s in enumerate(main_slides(prs), 1):
        names = [sh.name for sh in all_shapes(s.shapes)]
        discs = [x for x in names if x.startswith("Scheibe") or x.startswith("Icon ")]
        assert len(discs) == (3 if n == 1 else 0), (n, discs)
        if n in (1, 9):
            continue  # Kapitel- und Abschlussfolie
        label = next(sh for sh in s.shapes if sh.name == "Label").text_frame.text
        assert label == label.upper() and "KIKO" in label, (n, label)
        assert next(sh for sh in s.shapes if sh.name == "Titel").text_frame.text


def test_notizen_zum_ablesen_und_verstehen(deck):
    """Jede Hauptfolie: Sprechtext zum Ablesen (etwa eine Minute), Hintergrund, Stichworte, Rückfragen."""
    _, prs = deck
    total = 0
    for n, s in enumerate(main_slides(prs), 1):
        notes = s.notes_slide.notes_text_frame.text
        for head in ["SPRECHTEXT", "SO VERSTEHST DU ES", "STICHWORTE"]:
            assert head in notes, (n, head)
        speech = notes.split("SPRECHTEXT", 1)[1].split("SO VERSTEHST DU ES", 1)[0]
        words = len(speech.replace("[Pause]", "").split())
        assert (30 if n == 1 else 70) <= words <= 190, (n, words)
        total += words
    assert total <= 1250  # rund 8 bis 9 Minuten bei etwa 140 Wörtern pro Minute


def test_einleitung_mit_drei_farbigen_fragen(deck):
    _, prs = deck
    s = main_slides(prs)[0]
    tones = [sh.name for sh in all_shapes(s.shapes) if sh.name.startswith("Scheibe")]
    assert tones == ["Scheibe Vorhersage", "Scheibe Messen", "Scheibe Prüfhinweis"]
    t = texts(s)
    for needle in ["Folie 2", "Folie 5", "Folie 7"]:
        assert needle in t, needle


def test_google_slides_tauglich(deck):
    path, _ = deck
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        assert not [n for n in names if n.startswith("ppt/charts/") or n.lower().endswith(".svg")]
        fonts = set()
        for n in names:
            if n.startswith("ppt/slides/slide"):
                xml = z.read(n).decode("utf-8")
                assert "<a:alpha " not in xml and 'cap="all"' not in xml and ' spc="-' not in xml, n
                fonts |= set(re.findall(r'<a:latin typeface="([^"]+)"', xml))
        assert fonts <= {"IBM Plex Sans", "Geist Mono"}, fonts


def test_text_passt_mit_reserve(deck, mod):
    """Einzeilig: Text höchstens 97 % der Breite; mehrzeilig: Zeilen × Zeilenhöhe + 5 % passen."""
    _, prs = deck
    wrap_lines, measure = mod["K"]["wrap_lines"], mod["K"]["measure"]
    for n, s in enumerate(prs.slides, 1):
        for sh in all_shapes(s.shapes):
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            need, count, widest = 0.0, 0, 0.0
            for i, p in enumerate(sh.text_frame.paragraphs):
                runs = [r for r in p.runs if r.text]
                if not runs:
                    continue
                r0 = runs[0]
                indent = int(p._p.pPr.get("marL", 0)) / 9525 if p._p.pPr is not None else 0
                size, bold, font = r0.font.size.pt / 0.75, bool(r0.font.bold), r0.font.name
                line = "".join(r.text for r in runs)
                widest = max(widest, measure(line, size, bold, font) + indent)
                if sh.text_frame.word_wrap is False:
                    continue
                lines = wrap_lines(line, px(sh.width) - indent, size, bold, font)
                need += len(lines) * px(p.line_spacing) + (px(p.space_before) if i and p.space_before else 0)
                count += len(lines)
            where = (n, sh.text_frame.text[:50], need, px(sh.height), widest, px(sh.width))
            if sh.text_frame.word_wrap is False:
                assert widest * 1.08 <= px(sh.width) + 0.5, where
            elif count == 1:
                assert widest <= 0.97 * px(sh.width) and need <= px(sh.height) + 1, where
            elif count:
                assert need * 1.05 <= px(sh.height) + 1, where


def test_mindestschriftgroesse_auf_hauptfolien(deck):
    """Weniger Text, dafür lesbar: Fließtext auf Hauptfolien mindestens 16 px (Fußzeile ausgenommen)."""
    _, prs = deck
    for n, s in enumerate(main_slides(prs), 1):
        for sh in all_shapes(s.shapes):
            if not sh.has_text_frame or sh.top / 9525 >= 646:
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.text.strip():
                        assert r.font.size.pt / 0.75 >= 15.9, (n, r.text[:40], r.font.size.pt / 0.75)
