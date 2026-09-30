"""Tests für die gemeinsame Präsentation von Gruppe 6 (Google-Slides-tauglich)."""

from __future__ import annotations

import json
import runpy
from pathlib import Path

import pytest
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
TEAM_DIR = ROOT / "docs" / "presentation" / "team"
CHART_DIR = TEAM_DIR / "assets" / "charts"


@pytest.fixture(scope="module")
def team():
    return runpy.run_path(str(ROOT / "scripts" / "build_team_pptx.py"))


@pytest.fixture(scope="module")
def deck(team, tmp_path_factory):
    from pptx import Presentation

    path = team["build"](tmp_path_factory.mktemp("team") / "team.pptx")
    return path, Presentation(path)


def px(emu) -> float:
    return emu / 9525


def texts(slide) -> str:
    out = []
    for sh in slide.shapes:
        stack = [sh]
        while stack:
            cur = stack.pop()
            if cur.shape_type == 6:  # Gruppe
                stack.extend(cur.shapes)
            elif cur.has_text_frame:
                out.append(cur.text_frame.text)
            elif getattr(cur, "has_table", False) and cur.has_table:
                out.extend(c.text for r in cur.table.rows for c in r.cells)
    return "\n".join(out)


def all_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from all_shapes(sh.shapes)


def test_paket_ist_google_slides_tauglich(deck):
    """Keine Diagramm-Parts, keine SVGs, keine Transparenzen, keine Versalien-Formatierung."""
    import re
    import zipfile

    path, _ = deck
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        assert not [n for n in names if n.startswith("ppt/charts/") or n.startswith("ppt/embeddings/")]
        assert not [n for n in names if n.lower().endswith(".svg")]
        fonts = set()
        for n in names:
            if n.startswith("ppt/slides/slide") or n.startswith("ppt/theme/"):
                xml = z.read(n).decode("utf-8")
                if n.startswith("ppt/slides/"):  # Theme-Effektstile nutzt keine Form (kein p:style)
                    assert "<a:alpha " not in xml, n
                    assert 'cap="all"' not in xml, n
                fonts |= set(re.findall(r'<a:latin typeface="([^"]+)"', xml))
        assert fonts <= {"IBM Plex Sans", "Geist Mono"}, fonts


def test_kikos_folien_vollstaendig_und_in_reihenfolge(deck, team):
    _, prs = deck
    titles = [texts(s) for s in prs.slides]
    order = ["ML Canvas: vom Nutzen", "Fehler kosten Spotmarkt", "Methodik und Modell",
             "Vergangenheit erklärt Zukunft", "Das Modell lernt Vollaststunden", "Random Forest knapp vorn",
             "Aus dem Prognosefehler wird", "Die Schwelle wird vor 2025", "114 von 8.398",
             "ZL-00147: Im August", "Prüfen, bewerten und aus dem Feedback"]
    pos = [next(i for i, t in enumerate(titles) if needle in t) for needle in order]
    assert pos == sorted(pos) and pos[-1] - pos[0] == len(order) - 1


def test_kapitelleiste_nennt_sprecher(deck):
    _, prs = deck
    canvas = next(s for s in prs.slides if "Fehler kosten Spotmarkt" in texts(s))
    names = [sh.name for sh in canvas.shapes]
    assert [n for n in names if n.startswith("Kapitel ")] == [
        "Kapitel 01", "Kapitel 02", "Kapitel 03 aktiv", "Kapitel 04", "Kapitel 05", "Kapitel 06"]
    label = next(sh for sh in canvas.shapes if sh.name == "Kapitel-Label")
    assert label.text_frame.text == "03 · ML Canvas · Kiko"


def test_split_folie_ohne_kapitelleiste(deck):
    """Auf der Split-Folie nennt die große Kapitelnummer das Kapitel; die Leiste würde die Leitfrage überdecken."""
    _, prs = deck
    split = next(s for s in prs.slides if any(sh.name == "Split links" for sh in s.shapes)
                 and "Methodik und Modell" in texts(s))
    assert not [sh for sh in split.shapes if sh.name.startswith("Kapitel")]


def test_diagramme_als_bild_und_kartenrahmen_bleibt(deck):
    _, prs = deck
    slide = next(s for s in prs.slides if "Random Forest knapp vorn" in texts(s))
    pics = [sh for sh in slide.shapes if sh.name.startswith("Diagramm ")]
    assert len(pics) == 2
    for pic in pics:
        l, t, r, b = px(pic.left), px(pic.top), px(pic.left + pic.width), px(pic.top + pic.height)
        frames = [sh for sh in slide.shapes if sh.shape_type == 1 and px(sh.left) < l and px(sh.top) < t
                  and px(sh.left + sh.width) > r and px(sh.top + sh.height) > b]
        assert frames, "Kartenrahmen um das Diagramm fehlt"


def test_textfelder_ohne_umbruch_haben_reserve(deck, team):
    """Google Slides setzt Schrift minimal anders; einzeilige Felder brauchen 8 % Luft."""
    _, prs = deck
    measure = team["K"]["measure"]
    for n, slide in enumerate(prs.slides, 1):
        for sh in all_shapes(slide.shapes):
            if not sh.has_text_frame or sh.text_frame.word_wrap is not False:
                continue
            for p in sh.text_frame.paragraphs:
                width = sum(measure(r.text, r.font.size.pt / 0.75, bool(r.font.bold), r.font.name)
                            for r in p.runs if r.text)
                assert width * 1.08 <= px(sh.width) + 0.5, (n, p.text, width, px(sh.width))


def test_html_und_team_deck_laufen_nicht_auseinander(team):
    specs = team["kiko_specs"]()
    with pytest.raises(AssertionError, match="auseinander"):
        team["check_sync"](specs[:-1])


# --------------------------------------------------------------------------- Team-Folien

IANA, KIKO, PATRICK = [2, 3, 4], list(range(5, 16)), [16, 17, 18, 19]


def notes_of(slide) -> str:
    return slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""


def test_folienfolge_nach_agenda(deck):
    """Titel, Agenda, Iana, Kiko, Patrick, Schluss – Sprecherwechsel nur zweimal."""
    _, prs = deck
    t = [texts(s) for s in prs.slides]
    expected = {0: "Prognose des monatlichen Energieverbrauchs", 1: "Sechs Kapitel, drei Stimmen",
                2: "Planung per Hand, Auffälliges erst im Quartal", 3: "700 Zähler, 24 Monate",
                4: "Vier Problemarten", 5: "ML Canvas: vom Nutzen", 15: "Prüfen, bewerten und aus dem Feedback",
                16: "Genauer planen, früher prüfen", 17: "Das Verbrauchs-Cockpit führt", 18: "Schattenbetrieb",
                19: "Drei Rollen, drei Erkenntnisse", 20: "Ihre Fragen", 21: "Backup"}
    for i, needle in expected.items():
        assert needle in t[i], (i, needle)


def test_nummern_und_sprecher_in_der_fusszeile(deck):
    _, prs = deck
    nums = [next((sh.text_frame.text for sh in s.shapes if sh.name == "Seitenzahl"), "") for s in prs.slides]
    assert nums[1] == "2 / 21" and nums[5] == "6 / 21" and nums[20] == "21 / 21"
    backups = [n for n in nums if n.startswith("B")]
    assert backups == [f"B{i}" for i in range(1, 9)]
    for i in IANA:
        assert "Iana" in texts(prs.slides[i]), i
    for i in PATRICK:
        assert "Patrick" in texts(prs.slides[i]), i


def test_leitfragen_regel_auf_allen_folien(deck):
    _, prs = deck
    found = 0
    for s in prs.slides:
        for sh in s.shapes:
            if sh.name == "Leitfrage":
                q = sh.text_frame.text
                assert q.endswith("?") and len(q) <= 58, q
                found += 1
    assert found >= 19


def test_titel_nennt_gruppe_datum_und_alle_namen(deck):
    _, prs = deck
    t = texts(prs.slides[0])
    for needle in ["Iana Kraievska", "Kiko Ramon Lukas", "Patrick Olmo Hederer", "01.10.2026", "Gruppe 6",
                   "Wie können historische Verbrauchsdaten"]:
        assert needle in t, needle


def test_agenda_mit_kapiteln_minuten_und_personen(deck):
    _, prs = deck
    t = texts(prs.slides[1])
    for name in ["Ausgangssituation", "Daten und Datenqualität", "ML Canvas", "Methodik und Modell", "Ergebnisse",
                 "Empfehlungen und Ausblick", "Iana", "Kiko", "Patrick", "min"]:
        assert name in t, name


def test_ianas_zahlen_wie_in_ihren_folien(deck):
    _, prs = deck
    t = "\n".join(texts(prs.slides[i]) for i in IANA)
    for number in ["16.830", "16.800", "11.316", "505", "8.400", "700", "30", "0,18 %", "20", "0,12 %",
                   "180 Mio. EUR", "Stefan Lechtenberg", "Anke Bürger"]:
        assert number in t, number


def test_entwuerfe_fuer_iana_und_patrick_sind_markiert(deck):
    _, prs = deck
    for i in IANA + PATRICK:
        n = notes_of(prs.slides[i])
        assert "Entwurf aus Bericht" in n and "SPRECHTEXT" in n, i
    for i in KIKO:
        assert "Entwurf" not in notes_of(prs.slides[i]), i


def test_patrick_ergebnisse_und_fazit(deck):
    _, prs = deck
    t = texts(prs.slides[16])
    for needle in ["15,9", "9,5", "10 von 10", "92 %", "Henrik Maaß"]:
        assert needle in t, needle
    fazit = texts(prs.slides[19])
    for needle in ["Iana", "Patrick", "Kiko", "Validierung wichtiger", "Korrelation", "Monatsenergie"]:
        assert needle in fazit, needle
    emp = texts(prs.slides[18])
    for needle in ["Empfehlung 1", "Empfehlung 2", "Effekt", "Ausblick"]:
        assert needle in emp, needle


def test_cockpit_screenshot_echt_und_keine_designvorlage(deck):
    import hashlib

    _, prs = deck
    sha = lambda b: hashlib.sha1(b).hexdigest()  # noqa: E731
    real = sha((TEAM_DIR / "assets" / "verbrauchs-cockpit-prueffall.png").read_bytes())
    pics = {i: [sha(sh.image.blob) for sh in all_shapes(s.shapes) if sh.shape_type == 13]
            for i, s in enumerate(prs.slides)}
    assert real in pics[17]
    synthetic = ROOT / ".build" / "screenshots" / "dashboard.png"
    if synthetic.exists():
        assert all(sha(synthetic.read_bytes()) not in p for p in pics.values())


def test_backup_mit_ethik_und_datenqualitaet(deck):
    _, prs = deck
    t = [texts(s) for s in prs.slides]
    ethik = next(x for x in t if "Ethik" in x and "Datenschutz" in x)
    for needle in ["im Pilot festlegen", "Das Modell ersetzt keine Abrechnungsentscheidung", "Kundentyp"]:
        assert needle in ethik, needle
    dq = next(x for x in t if "Datenqualität im Detail" in x)
    for needle in ["672", "MWh → kWh", "504", "Post-hoc imputiert", "11.316", "67,3 %"]:
        assert needle in dq, needle


def test_diagramm_pngs_passen_zu_ihren_rechtecken():
    """Jedes Diagramm-PNG ist ein 3-facher Element-Screenshot seines SVG-Rechtecks."""
    charts = json.loads((CHART_DIR / "charts.json").read_text(encoding="utf-8"))
    assert len(charts) >= 8
    for c in charts:
        png = CHART_DIR / c["file"]
        assert png.exists(), c["file"]
        w, h = Image.open(png).size
        assert abs(w - 3 * c["w"]) <= 3 and abs(h - 3 * c["h"]) <= 3, c
        assert 0 <= c["x"] and c["x"] + c["w"] <= 1280 and 0 <= c["y"] and c["y"] + c["h"] <= 720, c
