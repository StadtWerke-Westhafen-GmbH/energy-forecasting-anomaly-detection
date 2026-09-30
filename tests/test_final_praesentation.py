"""Endversion der Team-Präsentation: einheitliches Design, Inhalte unverändert (außer Kikos Folie 17)."""

from __future__ import annotations

import re
import runpy
import zipfile
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "Copy of Praesentationsvorlage_IHK(3).pptx"
LABEL = re.compile(r"^(KAPITEL \d( UND \d)? · ([A-ZÄÖÜ0-9\- ]+ · )?(IANA|KIKO|PATRICK|TEAM)"
                   r"|BACKUP · [A-ZÄÖÜ0-9\- ]+ · (IANA|KIKO|PATRICK)|AGENDA · GRUPPE 6|DATENANALYSE-PROJEKT · GRUPPE 6)$")

pytestmark = pytest.mark.skipif(not SRC.exists(), reason="Team-Datei aus dem Drive fehlt")


@pytest.fixture(scope="module")
def mod():
    return runpy.run_path(str(ROOT / "scripts" / "build_final_praesentation.py"))


@pytest.fixture(scope="module")
def deck(mod, tmp_path_factory):
    from pptx import Presentation

    path = mod["build"](tmp_path_factory.mktemp("final") / "final.pptx")
    return path, Presentation(path)


@pytest.fixture(scope="module")
def source():
    from pptx import Presentation

    return Presentation(SRC)


def all_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from all_shapes(sh.shapes)


def texts(slide) -> str:
    out = []
    for sh in all_shapes(slide.shapes):
        if sh.has_text_frame:
            out.append(sh.text_frame.text)
        elif getattr(sh, "has_table", False) and sh.has_table:
            out.extend(c.text for r in sh.table.rows for c in r.cells)
    return "\n".join(out)


def norm(t: str) -> str:
    return re.sub(r"\s+", " ", t.replace("\x0b", " ")).strip()


def page_number(slide) -> str:
    for sh in slide.shapes:
        if sh.has_text_frame and sh.top / 9525 > 640 and sh.left / 9525 > 1000:
            t = sh.text_frame.text.strip()
            if re.fullmatch(r"\d+ / \d+|B\d+", t):
                return t
    return ""


def label_of(slide) -> str:
    for sh in slide.shapes:
        if sh.has_text_frame and sh.top / 9525 < 70:
            t = sh.text_frame.text.strip()
            if t and t == t.upper() and "·" in t:
                return t
    return ""


def test_folienfolge(deck):
    _, prs = deck
    t = [texts(s) for s in prs.slides]
    assert len(t) == 36  # 27 Hauptfolien, Backup-Trenner, 8 Backups
    expected = {0: "Prognose des monatlichen Energieverbrauchs", 1: "Agenda", 2: "Ausgangssituation",
                3: "Datenstruktur", 4: "Datenqualität", 5: "Typischer Stromverbrauch der Zähler",
                6: "Höchster monatlicher Stromverbrauch", 7: "Verbrauchsspitze im Verhältnis",
                8: "Machine Learning: vom Canvas zum Prüfhinweis", 15: "ZL-00147",
                16: "Vom Prüfhinweis zum Klassifikationsmodell", 17: "Ergebnisse", 25: "Persönliches Fazit",
                26: "Vielen Dank", 27: "Backup", 28: "Datenqualität im Detail", 32: "Acht kontrollierte",
                35: "ML Canvas im Wortlaut"}
    for i, needle in expected.items():
        assert needle in t[i], (i + 1, needle)


def test_seitenzahlen_durchgehend(deck):
    _, prs = deck
    nums = [page_number(s) for s in prs.slides]
    assert nums[1:27] == [f"{n} / 27" for n in range(2, 28)]
    assert [n for n in nums[27:] if n] == [f"B{i}" for i in range(1, 9)]
    assert not any("/ 9" in n or "/ 21" in n for n in nums)


def test_labels_folgen_dem_schema(deck):
    _, prs = deck
    for n, s in enumerate(prs.slides, 1):
        label = label_of(s)
        if n in (27, 28):  # Dank und Backup-Trenner ohne Label
            continue
        if n == 1:  # Titelfolie: Label steht laut Layout über dem Titel in der Folienmitte
            assert "DATENANALYSE-PROJEKT · GRUPPE 6" in texts(s)
            continue
        assert LABEL.match(label), (n, label)
    assert label_of(prs.slides[19]) == "KAPITEL 5 · TOP-3-TREIBER · PATRICK"
    assert label_of(prs.slides[2]) == "KAPITEL 1 · AUSGANGSSITUATION · IANA"


def test_fusszeile_mit_namen(deck):
    _, prs = deck
    names = {"IANA": "Iana Kraievska", "KIKO": "Kiko Ramon Lukas", "PATRICK": "Patrick Olmo Hederer",
             "TEAM": "Gruppe 6"}
    for n, s in enumerate(prs.slides, 1):
        label = label_of(s)
        who = next((k for k in names if label.endswith(k)), None)
        if not who or n == 1:
            continue
        foot = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.top / 9525 > 640
                and sh.text_frame.text.startswith("Verbrauchsprognose & Frühwarnung")]
        assert foot == [f"Verbrauchsprognose & Frühwarnung · {names[who]}"], (n, foot)


def test_inhalte_von_iana_woertlich(deck, source):
    """Tabellen und Texte aus der Team-Datei stehen unverändert in der Endversion."""
    _, prs = deck
    final = norm(" ".join(texts(s) for s in prs.slides))
    for i in (4, 30):  # Datenstruktur, Datenqualität im Detail
        table = next(sh for sh in source.slides[i - 1].shapes if getattr(sh, "has_table", False) and sh.has_table)
        for row in table.table.rows:
            for c in row.cells:
                if c.text.strip() and not c.text.startswith("Datensatz"):
                    assert norm(c.text) in final, (i, c.text)
    for needle in ["StadtWerke Westhafen GmbH", "Ungenaue Prognosen des Stromverbrauchs",
                   "Verspätete Erkennung von Anomalien", "Steigende Beschaffungskosten",
                   "Verzögerungen bei der Erkennung technischer Probleme",
                   "Ein automatisiertes Prognosesystem ist nicht vorhanden.",
                   "ML-Modell für Verbrauchsprognose und Anomalieerkennung",
                   "Datenbasierte Lösung der identifizierten Probleme.", "30 (0,18 %)", "8.400 (50 %)", "700 (4 %)",
                   "11.316 (67 %)", "505 (3 %)", "20 (0,12 %)", "Strukturell bedingt → nicht imputiert",
                   "16.830 Zeilen", "16.800 Zeilen", "Wir freuen uns auf das Fachgespräch.",
                   "Detail-Folien für mögliche Rückfragen.",
                   "Datenbereinigung ist keine rein technische Routine",
                   "Bei Zeitreihendaten ist die Validierung wichtiger als die Wahl des komplexesten Modells.",
                   "Normierung verschiebt die Korrelationen",
                   "15,9 % genauer als die einfache Fortschreibung",
                   "Große Abweichungen landen nach Monatsende auf einer Prüfliste – nicht erst im Quartal."]:
        assert norm(needle) in final, needle
    assert "Tipp: Eine Folie reicht" not in final  # Vorlagen-Hinweis entfällt
    assert "01.10.2026" in texts(prs.slides[0]) and "01.10.2024" not in final
    assert "kundentyp" in texts(prs.slides[4]) and "kunden_typen" not in final


def test_kiko_und_patrick_folien_unveraendert(deck, source):
    """Übernommene Folien behalten ihren Inhalt; geändert werden nur Label, Seitenzahl und Fußzeile."""
    _, prs = deck
    keep = {9: 8, 10: 9, 11: 10, 12: 11, 13: 12, 14: 13, 15: 14, 16: 15, 19: 17, 20: 18, 21: 19, 22: 20, 23: 21,
            24: 22, 25: 23, 26: 24, 31: 29, 32: 30, 33: 31, 34: 32, 35: 33, 36: 34, 37: 35}
    skip = re.compile(r"^(\d+( / \d+)?|B\d|Verbrauchsprognose & Frühwarnung.*|[A-ZÄÖÜ0-9 ·\-]+)$")
    for old, new in keep.items():
        a = [norm(x) for x in texts(source.slides[old - 1]).split("\n") if x.strip() and not skip.match(norm(x))]
        b = norm(texts(prs.slides[new]))
        for line in a:
            assert line in b, (old, line)


def test_folie_17_fuehrt_pruefen_und_ausblick_zusammen(deck):
    _, prs = deck
    s = prs.slides[16]
    t = texts(s)
    for needle in ["Heute im Cockpit", "Fallakte", "Messwert gültig", "Abweichung erklärt", "Prüfwarteschlange",
                   "bestätigte Anomalie", "Klassifikation", "Wahrscheinlichkeit", "%", "automatisiert",
                   "Das Modell ersetzt keine Abrechnungsentscheidung", "Patrick", "Retraining"]:
        assert needle in t, needle
    notes = s.notes_slide.notes_text_frame.text
    assert "SPRECHTEXT" in notes and "SO VERSTEHST DU ES" in notes


def test_backup_verweise_in_notizen_umnummeriert(deck):
    _, prs = deck
    assert "Backup B5" in prs.slides[13].notes_slide.notes_text_frame.text  # Kikos Modellwahl → Hyperparameter
    assert "Backup B7" in prs.slides[14].notes_slide.notes_text_frame.text  # Kikos Schwelle → Kalibrierung
    assert "Backup B4" in prs.slides[21].notes_slide.notes_text_frame.text  # Patricks Prüfliste → Prüffall


def test_ianas_diagramme_neu_im_designsystem(deck, mod):
    _, prs = deck
    for i in (5, 6, 7):
        pics = [sh for sh in prs.slides[i].shapes if sh.name.startswith("Diagramm ")]
        assert len(pics) == 1, i + 1
    stats = mod["eda_stats"]()
    assert stats["zaehler"] == 700 and round(stats["max_2024"]) == 675231


def test_neue_folien_google_tauglich(deck, mod):
    from lxml import etree

    path, prs = deck
    with zipfile.ZipFile(path) as z:
        assert not [n for n in z.namelist() if n.lower().endswith(".svg")]
    for i in mod["NEW_SLIDES"]:
        assert not list(prs.slides[i].placeholders), i + 1  # sonst „Klicken, um Titel hinzuzufügen“ in Google
        xml = etree.tostring(prs.slides[i]._element).decode("utf-8")
        assert "<a:alpha " not in xml and 'cap="all"' not in xml and ' spc="-' not in xml, i + 1
        assert set(re.findall(r'<a:latin typeface="([^"]+)"', xml)) <= {"IBM Plex Sans", "Geist Mono"}, i + 1
        assert ' i="1"' not in xml, i + 1  # keine Kursivschrift


def test_quelle_bleibt_unveraendert(mod, deck):
    import hashlib

    before = mod["SOURCE_SHA"]
    assert hashlib.sha1(SRC.read_bytes()).hexdigest() == before


def test_text_passt_auf_neuen_folien(deck, mod):
    """Neue Folien: einzeilige Felder mit 3 % Breitenreserve, mehrzeilige mit 5 % Höhenreserve."""
    _, prs = deck
    K = mod["kit"].K
    for i in mod["NEW_SLIDES"]:
        for sh in all_shapes(prs.slides[i].shapes):
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            need, count, widest = 0.0, 0, 0.0
            for j, p in enumerate(sh.text_frame.paragraphs):
                runs = [r for r in p.runs if r.text]
                if not runs:
                    continue
                r0 = runs[0]
                indent = int(p._p.pPr.get("marL", 0)) / 9525 if p._p.pPr is not None else 0
                size, bold, font = r0.font.size.pt / 0.75, bool(r0.font.bold), r0.font.name
                line = "".join(r.text for r in runs)
                widest = max(widest, K["measure"](line, size, bold, font) + indent)
                if sh.text_frame.word_wrap is False:
                    continue
                lines = K["wrap_lines"](line, sh.width / 9525 - indent, size, bold, font)
                need += len(lines) * p.line_spacing / 9525 + (p.space_before / 9525 if j and p.space_before else 0)
                count += len(lines)
            where = (i + 1, sh.text_frame.text[:50], need, sh.height / 9525)
            if sh.text_frame.word_wrap is False:
                assert widest * 1.08 <= sh.width / 9525 + 0.5, where
            elif count == 1:
                assert widest <= 0.97 * sh.width / 9525 and need <= sh.height / 9525 + 1, where
            elif count:
                assert need * 1.05 <= sh.height / 9525 + 1, where


def test_agenda_hebt_niemanden_hervor(deck):
    from lxml import etree

    _, prs = deck
    xml = etree.tostring(prs.slides[1]._element).decode("utf-8")
    assert "E3F3F7" not in xml  # alle Kapitelkarten gleich (weiß)


def test_seitenzahlen_einheitlich_gesetzt(deck):
    """Alle Seitenzahlen in IBM Plex Sans ohne Laufweite (Patricks standen in Geist Mono, „19  /  27“)."""
    _, prs = deck
    for n, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.has_text_frame and sh.top / 9525 > 640 and re.fullmatch(r"\d+ / \d+|B\d+", sh.text_frame.text):
                for r in sh.text_frame.paragraphs[0].runs:
                    assert r._r.rPr is None or not r._r.rPr.get("spc"), n
                    assert r.font.name in (None, "IBM Plex Sans"), (n, r.font.name)


def test_diagramm_achsentitel_lesbar(mod):
    for key, fig in mod["eda_figures"]().items():
        assert fig.layout.xaxis.title.font.size >= 18 and fig.layout.yaxis.title.font.size >= 18, key
