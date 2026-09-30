"""Absicherung für Kikos Präsentationsteil (scripts/build_kiko_praesentation.py)."""

from __future__ import annotations

import re
import runpy
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
BUILD = runpy.run_path(str(ROOT / "scripts" / "build_kiko_praesentation.py"))


@pytest.fixture(scope="module")
def facts():
    return BUILD["extract_facts"](BUILD["Notebook"].load())


@pytest.fixture(scope="module")
def slides(facts):
    return BUILD["build_slides"](facts)


@pytest.fixture(scope="module")
def deck(slides):
    return BUILD["render_deck"](slides)


def test_de_formatiert_deutsch():
    de, signed_de = BUILD["de"], BUILD["signed_de"]
    assert de(13272.407) == "13.272"
    assert de(144.3547, 2) == "144,35"
    assert de(0) == "0"
    assert signed_de(2.2987) == "+2,30"
    assert signed_de(-0.682) == "−0,68"
    assert signed_de(0) == "±0,00"


def test_parse_de_liest_tausender_und_komma():
    assert BUILD["parse_de"]("13.272") == 13272
    assert BUILD["parse_de"]("144,3547") == pytest.approx(144.3547)


def test_kennzahlen_entsprechen_notebook_13(facts):
    cv = {row["name"]: round(row["mean"]) for row in facts["cv"]}
    assert cv == {
        "Random Forest": 13272, "Lineare Regression": 13643,
        "Bis-zu-3-Monats-Mittel": 15483, "Vormonat": 18460,
    }
    assert all(len(row["folds"]) == 3 for row in facts["cv"])
    assert round(facts["test_rmse"]) == 9188
    assert facts["test_mae"] == 3725
    assert facts["test_n"] == 8398
    assert round(facts["baseline_gain_pct"], 1) == 15.9
    assert round(facts["vls_gain_pct"], 1) == 6.2
    assert round(facts["direct_kwh_rmse"]) == 9800


def test_schwelle_und_pruefvolumen(facts):
    assert facts["threshold"] == pytest.approx(144.3547, abs=1e-4)
    assert facts["n_calibration"] == 1397
    assert facts["n_above"] == 14
    assert facts["n_alerts"] == 114
    q99 = next(row for row in facts["workload"] if row["q"] == 99.0)
    assert q99 == {"q": 99.0, "schwelle": 144.4, "hinweise": 114, "je_monat": 9.5}
    assert [row["q"] for row in facts["workload"]] == [95.0, 97.5, 99.0, 99.5]


def test_fallbeispiel_zl_00147(facts):
    assert facts["case_kw"] == 49
    aug = 7
    assert facts["case_ist"][aug] == pytest.approx(23563.49, abs=0.01)
    assert facts["case_prognose"][aug] == pytest.approx(7303.90, abs=0.01)
    # Faktor lässt sich aus Ist, Prognose, kW und Schwelle nachrechnen
    recomputed = (facts["case_ist"][aug] - facts["case_prognose"][aug]) / facts["case_kw"] / facts["threshold"]
    assert recomputed == pytest.approx(facts["case_factor"][aug], abs=1e-3)
    assert [abs(f) > 1 for f in facts["case_factor"]].count(True) == 1


def test_fold_definition_entspricht_notebook(facts):
    assert facts["folds_in_notebook"] == BUILD["FOLD_DEFINITIONS"]


def test_hyperparameter_sieger(facts):
    best = facts["tuning"][0]
    assert (best["depth"], best["features"], best["leaf"]) == ("8", "0,7", 5)
    assert len(facts["tuning"]) == 8


def test_hauptteil_passt_mit_puffer_in_zehn_minuten(slides):
    main = [s for s in slides if not s.backup]
    assert len(main) == 11
    assert sum(s.minutes for s in main) * 60 == pytest.approx(570)  # 9:30 plus 30 s Puffer für Übergaben
    assert all(s.speech for s in main), "Jede Hauptfolie braucht einen Sprechtext"
    assert all(s.keywords for s in main), "Jede Hauptfolie braucht Stichworte für freies Sprechen"


def test_sprechtempo_laesst_zeit_fuer_grafiken(slides):
    # Frei gesprochen ~120 Wörter/min; darunter bleibt Luft für Pausen an den Grafiken
    for slide in slides:
        if not slide.backup:
            assert BUILD["words_per_minute"](slide) <= 115, slide.label


def test_reihenfolge_folgt_dem_roten_faden(slides):
    labels = [s.label for s in slides if not s.backup]
    assert labels[0] == "ML Canvas" and labels[1] == "Fehlerkosten" and labels[2] == "Roter Faden"
    assert labels.index("Zeitliche Validierung") < labels.index("Zielgröße VLS")
    assert labels[-1] == "Prüfen und bewerten"


def test_nummerierung_haupt_und_backup(slides):
    assert BUILD["slide_numbers"](slides) == [f"{n} / 11" for n in range(1, 12)] + [f"B{n}" for n in range(1, 7)]


def test_deck_enthaelt_alle_folien_und_kernzahlen(deck, slides):
    assert deck.count('<section data-minutes=') == len(slides)
    assert 'class="num">11 / 11<' in deck and 'class="num">B6<' in deck
    for text in ["13.272", "9.188", "3.725", "15,9", "144,4", "9,5", "ZL-00147", "6,2", "331,8", "114", "107",
                 "480,9", "149,1", "0,901", "7.256", "Stefan Lechtenberg", "Anke Bürger", "Henrik Maaß"]:
        assert text in deck, text
    assert "{NUM}" not in deck and "{LOGO}" not in deck
    assert "@font-face" in deck and "Geist Mono" in deck
    # Pfeile und Relationszeichen kommen aus dem eingebetteten Symbol-Subset
    assert "unicode-range:U+2192" in deck


def test_deutungen_aus_der_pruefung_bleiben_korrigiert(deck, slides):
    """Regressionstests für die Befunde der unabhängigen Prüfung."""
    speech = " ".join(" ".join(s.speech) for s in slides)
    # Faktor = Abweichung relativ zur Toleranz, nicht Verbrauch relativ zur Grenze
    assert "Verbrauch 2,3-fach über der Grenze" not in deck
    assert "Abweichung 2,3-mal so groß wie erlaubt" in deck
    # Finales Modell lernt auf ganz 2024; Nov/Dez nicht "nur" Kalibrierung
    assert "nutzen wir nur, um die Schwelle" not in speech
    assert "Lernen · ganz 2024" in deck
    # Recall braucht eine Stichprobe unauffälliger Fälle
    assert "Stichprobe unauffälliger" in speech and "Stichprobe unauffälliger" in deck
    assert "Precision/Recall-Werte" not in deck
    # Demo-Bewertung ist als Beispiel gekennzeichnet
    assert "Beispieleingabe" in deck
    # Regel mit Vorzeichen: plus oder minus eins
    assert "F ≥ +1 oder F ≤ −1" in deck
    # Nicht mehr als "Kapitel" nummeriert (passt nicht zur Team-Agenda)
    assert "Kapitel 2" not in deck and "Kapitel 5" not in deck


def test_beide_modellvergleiche_nutzen_dieselbe_skala(facts):
    left, right = BUILD["chart_cv"](facts), BUILD["chart_bench"](facts)
    ticks = lambda svg: re.findall(r'text-anchor="middle" >(\d+)</text>', svg)
    assert ticks(left) == ticks(right) == ["0", "7", "14", "21", "28"]


def test_streudiagramm_zeigt_alle_hinweise(facts):
    svg = BUILD["chart_scatter"](facts)
    inside = sum(1 for x, y in facts["scatter_alerts"] if 0 <= x <= 800 and 0 <= y <= 800)
    outside = facts["n_alerts"] - inside
    assert outside == 4
    assert f"▲ {outside} Hinweise über 800 VLS-h" in svg


def test_alerts_je_monat_summieren_sich(facts):
    assert sum(facts["alerts_by_month"]) == facts["n_alerts"] == 114
    assert (min(facts["alerts_by_month"]), max(facts["alerts_by_month"])) == (4, 16)


def test_exportierte_pngs_sind_exakt_16_zu_9():
    from PIL import Image

    pngs = sorted((ROOT / "docs" / "presentation" / "kiko" / "folien").glob("*.png"))
    if not pngs:
        pytest.skip("Noch kein Export (node scripts/export_kiko_praesentation.mjs)")
    assert len(pngs) == 17
    for png in pngs:
        image = Image.open(png)
        assert image.size == (2560, 1440), png.name
        assert image.convert("RGB").getpixel((5, 0)) != (35, 44, 54), f"Deck-Hintergrund in {png.name}"
        assert png.name.isascii()


def test_screenshot_platzhalter_bei_fehlender_datei():
    html = BUILD["screenshot"]("gibt-es-nicht.png", "Test")
    assert "Screenshot fehlt" in html


def test_sprechzettel_listet_jede_folie(slides):
    notes = BUILD["render_notes"](slides)
    assert notes.count('class="n"') == len(slides)
    assert "Backup" in notes and "Falls gefragt wird" in notes


def test_deck_bleibt_klein_genug_zum_teilen(deck):
    # Wortmarke steht auf jeder Folie; eingebettet in Originalgröße wären es > 7 MB
    assert len(deck.encode("utf-8")) < 3_000_000


def test_fold_werte_sind_farbig_mit_streuung(facts):
    svg = BUILD["chart_cv"](facts)
    for color, label in zip(BUILD["FOLD_COLORS"], BUILD["FOLD_LABELS"]):
        assert svg.count(f'fill="{color}"') == 4 + 1, label  # je Kandidat eine Raute plus Legende
        assert label in svg
    assert svg.count('stroke-width="1.5" fill="none"') == 4  # ±1 Std. je Kandidat


def test_fold_1_liegt_bei_allen_kandidaten_am_hoechsten(facts):
    """Aussage auf Folie 5: Fold 1 zieht jeden Mittelwert nach oben."""
    for row in facts["cv"]:
        assert row["folds"][0] == max(row["folds"]), row["name"]


def test_kalibrierung_erklaert_wirkung_auf_2025(facts, slides):
    assert facts["q99_2025_hindsight"] == pytest.approx(175.05, abs=0.01)
    assert facts["median_cal"] == pytest.approx(facts["median_2025"], abs=0.5)  # typische Fehler gleich groß
    cal = next(s for s in slides if s.label == "Kalibrierung")
    assert not cal.backup and "dark" in cal.html
    for text in ["eingefroren", "175,1 VLS-h", "144,4 VLS-h", "Leakage", "1,36 %"]:
        assert text in cal.html, text
    labels = [s.label for s in slides if not s.backup]
    assert labels.index("Vom Fehler zum Prüfhinweis") < labels.index("Kalibrierung") < labels.index(
        "Ist gegen Prognose 2025")


def test_kompletter_canvas_ist_hauptfolie_in_z_folge(slides):
    canvas = slides[0]
    assert canvas.label == "ML Canvas" and not canvas.backup
    order = [int(m) for m in re.findall(r'<span class="cv-nr">(\d\d)</span>', canvas.html)]
    assert order == [1, 7, 9, 3, 4, 2, 5, 6, 8, 10]  # Kohärenzpfad: Nutzen → Entscheidung → … → Monitoring
    assert "10 von 10 Feldern" in canvas.html


def test_canvas_texte_bleiben_kurz():
    for nr, (head, body, _) in BUILD["CANVAS_SHORT"].items():
        limit = 64 if nr not in ("08", "10") else 120
        assert len(body) <= limit, (nr, len(body))
        assert len(head) <= (22 if nr not in ("08", "10") else 48), (nr, head)


def test_canvas_wortlaut_entspricht_bericht_tab_a1():
    """Backup B1 zitiert Tab. A1; so passen Bericht und Präsentation zusammen (Rubrik)."""
    import docx

    doc = docx.Document(ROOT / "docs" / "IHK_Bericht_Gruppe_6_final.docx")
    cells = {row.cells[0].text.strip(): row.cells[2].text.strip()
             for table in doc.tables for row in table.rows if len(row.cells) >= 3}
    for nr, text in BUILD["CANVAS_A1"].items():
        assert cells.get(str(int(nr))) == text, nr


def test_bericht_zahlen_stehen_im_bericht():
    import docx

    doc = docx.Document(ROOT / "docs" / "IHK_Bericht_Gruppe_6_final.docx")
    text = "\n".join(p.text for p in doc.paragraphs)
    text += "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    b = BUILD["BERICHT"]
    assert f"R² von {b['r2']}" in text
    assert f"{b['streuung_zwischen_zaehlern']} % der Streuung" in text
    assert "zwischen {} und {} VLS-Stunden".format(*b["median_kundentyp"]) in text
    for _, value in b["importance"][:6]:
        assert BUILD["de"](value) in text, value


def test_designsprache(slides):
    main = [s for s in slides if not s.backup]
    for s in main:
        question = re.search(r'<div class="lq">.*?<p>(.*?)</p>', s.html, re.S)
        assert question and question.group(1).endswith("?") and len(question.group(1)) <= 58, s.label
    joined = "".join(s.html for s in slides)
    assert "border-left:5px" not in BUILD["CSS"] and "italic" not in BUILD["CSS"]
    assert joined.count('class="zw ') == 4  # Zählwerk höchstens viermal
    dark = [s.label for s in main if 'class="slide dark' in s.html]
    assert dark == ["Kalibrierung", "Prüfen und bewerten"]
    assert 'class="slide split"' in slides[2].html  # Roter Faden: links blau, rechts weiß


def test_icons_existieren(slides):
    names = set(re.findall(r'disc\("([a-z0-9-]+)"', (ROOT / "scripts" / "build_kiko_praesentation.py").read_text(encoding="utf-8")))
    names |= {icon for _, icon in BUILD["CANVAS"].values()} | set(BUILD["STEP_ICONS"])
    for name in names:
        assert (BUILD["ICON_DIR"] / f"{name}.svg").exists(), name


def test_abdeckung_der_vorgaben(slides):
    """Pflichtfragen aus Canvas-Vorlage, PowerPoint-Vorlage und Learning Journey sind sichtbar beantwortet."""
    html = "".join(s.html for s in slides if not s.backup)
    for needle in ["Was kostet es, wenn das Modell falsch liegt?",   # Canvas-Feld 08
                   "Nicht im Modell",                                  # bewusst weggelassene Merkmale
                   "Top-3-Treiber",                                    # PPT-Vorlage Ergebnisse
                   "R²",                                               # Metrik = Wert
                   "Das Modell ersetzt keine Abrechnungsentscheidung",  # Auftrag Abschn. 8
                   "Mein Fazit",                                       # persönliches Fazit
                   "Patrick: Ergebnisse"]:                             # Übergabe statt „Vielen Dank“
        assert needle in html, needle


def test_mittelwert_label_ueberdeckt_keinen_balkenwert(facts, slides):
    """Regression: das Mittelwert-Label lag im Diagramm auf dem Dezember-Wert „10“ (B6)."""
    svg = BUILD["chart_monthly"](facts)
    labels = re.findall(r'<text x="([\d.]+)" y="([\d.]+)"[^>]*>([^<]*)</text>', svg)
    assert not [t for _, _, t in labels if t.startswith("Ø")], "Mittelwert-Label gehört in die Kopfzeile der Karte"
    b6 = next(s for s in slides if s.label == "Backup Prüfaufwand")
    assert re.search(r'<span class="ch-avg">.*?Ø 9,5 je Monat</span>', b6.html)
