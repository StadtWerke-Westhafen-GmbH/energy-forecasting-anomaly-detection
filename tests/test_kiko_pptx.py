"""Absicherung für die PowerPoint-Fassung (scripts/build_kiko_pptx.py).

Die PPTX baut die 17 Folien der HTML-Fassung (maßgeblich) nativ nach. Die Tests prüfen
Struktur und Inhalte gegen dieselben Slide-Objekte und Fakten wie der HTML-Generator.
"""

from __future__ import annotations

import io
import re
import runpy
import shutil
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.shapes.group import GroupShape
from pptx.util import Emu

ROOT = Path(__file__).resolve().parents[1]
PPTX = runpy.run_path(str(ROOT / "scripts" / "build_kiko_pptx.py"))
HTML = PPTX["BUILD"]
PX = 9525  # EMU je CSS-Pixel
SHIPPED = ROOT / "docs" / "presentation" / "kiko" / "Kiko_ML_Canvas_Methodik_Prueffall.pptx"

# Folienindizes (0-basiert), Reihenfolge wie build_slides()
(CANVAS, FEHLERKOSTEN, FADEN, VALIDIERUNG, VLS, MODELL, SCHWELLE, KALIBRIERUNG, BAND, FALL,
 DASHBOARD, B1, B2, B3, B4, B5, B6) = range(17)


@pytest.fixture(scope="module")
def facts():
    return HTML["extract_facts"](HTML["Notebook"].load())


@pytest.fixture(scope="module")
def specs(facts):
    return HTML["build_slides"](facts)


@pytest.fixture(scope="module")
def deck(tmp_path_factory):
    path = tmp_path_factory.mktemp("pptx") / "kiko.pptx"
    PPTX["build"](path)
    return Presentation(path)


def _walk(shapes):
    for sh in shapes:
        yield sh
        if isinstance(sh, GroupShape):
            yield from _walk(sh.shapes)


def _texts(slide) -> str:
    return " ".join(sh.text_frame.text for sh in _walk(slide.shapes) if sh.has_text_frame)


def _named(slide, prefix):
    return [sh for sh in _walk(slide.shapes) if sh.name.startswith(prefix)]


def _is_disc(sh) -> bool:
    """Icon-Scheibe: Gruppe aus nativer Kreisform und Icon-Bild."""
    if not isinstance(sh, GroupShape):
        return False
    kids = list(sh.shapes)
    return (len(kids) == 2 and kids[0].shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
            and kids[0].auto_shape_type == MSO_SHAPE.OVAL and kids[1].shape_type == MSO_SHAPE_TYPE.PICTURE)


# --------------------------------------------------------------------------- Rahmen


def test_foliensatz_ist_16_zu_9_mit_17_folien(deck, specs):
    assert (deck.slide_width, deck.slide_height) == (Emu(1280 * PX), Emu(720 * PX))
    assert len(deck.slides) == len(specs) == 17
    assert sum(s.backup for s in specs) == 6


def test_leitfrage_und_titel_entsprechen_der_html_fassung(deck, specs):
    for slide, spec in zip(deck.slides, specs):
        question, title = PPTX["heading"](spec)
        assert question and title
        texts = _texts(slide)
        assert question in texts, question
        assert title in texts, title


def test_seitenzahlen_wie_html(deck, specs):
    numbers = HTML["slide_numbers"](specs)
    for slide, num in zip(deck.slides, numbers):
        box = next(sh for sh in slide.shapes if sh.name == "Seitenzahl")
        assert box.text_frame.text == num


def test_hintergruende_wie_html(deck, specs):
    colors = {"light": "FFFFFF", "mist": "F1F4F7", "dark": "063659", "split": "FFFFFF"}
    for slide, spec in zip(deck.slides, specs):
        kind = re.search(r'class="slide (\w+)"', spec.html).group(1)
        assert str(slide.background.fill.fore_color.rgb) == colors[kind], spec.label
    assert [i for i, s in enumerate(specs) if 'class="slide mist"' in s.html] == [FEHLERKOSTEN, VLS, SCHWELLE]
    assert [i for i, s in enumerate(specs) if 'class="slide dark"' in s.html] == [KALIBRIERUNG, DASHBOARD, B1, B5]
    panel = next(sh for sh in deck.slides[FADEN].shapes if sh.name == "Split links")
    assert (panel.width, str(panel.fill.fore_color.rgb)) == (Emu(512 * PX), "063659")  # 40 % navy


def test_notizen_mit_stichworten_sprechtext_lesehilfe_und_rueckfragen(deck, specs):
    for slide, spec in zip(deck.slides, specs):
        notes = slide.notes_slide.notes_text_frame.text
        assert "<" not in notes, "HTML-Reste in den Notizen"
        assert "SPRECHTEXT" in notes and PPTX["_plain"](spec.speech[0]) in notes
        if spec.keywords:
            assert notes.startswith("STICHWORTE") and PPTX["_plain"](spec.keywords[0]) in notes
        if spec.reading:
            assert "SO LIEST DU DIE FOLIE" in notes
        if spec.questions:
            assert "FALLS GEFRAGT WIRD" in notes and PPTX["_plain"](spec.questions[-1][0]) in notes
    assert all(s.keywords for s in specs if not s.backup)


def test_alle_formen_liegen_auf_der_folie(deck):
    width, height = deck.slide_width, deck.slide_height
    for n, slide in enumerate(deck.slides, start=1):
        for sh in _walk(slide.shapes):
            assert sh.left >= 0 and sh.top >= 0, (n, sh.name)
            assert sh.left + sh.width <= width + PX and sh.top + sh.height <= height + PX, (n, sh.name)


def test_nur_markenschriften_und_keine_kursivschrift(deck):
    fonts, italic = set(), []
    for n, slide in enumerate(deck.slides, start=1):
        frames = [sh.text_frame for sh in _walk(slide.shapes) if sh.has_text_frame]
        frames += [c.text_frame for sh in slide.shapes if sh.has_table for row in sh.table.rows for c in row.cells]
        for tf in frames:
            for run in (r for p in tf.paragraphs for r in p.runs):
                fonts.add(run.font.name)
                if run.font.italic:
                    italic.append((n, run.text))
    assert fonts <= {"IBM Plex Sans", "Geist Mono"}
    assert not italic


def test_keine_verlaeufe_und_keine_akzentstreifen(deck):
    for n, slide in enumerate(deck.slides, start=1):
        assert slide._element.find(".//" + qn("a:gradFill")) is None, n
        for sh in _walk(slide.shapes):
            if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.width <= 6 * PX and sh.height >= 20 * PX:
                # schmale Hochkant-Rechtecke nur als neutraler Faden, nie als farbiger Akzentstreifen
                assert str(sh.fill.fore_color.rgb) == "D2DAE2", (n, sh.name)


# --------------------------------------------------------------------------- Designsprache


def test_versal_labels_bleiben_als_text_editierbar(deck):
    """CSS text-transform:uppercase wird zu „Alle Großbuchstaben“ (cap="all"), der Text bleibt original."""
    runs = [r for sh in _walk(deck.slides[VALIDIERUNG].shapes) if sh.has_text_frame
            for p in sh.text_frame.paragraphs for r in p.runs if r.text == "Nicht im Modell"]
    assert runs and runs[0]._r.get_or_add_rPr().get("cap") == "all"


def test_icon_scheiben_wie_im_html(deck, specs):
    """Jede span.disc des HTML wird eine native Kreisform mit Icon-PNG (52 % des Durchmessers)."""
    for slide, spec in zip(deck.slides, specs):
        discs = [sh for sh in _walk(slide.shapes) if _is_disc(sh)]
        assert len(discs) == spec.html.count('class="disc"'), spec.label
        for grp in discs:
            ov, pic = list(grp.shapes)
            d = ov.width + (2 * PX if ov.line.fill.type is not None and ov.line.width else 0)
            assert abs(pic.width - 0.52 * d) <= 2 * PX, (spec.label, grp.name)


def test_icon_pngs_liegen_im_asset_ordner():
    icons = list((ROOT / "docs" / "presentation" / "kiko" / "assets" / "icons").glob("*__*.png"))
    assert len(icons) >= 50
    assert all(re.fullmatch(r"[a-z0-9-]+__[0-9A-F]{6}", p.stem) for p in icons)


def test_tracker_zeigt_aktive_schritte_wie_html(deck):
    expected = {VALIDIERUNG: {1}, VLS: {2, 3}, MODELL: {2, 3}, SCHWELLE: {4, 5, 6}, KALIBRIERUNG: {5, 6},
                BAND: {5, 6}, FALL: {4, 5, 6}, DASHBOARD: {6}}
    steps = HTML["STEPS"]
    for i, slide in enumerate(deck.slides):
        trackers = _named(slide, "Tracker ")
        if i not in expected:
            assert not trackers, i
            continue
        assert len(trackers) == 6, i
        active = {int(sh.name.split()[1]) for sh in trackers if sh.name.endswith(" aktiv")}
        assert active == expected[i], i
        label = next(sh for sh in slide.shapes if sh.name == "Tracker-Label").text_frame.text
        assert label.endswith(" · ".join(steps[k - 1] for k in sorted(active)))


def test_zaehlwerke_mit_ziffernzellen(deck):
    cases = {FADEN: ("23.563", "glass"), MODELL: ("−15,9%", "glass"), BAND: ("9,5", "glass"), FALL: ("+2,30", "red")}
    for i, slide in enumerate(deck.slides):
        groups = _named(slide, "Zählwerk ")
        if i not in cases:
            assert not groups, i
            continue
        value, tone = cases[i]
        (grp,) = groups
        assert grp.name == f"Zählwerk {value}"
        cells = [sh for sh in grp.shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
        assert len(cells) == sum(c.isdigit() for c in value)
        if tone == "red":
            assert {str(c.fill.fore_color.rgb) for c in cells} == {"B3261E"}


def test_canvas_folie_mit_allen_feldern(deck):
    text = _texts(deck.slides[CANVAS])
    for nr, (head, body, owner) in HTML["CANVAS_SHORT"].items():
        assert head in text and body in text, nr
        assert HTML["CANVAS"][nr][0] in text, nr
    assert "10 von 10 Feldern" in text
    dots = _named(deck.slides[CANVAS], "Canvas-Feld ")
    assert sorted(sh.name[-2:] for sh in dots) == sorted(HTML["CANVAS"])
    assert len(_named(deck.slides[CANVAS], "Chevron")) == 7
    assert _named(deck.slides[CANVAS], "Rücksprung zu 01")


def test_backup_canvas_im_wortlaut_von_tabelle_a1(deck):
    text = _texts(deck.slides[B1])
    for nr, wording in HTML["CANVAS_A1"].items():
        assert wording in text, nr
    assert "Kohärenz" in text


def test_fusszeile_mit_canvas_bezug_wie_html(deck, specs):
    for slide, spec in zip(deck.slides, specs):
        refs = re.findall(r'<span class="cref">.*?<b>(\d\d)</b>([^<]+)</span>', spec.html)
        text = _texts(slide)
        for nr, name in refs:
            assert nr in text and name.strip() in text, (spec.label, nr)
        assert "Verbrauchsprognose & Frühwarnung" in text and "Kiko" in text


def test_kernzahlen_stehen_auf_den_folien(deck):
    text = " ".join(_texts(s) for s in deck.slides)
    for value in ["13.272", "9.188", "3.725", "−15,9", "144,4 VLS-h", "+2,30", "ZL-00147", "−6,2 %",
                  "114", "107", "Beispieleingabe", "Stichprobe unauffälliger Fälle", "eingefroren", "175,1 VLS-h",
                  "480,9 − 149,1 = 331,8", "Nicht im Modell", "Ø 9,5 je Monat"]:
        assert value in text, value


# --------------------------------------------------------------------------- Diagramme


def _charts(slide):
    return [sh.chart for sh in slide.shapes if sh.has_chart]


def test_diagramme_sind_nativ_und_schemakonform(deck):
    per_slide = {i: len(_charts(s)) for i, s in enumerate(deck.slides)}
    assert per_slide[VLS] == 2 and per_slide[MODELL] == 2
    assert all(per_slide[i] == 1 for i in (SCHWELLE, BAND, FALL, B3, B4, B6))
    for slide in deck.slides:
        for chart in _charts(slide):
            for el in chart._chartSpace.iter(qn("c:axId"), qn("c:crossAx")):
                assert int(el.get("val")) >= 0, "Achsen-ID muss xsd:unsignedInt sein"
            assert chart.has_title is False


def test_streudiagramm_nutzt_standard_markerformat(deck):
    """Nur-Marker-Serien wie Excel: Linie noFill, eigenes Markerformat, keine QA-Variante."""
    (band,) = _charts(deck.slides[BAND])
    plot = band.plots[0]
    assert plot._element.find(qn("c:scatterStyle")).get("val") == "lineMarker"
    for series in list(plot.series)[:3]:
        ln = series._element.find(qn("c:spPr")).find(qn("a:ln"))
        assert ln.find(qn("a:noFill")) is not None, series.name
        assert series._element.find(qn("c:marker")).find(qn("c:symbol")).get("val") != "none"
    points = sum(len(list(s.values)) for s in list(plot.series)[:3])
    assert points == 8284 + 114 - 4 + 4  # alle Zähler-Monate im Ausschnitt plus Hinweise am Rand
    assert _named(deck.slides[BAND], "Toleranzband")


def test_fold_rauten_farbig(deck):
    names = [sh.name for sh in deck.slides[MODELL].shapes if sh.name.startswith("Fold ")]
    assert len(names) == 12  # 4 Kandidaten × 3 Folds
    colors = {str(sh.fill.fore_color.rgb) for sh in deck.slides[MODELL].shapes if sh.name.startswith("Fold ")}
    assert colors == {c.lstrip("#").upper() for c in HTML["FOLD_COLORS"]}


def test_treiber_diagramm_mit_berichtswerten(deck):
    (chart,) = _charts(deck.slides[B3])
    plot = chart.plots[0]
    assert list(plot.categories) == [name for name, _ in HTML["BERICHT"]["importance"]]
    assert list(plot.series[0].values) == [value for _, value in HTML["BERICHT"]["importance"]]


def test_pruefaufwand_als_saeulen_mit_durchschnittslinie(deck, facts):
    (chart,) = _charts(deck.slides[B6])
    bars, avg = list(chart.plots)
    assert bars._element.tag == qn("c:barChart") and avg._element.tag == qn("c:lineChart")
    assert list(bars.series[0].values) == facts["alerts_by_month"]
    assert list(avg.series[0].values) == pytest.approx([facts["n_alerts"] / 12] * 12)


def test_fallbeispiel_rechenweg_und_faktor(deck):
    text = _texts(deck.slides[FALL])
    for value in ["Ist − Prognose", "Residuum", "in VLS (49 kW)", "÷ Schwelle q99", "Aug: Faktor +2,30",
                  "September zum Vergleich"]:
        assert value in text, value


# --------------------------------------------------------------------------- Auslieferung


def test_ausgelieferte_datei_ist_aktuell(deck):
    if not SHIPPED.exists():
        pytest.skip("Noch nicht gebaut (python scripts/build_kiko_pptx.py)")
    prs = Presentation(SHIPPED)
    assert len(prs.slides) == 17
    (band,) = _charts(prs.slides[BAND])
    assert band.plots[0]._element.find(qn("c:scatterStyle")).get("val") == "lineMarker", "QA-Variante ausgeliefert"
    assert "Abweichung 2,3-mal so groß wie erlaubt" in _texts(prs.slides[FALL])
    assert [_texts(s) for s in prs.slides] == [_texts(s) for s in deck.slides], "PPTX neu bauen"


# --------------------------------------------------------------------------- Helfer


def test_rgba_und_hex_kurzform():
    rgba = PPTX["_rgba"]
    assert rgba("#fff") == ("#FFFFFF", None)
    assert rgba("#00718e") == ("#00718E", None)
    assert rgba("rgba(255,255,255,.07)") == ("#FFFFFF", pytest.approx(0.07))
    assert PPTX["mix"]("#6CC0D2", "#063659", .5) == "#397B96"


def test_pfadparser_absolut_relativ_und_boegen():
    points = PPTX["_path_points"]
    assert points("M1076 298 V306 H203 V314") == ([(1076, 298), (1076, 306), (203, 306), (203, 314)], False)
    pts, closed = points("M10,20 L15,25 L10,30 L5,25 z")
    assert closed and pts[-1] == (5, 25)
    pts, closed = points("M0,0 h10 a4,4 0 0 1 4,4 v2 z")  # Bogen als Sehne bis zum Endpunkt
    assert pts == [(0, 0), (10, 0), (14, 4), (14, 6)] and closed


def test_rich_fasst_leerraum_zusammen_und_bricht_bei_br():
    import lxml.html

    el = lxml.html.fragment_fromstring(
        '<p>weniger RMSE  als die Regel<br><span class="mono">Test 2025: 10.922 → 9.188 kWh</span></p>')
    paras = PPTX["rich"](el, [(".mono", {"font": "Geist Mono"})])
    assert [[t for t, _ in p] for p in paras] == [["weniger RMSE als die Regel"], ["Test 2025: 10.922 → 9.188 kWh"]]
    assert paras[1][0][1] == {"font": "Geist Mono"}


def test_icon_und_scheibenton_werden_aus_dem_html_erkannt():
    import lxml.html

    for name, tone, d in [("gauge", "teal", 34), ("history", "off", 26), ("lightbulb", "lamber", 30),
                          ("database", "ring", 36), ("shield-check", "glass", 28)]:
        span = lxml.html.fragment_fromstring(HTML["disc"](name, tone, d))
        assert PPTX["disc_of"](span) == PPTX["Disc"](name, tone, d)


def test_textmass_nutzt_ersatzschrift_fuer_pfeile():
    measure = PPTX["measure"]
    assert measure("→", 16) > 0.7 * 16  # DejaVu-Pfeil statt .notdef der Latin-Untermenge
    assert measure("Kiko", 16, bold=True) > measure("Kiko", 16)
    assert measure("ABC", 13, tracking=.08) == pytest.approx(measure("ABC", 13) + 3 * .08 * 13)
    assert PPTX["wrap_lines"]("Keine Precision und Recall", 240, 19, True) == ["Keine Precision und Recall"]


def test_zaehlwerk_setzt_ziffern_in_zellen():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    end = PPTX["zw"](slide, 0, 0, "23.563", "navy", 50)
    (grp,) = slide.shapes
    cells = [sh for sh in grp.shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
    assert len(cells) == 5 and end == pytest.approx(5 * .74 * 50 + .34 * 50 + 5 * .06 * 50)
    assert [sh.text_frame.text for sh in grp.shapes if sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX] == list("23.563")


def test_ring_bleibt_im_panel():
    prs = Presentation()
    prs.slide_width, prs.slide_height = PPTX["E"](1280), PPTX["E"](720)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    ring = PPTX["ring_arc"](slide, 502, 20, 209, (0, 0, 512, 653), "#6CC0D2", 2)
    assert ring.left >= 0 and ring.top >= 0 and ring.left + ring.width <= 512 * PX + PX


def test_fehlende_icons_werden_gesammelt_und_gerendert(tmp_path):
    module = PPTX["icon_png"].__globals__  # runpy liefert eine Kopie; die Funktionen lesen ihr eigenes Modul
    module["_MISSING"].clear()
    original = module["ICON_PNG_DIR"]
    module["ICON_PNG_DIR"] = tmp_path
    try:
        placeholder = PPTX["icon_png"]("gauge", "#084878")
        assert isinstance(placeholder, io.BytesIO)
        assert module["_MISSING"] == {("gauge", "#084878")}
        if shutil.which("node") is None:
            pytest.skip("Node fehlt: Rendern nicht prüfbar")
        PPTX["render_icons"](module["_MISSING"])
        assert Path(PPTX["icon_png"]("gauge", "#084878")) == tmp_path / "gauge__084878.png"
    finally:
        module["ICON_PNG_DIR"] = original
        module["_MISSING"].clear()
