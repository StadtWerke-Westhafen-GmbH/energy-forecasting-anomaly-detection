"""Protect colour semantics, generated outputs, and the actual chart APIs."""

import importlib.util
import json
from pathlib import Path

import numpy as np
import pandas as pd
import plotly.io as pio
import pytest
from energy_analytics.visualization import eda, theme

ROOT = Path(__file__).resolve().parents[1]


def luminance(hex_color):
    rgb = [int(hex_color[i : i + 2], 16) / 255 for i in (1, 3, 5)]
    linear = [v / 12.92 if v <= 0.04045 else ((v + 0.055) / 1.055) ** 2.4 for v in rgb]
    return sum(c * w for c, w in zip(linear, [0.2126, 0.7152, 0.0722]))


def contrast(a, b):
    light, dark = sorted([luminance(a), luminance(b)], reverse=True)
    return (light + 0.05) / (dark + 0.05)


@pytest.mark.parametrize(
    "fg,bg",
    [
        ("text-primary", "surface-card"),
        ("text-secondary", "surface-page"),
        ("text-muted", "surface-page"),
        ("text-inverse", "navy-700"),
        ("text-inverse", "teal-600"),
        ("text-inverse", "status-ok"),
        ("green-700", "status-ok-bg"),
        ("amber-700", "status-warn-bg"),
        ("red-600", "status-critical-bg"),
        ("teal-700", "status-info-bg"),
        ("status-critical-inverse", "navy-800"),
        ("status-ok-inverse", "navy-800"),
    ],
)
def test_text_contrast(fg, bg):
    assert contrast(theme.TOKENS[fg], theme.TOKENS[bg]) >= 4.5


def test_generated_outputs_are_current():
    spec = importlib.util.spec_from_file_location("build_tokens", ROOT / "scripts/build_tokens.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    for path, content in mod.compile_outputs().items():
        assert path.read_text(encoding="utf-8") == content, str(path)
    with pytest.raises(ValueError, match="Cyclic"):
        mod.resolve_tokens({"one": "var(--two)", "two": "var(--one)"})
    with pytest.raises(ValueError, match="Unknown"):
        mod.resolve_tokens({"one": "var(--missing)"})


def test_web_and_python_themes_are_identical():
    web = json.loads(
        (ROOT / "brand/design-system/dist/tokens.resolved.json").read_text(encoding="utf-8")
    )
    assert web["roles"] == theme.ROLE
    assert web["customers"] == theme.KUNDENTYP_COLORS
    assert web["plotly"] == theme.template().to_plotly_json()
    # Forecasts/alerts must never acquire meaning as arbitrary categorical colours.
    assert theme.ROLE["anomalie"] not in theme.QUALITATIVE
    assert theme.ROLE["prognose"] not in theme.QUALITATIVE


def test_forecast_helper_supports_iterators_and_validates_bounds():
    traces = theme.forecast_traces(iter(["a", "b"]), [1, 2], [1.1, 1.9], [0.5, 1.5], [1.5, 2.5])
    assert list(traces[0].x) == ["a", "b", "b", "a"]
    assert traces[-1].line.dash == "4,2"
    with pytest.raises(ValueError):
        theme.forecast_traces([1], [2], [3], lower=[1])
    with pytest.raises(ValueError):
        theme.forecast_traces([1], [2], [3], [4], [1])


def test_all_fifteen_chart_builders():
    eda.setup()
    x = ["01/2025", "02/2025", "03/2025"]
    values = [100.0, 120.0, 115.0]
    groups = {"Gewerbe": values, "Industrie": [200.0, 240.0, 230.0]}
    frames = [
        eda.timeseries_forecast(x, values, [99, 119, 116]),
        eda.seasonality(x, groups),
        eda.histogram(groups),
        eda.boxplot(groups),
        eda.scatter_trend({"Gewerbe": ([1, 2, 3], values)}),
        eda.heatmap(x, ["Gewerbe"], [values]),
        eda.correlation(pd.DataFrame([[1, 0.2], [0.2, 1]], columns=["A", "B"])),
        eda.residual_hist([0, 1, -2], 1),
        eda.residual_bars(x, [0, 1, -2], 1),
        eda.feature_importance(["A", "B"], [0.7, 0.3]),
        eda.missing_values(["A", "B"], [1, 5]),
        eda.parity(values, [99, 121, 114]),
        eda.anomaly_stack(x, groups),
        eda.top_n(["A", "B"], [1, 2]),
        eda.share_area(x, groups),
    ]
    for fig in frames:
        assert fig.data
        assert fig.layout.template.layout.font.family == theme.FONT
        pio.to_json(fig, validate=True)
    assert frames[8].data[0].marker.color[-1] == theme.ROLE["anomalie"]
    assert np.isclose(frames[9].data[0].x[-1], 0.7)


def test_matplotlib_draws_and_uses_local_font(tmp_path):
    import matplotlib

    matplotlib.use("Agg")
    plt = theme.apply_matplotlib()
    fig, ax = plt.subplots()
    ax.plot([1, 2, 3], [1000, 2000, 3000])
    theme.format_matplotlib_axes(ax)
    assert ax.yaxis.get_major_formatter()(12000, 0) == "12.000"
    assert ax.lines[0].get_color() == theme.QUALITATIVE[0]
    fig.savefig(tmp_path / "chart.png")
    plt.close(fig)


def test_chart_builders_do_not_require_global_template():
    previous = pio.templates.default
    try:
        pio.templates.default = "plotly"
        fig = eda.timeseries_forecast([1, 2], [0.1, 0.2])
        assert fig.layout.template.layout.font.family == theme.FONT
        assert fig.layout.template.layout.yaxis.tickformat == ",~g"
        assert pio.templates.default == "plotly"
    finally:
        pio.templates.default = previous


def test_slide_export_does_not_mutate_notebook_figure(monkeypatch):
    import plotly.graph_objects as go

    fig = eda.timeseries_forecast([1, 2], [100, 200])
    before = pio.to_json(fig)
    captured = []
    monkeypatch.setattr(
        go.Figure, "write_image", lambda self, *args, **kwargs: captured.append(self)
    )
    eda.save_for_slide(fig, "unused.png")
    assert pio.to_json(fig) == before
    assert captured[0].layout.font.size == 20


def test_office_and_notebook_templates():
    import zipfile
    import nbformat
    from pptx import Presentation
    from docx import Document

    folder = ROOT / "brand/templates"
    prs = Presentation(folder / "presentations/sww-project-template.pptx")
    assert len(prs.slides) == 10
    assert any(shape.has_chart for slide in prs.slides for shape in slide.shapes)
    with zipfile.ZipFile(folder / "presentations/sww-project-template.potx") as z:
        assert b"presentationml.template.main+xml" in z.read("[Content_Types].xml")
    assert len(Document(folder / "documents/sww-report-template.docx").tables) == 1
    nb = nbformat.read(folder / "notebooks/sww-analysis-template.ipynb", as_version=4)
    nbformat.validate(nb)
    assert all(not c.get("outputs") for c in nb.cells)
