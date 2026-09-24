"""Erstellt eine Berichtskopie mit genau drei ausgearbeiteten ML-Bereichen.

Geaendert werden ausschliesslich:
- ML-Canvas-Feld 6 "Evaluation / Ueberpruefung"
- Kapitel 4.3 "Ergebnisse"
- Kapitel 4.4 "Bewertung der Ergebnisqualitaet"

Die Quelldatei bleibt unveraendert. Alle anderen Platzhalter und Inhalte der
Berichtsvorlage werden bewusst erhalten.
"""

from __future__ import annotations

from hashlib import sha256
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "Berichtsvorlage_IHK(3).docx"
OUTPUT = (
    ROOT
    / "docs"
    / "Berichtsvorlage_IHK(3)_Kapitel_4_3_4_4_und_Canvas_6.docx"
)
EXPECTED_SOURCE_SHA256 = (
    "a02fa0f2903116a5f0cadcb56a854adf7da5c0f4f50190ba17476b294d1e47e0"
)


CANVAS_EVALUATION = (
    "Evaluation / Überprüfung: Die Prognosemodelle werden in drei zeitlich "
    "vorwärts laufenden Folds des Jahres 2024 geprüft; spätere Monate gelangen "
    "dabei nie in frühere Trainingsdaten. Hauptmetrik ist der RMSE in kWh, "
    "ergänzt um MAE und R². Als Baselines dienen der Vormonatswert und das "
    "Mittel aus bis zu drei Vormonaten. Die gewählte Modellkonfiguration wird "
    "anschließend retrospektiv und monatsweise auf 2025 bewertet. Die "
    "Anomalieschwelle gehört nicht zum Modelltraining: Sie wird aus 1.397 "
    "absoluten VLS-Prognosefehlern von November und Dezember 2024 als 99. "
    "Perzentil bei 144,4 VLS-Stunden kalibriert. Da bestätigte fachliche Labels "
    "fehlen, werden noch keine Precision oder Recall ausgewiesen."
)


RESULTS_BLOCK = [
    (
        "Heading 3",
        "4.3.1   Zielvariable und zeitliche Validierung",
    ),
    (
        "normal",
        "Die bereinigte Datengrundlage umfasst 16.800 Zähler-Monate von 700 "
        "Zählern aus den Jahren 2024 und 2025. Als interne Zielgröße dienen "
        "Vollaststunden (VLS). Sie ergeben sich aus dem Monatsverbrauch in kWh "
        "geteilt durch die Vertragsleistung in kW und machen Zähler "
        "unterschiedlicher Größe vergleichbarer. Für die fachliche Bewertung "
        "werden die Prognosen wieder mit der jeweiligen Vertragsleistung in "
        "kWh zurückgerechnet.",
    ),
    (
        "normal",
        "Da es sich um Zeitreihendaten handelt, erfolgte keine zufällige "
        "Aufteilung. Drei vorwärts laufende Folds bildeten die spätere Anwendung "
        "nach: Fold 1 lernte aus Januar bis April 2024 und bewertete Mai bis "
        "Juni, Fold 2 lernte aus Januar bis Juni und bewertete Juli bis August, "
        "Fold 3 lernte aus Januar bis August und bewertete September bis "
        "Oktober. Hauptmetrik war der RMSE in kWh; MAE und R² ergänzten die "
        "Einordnung. Als fachliche Baselines dienten der Vormonatswert und das "
        "Mittel aus bis zu drei Vormonaten.",
    ),
    (
        "Heading 3",
        "4.3.2   Modellwahl im Jahr 2024",
    ),
    (
        "normal",
        "Für den Random Forest wurden acht kontrollierte "
        "Hyperparameterkombinationen ausschließlich innerhalb dieser drei Folds "
        "verglichen. Die beste Kombination verwendet 300 Bäume, eine maximale "
        "Tiefe von 8, mindestens 5 Beobachtungen je Blatt, 70 Prozent der "
        "Merkmale je Aufteilung und den Zufallsstartwert 42.",
    ),
    (
        "normal",
        "Mit einem mittleren CV-RMSE von 13.272 kWh erreicht der Random Forest "
        "den niedrigsten Fehler. Es folgen die lineare Regression mit 13.643 "
        "kWh, das Mittel aus bis zu drei Vormonaten mit 15.483 kWh und der "
        "Vormonat mit 18.460 kWh. Damit liegt der Random Forest 2,7 Prozent vor "
        "dem linearen Modell und 14,3 Prozent vor der stärkeren Baseline. Seine "
        "Standardabweichung über die Folds beträgt 3.843 kWh und zeigt, dass die "
        "Ergebnisgüte vom Bewertungszeitraum abhängt.",
    ),
    (
        "Heading 3",
        "4.3.3   Retrospektiver Benchmark 2025",
    ),
    (
        "normal",
        "Die gewählte Konfiguration wurde anschließend in einem retrospektiven "
        "One-Step-Ahead-Benchmark auf 8.398 auswertbaren Zähler-Monaten des "
        "Jahres 2025 geprüft. Für jeden Monat wurden nur zu diesem Zeitpunkt "
        "verfügbare Merkmale verwendet.",
    ),
    (
        "normal",
        "Der Random Forest erzielt 9.188 kWh RMSE, 3.725 kWh MAE und ein R² von "
        "0,901. Die lineare Regression erreicht 9.414 kWh, das "
        "Drei-Monats-Mittel 10.922 kWh und der Vormonat 12.386 kWh RMSE. Der "
        "Vorteil beträgt damit 15,9 Prozent gegenüber der besten Baseline und "
        "2,4 Prozent gegenüber der linearen Regression.",
    ),
    (
        "normal",
        "Ein direkt auf kWh trainierter Random Forest erreicht 9.800 kWh RMSE. "
        "Die VLS-Variante mit Rückrechnung liegt mit 9.188 kWh um 6,2 Prozent "
        "niedriger. Dies belegt keine allgemeine Überlegenheit von VLS, spricht "
        "für diesen Datensatz jedoch für die Normalisierung nach "
        "Vertragsleistung.",
    ),
    (
        "Heading 3",
        "4.3.4   Einfluss der Merkmalsgruppen",
    ),
    (
        "normal",
        "Die gruppierte Permutationsanalyse misst, wie stark der RMSE steigt, "
        "wenn die Information einer Merkmalsgruppe zufällig gemischt wird. Den "
        "größten Anstieg verursacht die Verbrauchshistorie mit 7.256 kWh. Es "
        "folgen Produktionsplan mit 798 kWh, geplante Wartung mit 591 kWh, "
        "Kalender mit 276 kWh, Wetterprognose mit 154 kWh, Jahreszeit mit 121 "
        "kWh und Kundentyp mit 3 kWh. Die Werte beschreiben Prognosebeiträge, "
        "jedoch keine Kausalität; korrelierte Gruppen können sich Bedeutung "
        "teilen.",
    ),
    (
        "Heading 3",
        "4.3.5   Residuenbasierte Prüfhinweise",
    ),
    (
        "normal",
        "Die Anomalieschwelle wurde getrennt vom Modelltraining aus 1.397 "
        "absoluten Out-of-Fold-Fehlern in VLS für November und Dezember 2024 "
        "bestimmt. Das 99. Perzentil liegt bei 144,4 VLS-Stunden; rund 99 Prozent "
        "der Kalibrierungsfehler sind höchstens so groß. Ein Zähler-Monat erhält "
        "einen Prüfhinweis, wenn sein absoluter Fehler diese feste Grenze "
        "erreicht oder überschreitet. Die zugehörige kWh-Grenze hängt von der "
        "Vertragsleistung des Zählers ab.",
    ),
    (
        "normal",
        "Damit entstehen 2025 insgesamt 114 Hinweise für 107 Zähler, "
        "durchschnittlich 9,5 pro Monat; 68 Abweichungen liegen oberhalb und 46 "
        "unterhalb der Prognose. Zehn Hinweise überschneiden sich mit "
        "Datenqualitätskennzeichen. Die Sensitivitätsanalyse ergibt bei 95, "
        "97,5, 99 und 99,5 Prozent insgesamt 400, 256, 114 beziehungsweise 69 "
        "Hinweise. Das 99. Perzentil wurde als handhabbarer Pilotwert gewählt, "
        "nicht als mathematisch einzig richtige Grenze.",
    ),
]


QUALITY_BLOCK = [
    (
        "Heading 3",
        "4.4.1   Aussagekraft der Prognoseergebnisse",
    ),
    (
        "normal",
        "Der Random Forest liegt sowohl in der zeitlichen Validierung 2024 als "
        "auch im Benchmark 2025 vor beiden Baselines. Dies zeigt zusätzlichen "
        "Prognosenutzen gegenüber einfachen Fortschreibungen. Ein eindeutiges "
        "Overfitting-Signal ist nicht erkennbar: Die Baumtiefe ist auf acht "
        "begrenzt, und der Vorteil bleibt im späteren Jahr bestehen. Wegen der "
        "kurzen Zeitreihe lässt sich Overfitting dennoch nicht vollständig "
        "ausschließen.",
    ),
    (
        "normal",
        "Das R² von 0,901 bedeutet, dass das Modell einen großen Anteil der "
        "Streuung im Benchmarkjahr erklärt; es ist weder eine Trefferquote noch "
        "eine Anomaliewahrscheinlichkeit. Dass der RMSE deutlich über dem MAE "
        "liegt, weist auf wenige besonders große Fehler hin. Genau diese Fälle "
        "sind für die nachgelagerte Prüfung relevant.",
    ),
    (
        "Heading 3",
        "4.4.2   Grenzen der Anomaliebewertung",
    ),
    (
        "normal",
        "Der Benchmark 2025 ist retrospektiv und kein künftig ungesehener "
        "Blindtest. Seine Werte belegen die Leistung im gewählten "
        "Versuchsaufbau, garantieren jedoch keine unveränderte Produktivleistung. "
        "Zudem stützt sich die Schwelle nur auf zwei Wintermonate. Saisonale "
        "Veränderungen können deshalb zu einer anderen Fehlerverteilung führen.",
    ),
    (
        "normal",
        "Für die Prüfhinweise liegen keine vollständigen fachlichen Labels vor. "
        "Precision, Recall sowie Falsch-Positiv- und Falsch-Negativ-Rate können "
        "daher nicht belastbar berechnet werden. Ein großer Modellfehler ist "
        "zunächst eine Modellanomalie, nicht automatisch ein Defekt oder eine "
        "fehlerhafte Abrechnung. Auch die zehn Überschneidungen mit "
        "Datenqualitätskennzeichen sind ein Plausibilitätshinweis, aber kein "
        "Gütenachweis. Die starke Bedeutung der Verbrauchshistorie zeigt "
        "außerdem ein Kaltstart-Risiko; frühe Zeiträume ohne Historie aus 2023 "
        "sind vorsichtig zu interpretieren.",
    ),
    (
        "Heading 3",
        "4.4.3   Bewertung für den geplanten Einsatz",
    ),
    (
        "normal",
        "Für einen kontrollierten Pilotbetrieb ist der Ansatz geeignet, wenn er "
        "als Priorisierungshilfe eingesetzt wird. Das Modell erzeugt keine "
        "automatische Fachentscheidung, sondern markiert Zähler-Monate für die "
        "Prüfung von Datenqualität, Produktion, Wartung und weiteren "
        "Kontextinformationen.",
    ),
    (
        "normal",
        "Empfohlen wird ein prospektiver Schattenbetrieb mit neuen Monaten. "
        "Dabei sollten Fehlerniveau, Merkmalsverteilungen, Hinweisvolumen und "
        "Bearbeitungszeit überwacht sowie bestätigte und verworfene Hinweise "
        "strukturiert gespeichert werden. Erst diese Ground Truth ermöglicht "
        "Precision und Recall, eine fachlich optimierte Schwelle und "
        "gegebenenfalls ein zusätzliches Klassifikationsmodell.",
    ),
]


def file_sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def find_paragraph(doc: Document, exact_text: str):
    matches = [p for p in doc.paragraphs if p.text.strip() == exact_text]
    if len(matches) != 1:
        raise ValueError(
            f"Erwartet genau einen Absatz {exact_text!r}, gefunden: {len(matches)}"
        )
    return matches[0]


def clear_paragraph(paragraph) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)


def set_text(paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.add_run(text)


def add_after(doc: Document, anchor, text: str, style: str):
    paragraph = doc.add_paragraph(style=style)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.add_run(text)
    anchor._p.addnext(paragraph._p)
    return paragraph


def delete_paragraph(paragraph) -> None:
    parent = paragraph._p.getparent()
    parent.remove(paragraph._p)


def replace_section_block(
    doc: Document,
    heading_text: str,
    next_heading_text: str,
    expected_placeholder_start: str,
    block: list[tuple[str, str]],
) -> None:
    heading = find_paragraph(doc, heading_text)
    next_heading = find_paragraph(doc, next_heading_text)
    body = heading._p.getparent()
    start = body.index(heading._p)
    end = body.index(next_heading._p)
    between = list(body)[start + 1 : end]
    visible_text = "".join(
        "".join(node.text or "" for node in element.iter()) for element in between
    )
    if expected_placeholder_start not in visible_text:
        raise AssertionError(
            f"Unerwarteter Ausgangsinhalt zwischen {heading_text!r} und "
            f"{next_heading_text!r}"
        )
    for element in between:
        body.remove(element)
    anchor = heading
    for style, text in block:
        anchor = add_after(doc, anchor, text, style)


def validate(doc: Document) -> None:
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    required = [
        "Evaluation / Überprüfung: Die Prognosemodelle werden in drei zeitlich",
        "4.3.1   Zielvariable und zeitliche Validierung",
        "13.272 kWh",
        "9.188 kWh",
        "R² von 0,901",
        "1.397 absoluten Out-of-Fold-Fehlern",
        "Verbrauchshistorie mit 7.256 kWh",
        "114 Hinweise",
        "4.4.2   Grenzen der Anomaliebewertung",
        "Precision und Recall",
        "prospektiver Schattenbetrieb",
    ]
    for phrase in required:
        if phrase not in text:
            raise AssertionError(f"Pflichtinhalt fehlt: {phrase}")
    forbidden = [
        "Evaluation / Überprüfung: [Ihre Antwort.]",
        "[Beispielsatz: Auf dem Testdatensatz",
        "[Beispielsatz: Die Ergebnisqualität",
        "99,7",
        "99.7",
        "WAPE",
    ]
    for phrase in forbidden:
        if phrase in text:
            raise AssertionError(f"Unerwünschter Inhalt gefunden: {phrase}")
    if len(doc.sections) != 1:
        raise AssertionError("Die Abschnittsstruktur wurde unerwartet verändert")
    if len(doc.tables) != 3:
        raise AssertionError("Die vorhandenen Tabellen wurden unerwartet verändert")

    paragraphs = doc.paragraphs
    results_start = next(
        index
        for index, paragraph in enumerate(paragraphs)
        if paragraph.text.strip() == "4.3   Ergebnisse"
    )
    results_end = next(
        index
        for index, paragraph in enumerate(paragraphs)
        if paragraph.text.strip() == "4.5   Visualisierungen und Dashboard"
    )
    report_words = sum(
        len(paragraph.text.split())
        for paragraph in paragraphs[results_start:results_end]
    )
    if report_words > 900:
        raise AssertionError(
            f"Kapitel 4.3 und 4.4 sind mit {report_words} Wörtern zu lang"
        )


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    before_hash = file_sha256(SOURCE)
    if before_hash != EXPECTED_SOURCE_SHA256:
        raise AssertionError(
            "Die Berichtsvorlage hat sich seit der Prüfung verändert; bitte neu prüfen."
        )

    doc = Document(SOURCE)
    evaluation = find_paragraph(doc, "Evaluation / Überprüfung: [Ihre Antwort.]")
    set_text(evaluation, CANVAS_EVALUATION)

    replace_section_block(
        doc,
        "4.3   Ergebnisse",
        "4.4   Bewertung der Ergebnisqualität",
        "[Beispielsatz: Auf dem Testdatensatz",
        RESULTS_BLOCK,
    )
    replace_section_block(
        doc,
        "4.4   Bewertung der Ergebnisqualität",
        "4.5   Visualisierungen und Dashboard",
        "[Beispielsatz: Die Ergebnisqualität",
        QUALITY_BLOCK,
    )

    validate(doc)
    doc.save(OUTPUT)

    reopened = Document(OUTPUT)
    validate(reopened)
    if file_sha256(SOURCE) != before_hash:
        raise AssertionError("Die Quelldatei wurde verändert")

    print(f"Erstellt: {OUTPUT}")
    print(f"Quelle unverändert: {before_hash}")
    print(f"Absätze: {len(reopened.paragraphs)} | Tabellen: {len(reopened.tables)}")


if __name__ == "__main__":
    main()
