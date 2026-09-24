from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "IHK_Bericht_Gruppe_6_Patrick.docx"
OUTPUT = ROOT / "docs" / "IHK_Bericht_Gruppe_6_Patrick_ergaenzt.docx"


NAVY = "17365D"
PALE_BLUE = "EAF2F8"
WHITE = "FFFFFF"
LIGHT_GRAY = "D9D9D9"
TEXT = "000000"


def find_paragraph(document: Document, exact_text: str) -> Paragraph:
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text.strip() == exact_text]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph {exact_text!r}, found {len(matches)}")
    return matches[0]


def replace_paragraph(paragraph: Paragraph, text: str, *, bold: bool = False) -> Paragraph:
    paragraph.clear()
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    paragraph.style = "normal"
    return paragraph


def clear_highlight(paragraph: Paragraph) -> None:
    """Remove draft highlighting while preserving all other run formatting."""
    for run in paragraph.runs:
        r_pr = run._element.get_or_add_rPr()
        highlight = r_pr.find(qn("w:highlight"))
        if highlight is not None:
            r_pr.remove(highlight)


def replace_text_in_runs(paragraph: Paragraph, old: str, new: str) -> None:
    """Replace text without flattening existing inline formatting."""
    if old not in paragraph.text:
        return
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            return
    raise RuntimeError(f"Text spans multiple runs and cannot be replaced safely: {old!r}")


def compact_cover_spacing(document: Document, blanks_to_keep: int = 3) -> None:
    """Keep the cover metadata table together on page one."""
    body = document._body._body
    first_table = next(element for element in body if element.tag == qn("w:tbl"))
    blank_paragraphs = []
    element = first_table.getprevious()
    while element is not None and element.tag == qn("w:p"):
        paragraph = Paragraph(element, document._body)
        if paragraph.text.strip():
            break
        blank_paragraphs.append(element)
        element = element.getprevious()

    for paragraph_element in blank_paragraphs[: max(0, len(blank_paragraphs) - blanks_to_keep)]:
        paragraph_element.getparent().remove(paragraph_element)


def keep_existing_table_rows_together(document: Document) -> None:
    """Prevent rows in source tables from being split across page boundaries."""
    for table in document.tables:
        for row in table.rows:
            keep_row_together(row)


def paragraph_after(
    document: Document,
    element,
    text: str = "",
    *,
    bold: bool = False,
    italic: bool = False,
    keep_with_next: bool = False,
) -> Paragraph:
    paragraph_element = OxmlElement("w:p")
    element.addnext(paragraph_element)
    paragraph = Paragraph(paragraph_element, document._body)
    paragraph.style = "normal"
    if text:
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = "Arial"
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    paragraph.paragraph_format.keep_with_next = keep_with_next
    return paragraph


def add_bullet_after(document: Document, element, text: str) -> Paragraph:
    paragraph = paragraph_after(document, element, text)
    paragraph.paragraph_format.left_indent = Cm(0.65)
    paragraph.paragraph_format.first_line_indent = Cm(-0.45)
    first_run = paragraph.runs[0]
    first_run.text = f"•\t{text}"
    return paragraph


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top=100, start=110, bottom=100, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), LIGHT_GRAY)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def format_table(table, widths_cm: list[float]) -> None:
    table.autofit = False
    set_table_borders(table)
    for row_index, row in enumerate(table.rows):
        keep_row_together(row)
        if row_index == 0:
            repeat_header(row)
        for col_index, cell in enumerate(row.cells):
            cell.width = Cm(widths_cm[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, NAVY if row_index == 0 else (PALE_BLUE if row_index % 2 == 0 else WHITE))
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
                    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor.from_string(WHITE if row_index == 0 else TEXT)
                    run.bold = row_index == 0


def table_after(document: Document, element, headers: list[str], rows: list[list[str]], widths_cm: list[float]):
    table = document.add_table(rows=1, cols=len(headers))
    for column, value in enumerate(headers):
        table.cell(0, column).text = value
    for values in rows:
        cells = table.add_row().cells
        for column, value in enumerate(values):
            cells[column].text = value
    format_table(table, widths_cm)
    element.addnext(table._tbl)
    return table


def update_sprint_three(document: Document) -> None:
    placeholder = find_paragraph(
        document,
        "[Sprint-Ziel, Aufgaben, Ergebnisse, Herausforderungen]",
    )
    replace_paragraph(placeholder, "Sprint-Ziel", bold=True)
    cursor = placeholder._p

    body = paragraph_after(
        document,
        cursor,
        "Ziel des dritten Sprints war es, aus der bereinigten Modellierungsbasis ein "
        "nachvollziehbares Regressionsmodell für die Folgemonatsprognose auszuwählen, "
        "zeitlich sauber zu bewerten und aus seinen späteren Prognosefehlern eine vor "
        "dem Benchmarkjahr festgelegte Schwelle für Prüfhinweise abzuleiten.",
    )
    cursor = body._p
    label = paragraph_after(document, cursor, "Aufgaben", bold=True, keep_with_next=True)
    cursor = label._p
    for item in [
        "Vollaststunden als interne Zielgröße bilden und jede Prognose für die Bewertung wieder in kWh zurückrechnen.",
        "Lineare Regression und Random Forest in drei vorwärts laufenden Zeit-Folds trainieren und mit Vormonat sowie dem Mittel aus bis zu drei Vormonaten vergleichen.",
        "Acht kontrollierte Hyperparameterkombinationen des Random Forest ausschließlich innerhalb der 2024er Folds prüfen.",
        "Die Fehlerschwelle mit November und Dezember 2024 kalibrieren, den retrospektiven Benchmark 2025 durchführen und die Ergebnisse für das Dashboard exportieren.",
    ]:
        bullet = add_bullet_after(document, cursor, item)
        cursor = bullet._p

    label = paragraph_after(document, cursor, "Ergebnisse", bold=True, keep_with_next=True)
    cursor = label._p
    result = paragraph_after(
        document,
        cursor,
        "Der Random Forest erzielte mit 13.272 kWh den niedrigsten mittleren "
        "Validierungs-RMSE und wurde als finales Modell ausgewählt. Die lineare "
        "Regression erreichte 13.643 kWh, die stärkere einfache Baseline 15.483 kWh. "
        "Im retrospektiven Benchmark 2025 lag der RMSE des Random Forest bei 9.188 kWh "
        "und damit 15,9 % unter der besten Baseline. Das 99. Perzentil der 1.397 "
        "Kalibrierungsfehler ergab eine Schwelle von 144,4 VLS-Stunden und 114 "
        "Prüfhinweise im Jahr 2025, durchschnittlich 9,5 pro Monat.",
    )
    cursor = result._p

    label = paragraph_after(document, cursor, "Herausforderungen", bold=True, keep_with_next=True)
    cursor = label._p
    paragraph_after(
        document,
        cursor,
        "Die zentrale Herausforderung war die strikte zeitliche Trennung von "
        "Modellauswahl, Schwellenkalibrierung und Benchmark. Zusätzlich musste die "
        "Verbrauchshistorie je Zähler ohne zukünftige Informationen neu berechnet "
        "werden. Da fachlich bestätigte Anomalie-Labels fehlen, werden die Ergebnisse "
        "bewusst als Prüfhinweise und nicht als erkannte Defekte bezeichnet.",
    )


def update_personal_outlook(document: Document) -> None:
    anchor = find_paragraph(
        document,
        "Ich habe gelernt, dass im Data-Science-Kontext die saubere Validierung wichtiger ist "
        "als die Wahl des komplexesten Modells. Besonders herausfordernd war das Vermeiden "
        "von Data Leakage. Das gelieferte 3-Monats-Mittel war fehlerhaft und musste je Zähler "
        "ausschließlich aus vergangenen Werten neu berechnet werden. Die Modellauswahl durfte "
        "nur auf den zeitlichen Prüfzeiträumen in 2024 beruhen. Wertvoll war die Erfahrung, "
        "dass eine einfache Baseline und eine lineare Regression als Referenz unverzichtbar "
        "sind. Nur daran ließ sich zeigen, dass der Random Forest einen echten Mehrwert bringt "
        "(RMSE 15,9 % unter der besten einfachen Baseline). Aus dem Tutor-Gespräch habe ich "
        "mitgenommen, dass für die Fachanwender eine verstandene Entscheidung mehr zählt als "
        "das letzte Promille Modellgüte.",
    )
    paragraph = paragraph_after(
        document,
        anchor._p,
        "Mein persönlicher Ausblick ist ein strukturierter Feedbackprozess: Nach jeder "
        "fachlichen Prüfung sollte gespeichert werden, ob der Hinweis tatsächlich eine "
        "relevante Anomalie war und welche Ursache vorlag. Sobald genügend belastbare "
        "Ja/Nein-Labels vorhanden sind, könnte das Regressionsmodell um ein "
        "Klassifikationsmodell ergänzt werden. Das Regressionsmodell würde weiterhin den "
        "erwarteten Verbrauch und die Abweichung bestimmen; das Klassifikationsmodell könnte "
        "anschließend eine kalibrierte Wahrscheinlichkeit in Prozent dafür schätzen, dass "
        "ein Prüfhinweis fachlich bestätigt wird. Eine solche Prozentangabe wäre erst nach "
        "getrennter Validierung und einer Bewertung mit Precision und Recall belastbar.",
    )
    paragraph.paragraph_format.space_before = Pt(6)


def update_workflow_step_six(document: Document) -> None:
    placeholder = find_paragraph(
        document,
        "[Ergänzung durch die Modellierung: Modelltyp, Trainingsablauf, Hyperparameter, "
        "Evaluierungsmetriken und Baselines]",
    )
    replace_paragraph(
        placeholder,
        "Für die eigentliche Modellauswahl wurden ausschließlich zwei überwachte "
        "Regressionsmodelle trainiert. Vormonat und Bis-zu-3-Monats-Mittel dienten als "
        "untrainierte Baselines und wurden deshalb nicht als Modelle ausgewiesen.",
    )
    placeholder.paragraph_format.space_before = Pt(6)
    placeholder.paragraph_format.space_after = Pt(6)
    cursor = placeholder._p

    caption = paragraph_after(document, cursor, "Trainierte Modelle", bold=True, keep_with_next=True)
    cursor = caption._p
    model_table = table_after(
        document,
        cursor,
        ["Trainiertes Modell", "Rolle", "Entscheidungsrelevantes Ergebnis"],
        [
            [
                "Lineare Regression",
                "Verständliche lineare Referenz; numerische Merkmale werden standardisiert.",
                "Mittlerer CV-RMSE 2024: 13.643 kWh.",
            ],
            [
                "Random Forest Regressor",
                "Nichtlinearer Vergleich und finales Prognosemodell.",
                "13.272 kWh; gewählt mit 300 Bäumen, Tiefe 8, min. 5 Fällen je Blatt, max_features 0,7 und random_state 42.",
            ],
        ],
        [3.7, 5.6, 6.6],
    )
    cursor = model_table._tbl

    intro = paragraph_after(
        document,
        cursor,
        "Beide Modelle erhielten dasselbe fachlich begründete Merkmalset. Es enthält nur "
        "Informationen, die vor Beginn des vorherzusagenden Monats verfügbar sind.",
    )
    intro.paragraph_format.space_before = Pt(7)
    cursor = intro._p
    caption = paragraph_after(document, cursor, "Verwendete Modellmerkmale", bold=True, keep_with_next=True)
    cursor = caption._p
    feature_table = table_after(
        document,
        cursor,
        ["Merkmalsgruppe", "Verwendete Spalten", "Information für das Modell"],
        [
            ["Verbrauchshistorie", "vormonat_vls; letzte_3_monate_vls", "Letzter gültiger Monat und Mittel aus bis zu drei verfügbaren Vormonaten."],
            ["Kalender und Jahreszeit", "monat_idx; arbeitstage; feiertage_im_monat", "Saisonale Lage sowie verfügbare Arbeits- und Feiertage."],
            ["Wetter", "heizgradtage", "Temperaturbezogener Bedarf; im späteren Betrieb als Prognosewert bereitzustellen."],
            ["Betrieb und Planung", "produktionsplan_index; wartung_aktiv", "Erwartete Produktionsaktivität und geplanter Stillstand."],
            ["Kundengruppe", "kundentyp", "Gewerbe, Industrie oder Kommunal; für das Modell One-Hot-kodiert."],
        ],
        [3.7, 5.5, 6.7],
    )
    cursor = feature_table._tbl

    note = paragraph_after(
        document,
        cursor,
        "Die Zielvariable war vollaststunden. Die Vertragsleistung wurde nicht vom Modell "
        "als Merkmal gelernt, sondern ausschließlich danach zur exakten Rückrechnung "
        "Prognose-VLS × Vertragsleistung = Prognose-kWh verwendet. Fehlende numerische "
        "Werte wurden innerhalb jedes Trainingsfolds per Median und Fehlwertindikator "
        "behandelt; der Kundentyp wurde One-Hot-kodiert.",
    )
    note.paragraph_format.space_before = Pt(7)
    cursor = note._p

    paragraph_after(
        document,
        cursor,
        "Die Auswahl erfolgte in drei vorwärts laufenden Folds aus 2024. Acht Kombinationen "
        "aus Baumtiefe, Mindestanzahl je Blatt und Merkmalsanteil wurden nur dort verglichen; "
        "2025 entschied keinen Hyperparameter. Hauptmetrik war der RMSE in kWh, ergänzt um "
        "MAE und R². Anschließend wurde die Fehlerschwelle aus November und Dezember 2024 "
        "kalibriert und erst danach auf den retrospektiven Benchmark 2025 angewendet.",
    )


def make_small_consistency_fixes(document: Document) -> None:
    dashboard = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith("Ein Prüfhinweis entsteht, wenn die Abweichung")
    )
    replace_paragraph(
        dashboard,
        "Ein Prüfhinweis entsteht, wenn die Abweichung zwischen Ist und Prognose, bezogen "
        "auf die Vertragsleistung, mindestens 144,4 Vollaststunden beträgt. Die Schwelle "
        "entspricht dem 99. Perzentil der absoluten Prognosefehler im Kalibrierungszeitraum "
        "11/2024–12/2024. Im retrospektiven Benchmark 2025 ergeben sich daraus 114 Hinweise, "
        "durchschnittlich 9,5 pro Monat. Diese Menge kann das Netzmanagement nach jedem "
        "Monatsabschluss prüfen.",
    )

    # These body paragraphs were accidentally assigned heading styles in the source.
    for exact_text in [
        "Die folgenden Abschnitte geben die persönlichen Lernerfahrungen der drei "
        "Gruppenmitglieder entlang ihrer Schwerpunkte wieder.",
    ]:
        find_paragraph(document, exact_text).style = "normal"

    finalisation_body = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith("In der Finalisierungsphase ging es darum")
    )
    finalisation_body.style = "normal"

    # These two placeholders have now been completed and should no longer look unfinished.
    for heading_prefix in ("3.3.3", "(6) Modelltraining"):
        heading = next(
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text.strip().startswith(heading_prefix)
        )
        clear_highlight(heading)

    finalisation_text = next(
        paragraph
        for paragraph in document.paragraphs
        if "vorstellt.Die individuelle Vorbereitung" in paragraph.text
    )
    replace_text_in_runs(finalisation_text, "vorstellt.Die", "vorstellt. Die")

    data_basis = next(
        paragraph
        for paragraph in document.paragraphs
        if "16.800 Zeilen und 19 Spalten" in paragraph.text
    )
    replace_text_in_runs(
        data_basis,
        "16.800 Zeilen und 19 Spalten",
        "16.800 Zeilen und 20 Spalten",
    )

    data_model_intro = next(
        paragraph
        for paragraph in document.paragraphs
        if "Datenmodell besteht aus 17 Spalten" in paragraph.text
    )
    replace_paragraph(
        data_model_intro,
        "Die aktuelle Modellierungsbasis umfasst 20 Spalten. Die nachfolgende Tabelle "
        "bildet die fachlichen Ausgangs- und Analysefelder aus der Datenaufbereitung ab; "
        "das tats\u00e4chlich trainierte Merkmalset ist in Schritt (6) separat und verbindlich "
        "ausgewiesen. Als Zielmerkmal wurde vollaststunden erg\u00e4nzt:",
    )


def main() -> None:
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)

    document = Document(SOURCE)
    compact_cover_spacing(document)
    keep_existing_table_rows_together(document)
    update_sprint_three(document)
    update_personal_outlook(document)
    update_workflow_step_six(document)
    make_small_consistency_fixes(document)

    # Ask Word to refresh TOC/page-number fields when the document is opened.
    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
