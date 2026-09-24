"""Ergaenzt den Machine-Learning-Layer in der aktuellen IHK-Berichtsvorlage.

Die Quelldatei bleibt unveraendert. Das Skript erzeugt eine neue DOCX-Datei,
fuellt den ML Canvas, dokumentiert die layerorientierte Verantwortung und
ersetzt die Platzhalter in den Modellierungs- und Ergebnisabschnitten.
"""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "Berichtsvorlage_IHK(3).docx"
OUTPUT = ROOT / "docs" / "Berichtsvorlage_IHK(3)_mit_ML_Teil.docx"

NAVY = "17365D"
PALE_BLUE = "F3F6FA"
GRID = "D9D9D9"
WHITE = "FFFFFF"
BLACK = "000000"
MONO = "Roboto Mono"


def find_paragraph(doc: Document, starts_with: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(starts_with):
            return paragraph
    raise ValueError(f"Absatz nicht gefunden: {starts_with!r}")


def find_paragraph_with_style(doc: Document, starts_with: str, style_name: str):
    for paragraph in doc.paragraphs:
        if (
            paragraph.style.name == style_name
            and paragraph.text.strip().startswith(starts_with)
        ):
            return paragraph
    raise ValueError(
        f"Absatz nicht gefunden: {starts_with!r} mit Stil {style_name!r}"
    )


def clear_paragraph(paragraph):
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)


def set_paragraph(paragraph, text: str, *, bold_lead: str | None = None):
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
        paragraph.add_run(text[len(bold_lead) :])
    else:
        paragraph.add_run(text)
    return paragraph


def add_paragraph_after(doc: Document, anchor, text: str = "", style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if text:
        paragraph.add_run(text)
    anchor.addnext(paragraph._p)
    return paragraph


def add_labeled_paragraph_after(doc: Document, anchor, label: str, body: str):
    paragraph = add_paragraph_after(doc, anchor)
    run = paragraph.add_run(label)
    run.bold = True
    paragraph.add_run(body)
    return paragraph


def delete_elements_between(start_element, end_element):
    parent = start_element.getparent()
    if end_element.getparent() is not parent:
        raise ValueError("Start- und Endelement liegen nicht im selben Dokumentteil")
    start_index = parent.index(start_element)
    end_index = parent.index(end_element)
    if end_index <= start_index:
        raise ValueError("Endelement liegt nicht hinter dem Startelement")
    for element in list(parent)[start_index + 1 : end_index]:
        parent.remove(element)


def delete_paragraph(paragraph):
    parent = paragraph._p.getparent()
    parent.remove(paragraph._p)


def delete_table(table):
    parent = table._tbl.getparent()
    parent.remove(table._tbl)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")


def set_cell_width(cell, width_cm: float):
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(round(width_cm / 2.54 * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), GRID)


def set_table_layout_fixed(table):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    """Haelt eine Tabellenzeile auf derselben Seite."""
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        no_split = OxmlElement("w:cantSplit")
        no_split.set(qn("w:val"), "true")
        tr_pr.append(no_split)


def format_cell(
    cell,
    *,
    header: bool = False,
    shade: bool = False,
    align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
    mono: bool = False,
):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, NAVY if header else (PALE_BLUE if shade else WHITE))
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.15
        for run in paragraph.runs:
            run.font.name = MONO if mono else "Arial"
            run._element.get_or_add_rPr().rFonts.set(
                qn("w:ascii"), MONO if mono else "Arial"
            )
            run._element.get_or_add_rPr().rFonts.set(
                qn("w:hAnsi"), MONO if mono else "Arial"
            )
            run.font.size = Pt(12)
            run.font.bold = bool(header)
            run.font.color.rgb = RGBColor.from_string(WHITE if header else BLACK)


def style_table(table, widths: list[float], mono_first_col: bool = False):
    set_table_layout_fixed(table)
    set_table_borders(table)
    mark_header_row(table.rows[0])
    for row_index, row in enumerate(table.rows):
        prevent_row_split(row)
        for col_index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[col_index])
            format_cell(
                cell,
                header=row_index == 0,
                shade=row_index > 0 and row_index % 2 == 0,
                mono=mono_first_col and row_index > 0 and col_index == 0,
            )


def compact_cover(doc: Document):
    """Entfernt ueberzaehlige Leerzeilen, damit das Deckblatt eine Seite belegt."""
    cover_table = doc.tables[0]
    empty_predecessors = []
    element = cover_table._tbl.getprevious()
    while element is not None and element.tag == qn("w:p"):
        text = "".join(node.text or "" for node in element.iter(qn("w:t"))).strip()
        if text:
            break
        empty_predecessors.append(element)
        element = element.getprevious()
    # Zwei Leerzeilen zwischen Untertitel und Metadaten erhalten.
    for paragraph_element in empty_predecessors[2:]:
        paragraph_element.getparent().remove(paragraph_element)


def replace_across_runs(paragraph, old: str, new: str):
    """Ersetzt Text auch dann, wenn Word ihn auf mehrere Runs verteilt hat."""
    while old in "".join(run.text for run in paragraph.runs):
        full_text = "".join(run.text for run in paragraph.runs)
        start = full_text.index(old)
        end = start + len(old)
        cursor = 0
        start_run = start_offset = end_run = end_offset = None
        for index, run in enumerate(paragraph.runs):
            next_cursor = cursor + len(run.text)
            if start_run is None and start < next_cursor:
                start_run, start_offset = index, start - cursor
            if end <= next_cursor:
                end_run, end_offset = index, end - cursor
                break
            cursor = next_cursor
        if start_run is None or end_run is None:
            raise RuntimeError(f"Text konnte nicht ersetzt werden: {old!r}")
        first = paragraph.runs[start_run]
        last = paragraph.runs[end_run]
        prefix = first.text[:start_offset]
        suffix = last.text[end_offset:]
        if start_run == end_run:
            first.text = prefix + new + suffix
        else:
            first.text = prefix + new
            for index in range(start_run + 1, end_run):
                paragraph.runs[index].text = ""
            last.text = suffix


def apply_consistency_edits(doc: Document):
    replacements = {
        "Training (2024) und Test (2025)": (
            "Modellentwicklung (2024) und retrospektiven Benchmark (2025)"
        ),
        "Prognose- und Anomalie-Modells": "Prognose- und Prüfhinweissystems",
        "Anomalie-Alerts": "Prüfhinweise",
        "residuenbasierten Anomalie-Erkennung": "residuenbasierten Prüfhinweisregel",
        "Prognosegüte, Anomalie-Erkennung und Datenqualität": "Prognosegüte, Qualität der Prüfhinweise und Datenqualität",
        "frühzeitig Anomalie-Hinweise pro Zähler": "monatliche Prüfhinweise pro Zähler",
        "nachvollziehbaren Anomalie-Definition": "nachvollziehbaren Prüfhinweisregel",
    }
    for paragraph in doc.paragraphs:
        for old, new in replacements.items():
            if old in paragraph.text:
                replace_across_runs(paragraph, old, new)

        if paragraph.text.strip().startswith("Dabei sollten Prognosegüte"):
            replace_across_runs(
                paragraph, "Anomalie-Erkennung", "Qualität der Prüfhinweise"
            )

    rubric = find_paragraph(doc, "Punktevergabe: 20 von 70")
    delete_paragraph(rubric)


def update_static_toc(doc: Document):
    """Aktualisiert das manuelle Inhaltsverzeichnis fuer den geprueften Renderstand."""
    entries = {
        "1   Ausgangssituation": "1   Ausgangssituation und Ist-Zustand ............................... S. 4",
        "2   Machine Learning Canvas": "2   Machine Learning Canvas ............................................. S. 5",
        "3   Projektplanung": "3   Projektplanung und -durchführung ................................. S. 7",
        "4   Workflow": "4   Workflow, Datenmodell und Ergebnisse ........................ S. 12",
        "Literaturverzeichnis": "Literaturverzeichnis ........................................................ S. 20",
        "Abkürzungsverzeichnis": "Abkürzungsverzeichnis .................................................. S. 3",
    }
    for starts_with, replacement in entries.items():
        paragraph = find_paragraph(doc, starts_with)
        set_paragraph(paragraph, replacement)


def add_table_after(
    doc: Document,
    anchor,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    *,
    mono_first_col: bool = False,
):
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    style_table(table, widths, mono_first_col=mono_first_col)
    anchor.addnext(table._tbl)
    return table


def add_bullet_after(doc: Document, anchor, text: str):
    paragraph = add_paragraph_after(doc, anchor, style="List Bullet")
    paragraph.add_run(text)
    paragraph.paragraph_format.left_indent = Cm(0.65)
    paragraph.paragraph_format.first_line_indent = Cm(-0.35)
    return paragraph


def add_number_after(doc: Document, anchor, text: str):
    paragraph = add_paragraph_after(doc, anchor, style="List Number")
    paragraph.add_run(text)
    paragraph.paragraph_format.left_indent = Cm(0.65)
    paragraph.paragraph_format.first_line_indent = Cm(-0.35)
    return paragraph


def paragraph_after_table(doc: Document, table, text: str = "", style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    if text:
        paragraph.add_run(text)
    table._tbl.addnext(paragraph._p)
    return paragraph


def build_canvas(doc: Document):
    overview = find_paragraph(doc, "Der Machine Learning Canvas ist")
    set_paragraph(
        overview,
        "Der Machine Learning Canvas verbindet das geschäftliche Ziel mit der Datengrundlage, "
        "der Modellierung und dem späteren Arbeitsprozess. Im Projekt werden zwei zeitlich "
        "getrennte Aufgaben unterschieden: Vor Beginn eines Monats wird der erwartete "
        "Stromverbrauch je Zähler prognostiziert. Nach Eingang des tatsächlichen Monatswerts "
        "wird die Abweichung zwischen Istwert und Prognose geprüft. Eine ungewöhnlich große "
        "Abweichung erzeugt einen Prüfhinweis, aber noch keine bestätigte Anomalie oder "
        "automatische Maßnahme.",
    )

    heading = find_paragraph(doc, "2.2")
    chapter_three = find_paragraph_with_style(doc, "3   Projektplanung", "Heading 1")
    delete_elements_between(heading._p, chapter_three._p)

    canvas_rows = [
        [
            "1. Value Proposition / Mehrwert",
            "Die Energiebeschaffung erhält vor Monatsbeginn eine belastbarere Verbrauchsprognose. "
            "Das Netzmanagement kann ungewöhnliche Zählerwerte monatlich statt erst beim "
            "Quartalsabschluss prüfen.",
        ],
        [
            "2. Data Sources / Datenquellen",
            "Monatliche Zähler- und Verbrauchsdaten, Vertragsleistung, Kundentyp, Kalenderdaten, "
            "Heizgradtage beziehungsweise Wetterprognose, Produktionsplan und geplante Wartung. "
            "Die Modellierungsbasis umfasst 700 Zähler von Januar 2024 bis Dezember 2025.",
        ],
        [
            "3. Prediction / Vorhersage",
            "Überwachte Regression der Vollaststunden des Folgemonats je Zähler. Die Prognose wird "
            "über die jeweilige Vertragsleistung wieder in kWh umgerechnet. Nach Monatsabschluss "
            "bildet das Residuum die Grundlage für einen Prüfhinweis.",
        ],
        [
            "4. Features / Merkmale",
            "Vertragsleistung, Monatsnummer, Arbeitstage, Feiertage, Heizgradtage, Produktionsplan, "
            "geplante Wartung, Kundentyp, Vormonatswert und Mittel aus bis zu drei vorherigen "
            "Monaten. Verwendet werden nur Informationen, die zum Prognosezeitpunkt vorliegen.",
        ],
        [
            "5. Learning Approach / Lernansatz",
            "Eine lineare Regression dient als verständliche Referenz. Ein Random Forest bildet "
            "zusätzlich nichtlineare Zusammenhänge ab. Beide Modelle werden mit einfachen "
            "historischen Baselines verglichen.",
        ],
        [
            "6. Evaluation / Überprüfung",
            "Hauptmetrik ist der RMSE in kWh, da große Mengenfehler stärker gewichtet werden. Der "
            "MAE ergänzt die Bewertung als durchschnittliche absolute Abweichung. Die Modellwahl "
            "erfolgt über drei zeitlich vorwärts laufende Prüfungen aus 2024. Als Baselines dienen "
            "der Vormonat und das Mittel aus bis zu drei Vormonaten.",
        ],
        [
            "7. Decision / Entscheidung",
            "Die kWh-Prognose unterstützt die Mengenplanung. Überschreitet das spätere absolute "
            "VLS-Residuum die kalibrierte Schwelle, wird der Zähler-Monat priorisiert geprüft. "
            "Messwert, Datenqualität, Produktionsplan, Wartung und betrieblicher Kontext werden "
            "durch eine Fachkraft bewertet.",
        ],
        [
            "8. Impact / Auswirkung",
            "Gemessen werden Prognosefehler, monatliches Hinweisvolumen, Bearbeitungszeit und "
            "spätere Bestätigungsquote. Im retrospektiven Benchmark liegt der RMSE 15,9 Prozent "
            "unter der besten einfachen Baseline. Beim gewählten 99. Perzentil entstehen 9,5 "
            "Prüfhinweise pro Monat. Ein finanzieller Nutzen ist ohne Preis-, Maßnahmen- und "
            "Störungsdaten noch nicht belastbar berechenbar.",
        ],
        [
            "9. Prediction Timing / Vorhersagezeitpunkt",
            "Die Folgemonatsprognose wird vor Beginn des Zielmonats erstellt. Ein Prüfhinweis kann "
            "erst nach Eingang des realisierten Monatsverbrauchs entstehen. Die Anwendung ist "
            "daher bewusst kein Echtzeitsystem.",
        ],
        [
            "10. Monitoring & Maintenance / Wartung",
            "Monatlich werden Datenqualität, Prognosefehler, Merkmalsverteilungen, Hinweisvolumen "
            "und fachliche Rückmeldungen überwacht. Modell und Schwelle werden erst bei "
            "dokumentierter Verschlechterung oder veränderten Rahmenbedingungen neu trainiert "
            "beziehungsweise kalibriert.",
        ],
    ]
    table = add_table_after(
        doc,
        heading._p,
        ["Feld", "Projektinhalt"],
        canvas_rows,
        [4.25, 11.75],
    )
    coherence = paragraph_after_table(doc, table, "2.3   Kohärenz des Canvas", "Heading 2")
    anchor = add_paragraph_after(
        doc,
        coherence._p,
        "Die Felder Vorhersage, Lernansatz und Überprüfung bilden eine konsistente Einheit: Eine "
        "kontinuierliche Zielgröße wird durch überwachte Regression geschätzt und mit "
        "Regressionsmetriken in kWh bewertet. Die Vollaststunden dienen als interne Normierung; "
        "das Ergebnis für den Fachbereich bleibt eine Verbrauchsprognose in kWh.",
    )
    anchor = add_paragraph_after(
        doc,
        anchor._p,
        "Die Felder Entscheidung und Vorhersagezeitpunkt trennen Prognose und Anomalieprüfung "
        "bewusst: Die Prognose entsteht vor dem Zielmonat, der Prüfhinweis erst nach Eingang des "
        "Istwerts. Die Wirkung wird zunächst über belegbare Prognose- und Prozesskennzahlen "
        "gemessen. Nicht nachgewiesene Kosteneinsparungen werden nicht angenommen.",
    )
    chain = add_paragraph_after(
        doc,
        anchor._p,
        "Daten  →  Vollaststunden  →  Regression  →  Rückrechnung in kWh  →  Residuum  →  "
        "Prüfhinweis  →  fachliche Entscheidung",
    )
    chain.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in chain.runs:
        run.bold = True


def build_project_organization(doc: Document):
    intro = find_paragraph(doc, "Die Gruppe bestand aus drei Mitgliedern")
    set_paragraph(
        intro,
        "Die Projektarbeit wurde layerorientiert organisiert. Dadurch blieb jedes "
        "Gruppenmitglied über alle Projektphasen für einen fachlichen Bereich "
        "hauptverantwortlich.",
    )
    roles = [
        [
            "Person A – Data Management",
            "Datenquelle, Import, Datenqualitätsprüfung, Bereinigung und Bereitstellung der "
            "analytischen Tabelle",
        ],
        [
            "Person B – Analytics und Visualisierung",
            "Explorative Datenanalyse, fachliche Visualisierungen und Gestaltung des Dashboards",
        ],
        [
            "Person C – Machine Learning",
            "Kapitel 2 sowie Zielvariable, Feature Engineering, zeitliche Validierung, "
            "Hyperparametervergleich, Modellwahl, Evaluation, Permutation Importance, "
            "Schwellenkalibrierung und Datenschnittstelle zum Dashboard",
        ],
    ]
    table = add_table_after(
        doc,
        intro._p,
        ["Rolle", "Hauptverantwortung"],
        roles,
        [4.6, 11.4],
    )
    followup = find_paragraph(doc, "Die Aufgaben wurden zu Projektbeginn")
    set_paragraph(
        followup,
        "Die Rollen wurden zu Projektbeginn vereinbart und über alle drei Sprints beibehalten. "
        "Im Sprint-Log ist jede Aufgabe genau einer hauptverantwortlichen Person zugeordnet; "
        "unterstützende Beiträge werden getrennt dokumentiert. Die Ergebnisse der drei Layer "
        "wurden anschließend zu Bericht, Präsentation und Live-Demo zusammengeführt.",
    )

    target = find_paragraph(doc, "Ziel des Projekts war es")
    set_paragraph(
        target,
        "Ziel des Projekts war es, für jeden Zähler die Vollaststunden des Folgemonats zu "
        "prognostizieren und das Ergebnis über die Vertragsleistung wieder in kWh umzurechnen. "
        "Die Prognose soll die monatliche Mengenplanung unterstützen. Nach Eingang des "
        "tatsächlichen Verbrauchs werden ungewöhnlich große Abweichungen als Prüfhinweise an das "
        "Netzmanagement übergeben. Das Modell bestätigt dabei weder einen Defekt noch eine "
        "Wartungsmaßnahme.",
    )


def build_sprint_three(doc: Document):
    heading = find_paragraph(doc, "3.3.3")
    heading.text = "3.3.3   Sprint 3 — Modellierung und Evaluierung"
    heading.style = "Heading 3"
    next_heading = find_paragraph(doc, "3.3.4")
    delete_elements_between(heading._p, next_heading._p)
    anchor = add_labeled_paragraph_after(
        doc,
        heading._p,
        "Sprint-Ziel. ",
        "Ziel des dritten Sprints war eine reproduzierbare Modellierung, die zukünftige Monate "
        "ausschließlich aus zuvor verfügbaren Informationen prognostiziert. Zusätzlich sollte "
        "eine transparente Regel große Prognoseabweichungen in fachlich prüfbare Hinweise "
        "übersetzen.",
    )
    anchor = add_labeled_paragraph_after(doc, anchor._p, "Aufgaben. ", "Der ML-Layer umfasste:")
    for item in [
        "Vollaststunden als interne Zielvariable bilden und Prognosen wieder in kWh zurückrechnen",
        "Historienmerkmale je Zähler ohne Nutzung des aktuellen oder eines zukünftigen Monats rekonstruieren",
        "Vormonat und Mittel aus bis zu drei Vormonaten als Baselines definieren",
        "lineare Regression und Random Forest in drei zeitlichen Folds vergleichen",
        "acht kontrollierte Random-Forest-Konfigurationen ausschließlich anhand der 2024-Folds bewerten",
        "die Prüfhinweisschwelle aus November und Dezember 2024 kalibrieren und die Ergebnisse an das Dashboard übergeben",
    ]:
        anchor = add_bullet_after(doc, anchor._p, item)
    anchor = add_labeled_paragraph_after(
        doc,
        anchor._p,
        "Ergebnisse. ",
        "Der Random Forest erreichte mit 13.272 kWh den niedrigsten mittleren RMSE der drei "
        "zeitlichen Prüfungen. Die gewählte Konfiguration verwendet 300 Bäume, eine maximale "
        "Tiefe von 8, mindestens 5 Fälle pro Blatt und 70 Prozent der Merkmale je Aufteilung. Im "
        "retrospektiven Benchmark 2025 erreichte das Modell 9.188 kWh RMSE und 3.725 kWh MAE. "
        "Gegenüber dem Mittel aus bis zu drei Vormonaten verringerte sich der RMSE um 15,9 Prozent.",
    )
    add_labeled_paragraph_after(
        doc,
        anchor._p,
        "Herausforderungen. ",
        "Die Historie durfte keine Werte aus dem vorherzusagenden oder einem späteren Monat "
        "enthalten. Sie wurde deshalb je Zähler mit einer zeitlichen Verschiebung neu aufgebaut. "
        "Dadurch konnten 1.453 zuvor fehlende Historienwerte genutzt werden, ohne Zukunftswerte "
        "einzubauen. Die verbleibenden 700 fehlenden Werte entsprechen dem ersten Monat jedes "
        "Zählers und wurden erst innerhalb des jeweiligen Trainingsfensters imputiert.",
    )


def build_workflow(doc: Document):
    overview = find_paragraph(doc, "[Beispielsatz: Der Workflow umfasste")
    set_paragraph(
        overview,
        "Der Workflow folgt sieben nachvollziehbaren Schritten: Quelle, Import, Bereinigung, "
        "Transformation, Modellierung, Visualisierung und Ergebnisübergabe. Die folgenden "
        "Abschnitte dokumentieren pro Schritt nicht nur die technische Ausführung, sondern auch "
        "die fachliche Entscheidung und ihre Begründung.",
    )

    source_heading = find_paragraph(doc, "Quelle")
    source_heading.text = "4.1.1   Quelle"
    source_heading.style = "Heading 3"
    source_text = find_paragraph(doc, "Als Datenquelle wurde")
    set_paragraph(
        source_text,
        "Als Rohquelle wurde die Datei data/raw/verbrauch.csv verwendet. Sie enthielt 16.830 "
        "monatliche Beobachtungen von 700 Zählern für Januar 2024 bis Dezember 2025. Nach der "
        "Bereinigung von 30 doppelten Zähler-Monat-Kombinationen umfasst die Modellierungsbasis "
        "16.800 Zeilen. Für den ML-Layer wurde daraus die reproduzierbare Datei "
        "data/processed/modellierung_basis_bis_3_monate.csv abgeleitet.",
    )

    import_heading = find_paragraph(doc, "Import")
    import_heading.text = "4.1.2   Import"
    import_heading.style = "Heading 3"
    cleanup_heading = find_paragraph(doc, "Bereinigung")
    cleanup_heading.text = "4.1.3   Bereinigung"
    cleanup_heading.style = "Heading 3"

    model_data_heading = find_paragraph(doc, "4.2")
    anchor = model_data_heading._p.getprevious()
    while anchor is not None and anchor.tag != qn("w:p"):
        anchor = anchor.getprevious()
    if anchor is None:
        raise RuntimeError("Einfügeposition vor 4.2 wurde nicht gefunden")

    transform_heading = add_paragraph_after(
        doc, anchor, "4.1.4   Transformation und Feature Engineering", "Heading 3"
    )
    anchor_p = add_paragraph_after(
        doc,
        transform_heading._p,
        "Als interne Zielvariable wurden Vollaststunden (VLS) gebildet. Sie ergeben sich aus dem "
        "Verbrauch in kWh geteilt durch die Vertragsleistung in kW. Damit wird der Verbrauch "
        "relativ zur Anschlussgröße betrachtet. VLS sind hier eine rechnerische Normierung und "
        "keine gemessene Laufzeit. Für Fachbereich und Modellbewertung wird jede Prognose wieder "
        "in kWh zurückgerechnet.",
    )
    anchor_p = add_paragraph_after(
        doc,
        anchor_p._p,
        "Die Historienmerkmale wurden chronologisch und getrennt je Zähler berechnet. Der aktuelle "
        "Zielmonat wurde durch eine Verschiebung um einen Monat ausgeschlossen. Im ersten Monat "
        "blieb die Historie leer, im zweiten Monat stand ein Vormonat zur Verfügung und ab dem "
        "vierten Monat flossen bis zu drei gültige Vormonate ein. Als unmöglich markierte "
        "Messwerte wurden nicht in spätere Historienfenster übernommen.",
    )
    anchor_p = add_paragraph_after(
        doc,
        anchor_p._p,
        "Fehlende Merkmalswerte wurden mit einem Median-Imputer behandelt, der in jedem Fold "
        "ausschließlich auf den jeweiligen Trainingsdaten gelernt wurde. Der Vorjahresverbrauch "
        "blieb im Hauptmodell ungenutzt, weil für 2024 keine Historie aus 2023 vorlag. Auch die "
        "regelbasierte EDA-Spalte anomalie wurde weder als Merkmal noch als Ziel verwendet, da sie "
        "keine fachlich bestätigte Ground Truth darstellt.",
    )

    validation_heading = add_paragraph_after(
        doc, anchor_p._p, "4.1.5   Modellierung und zeitliche Validierung", "Heading 3"
    )
    anchor_p = add_paragraph_after(
        doc,
        validation_heading._p,
        "Die Modellwahl erfolgte ausschließlich mit Daten aus Januar bis Oktober 2024. Ein "
        "zufälliger Split wurde vermieden, weil sonst spätere Monate in das Training früherer "
        "Prognosen gelangen könnten. Stattdessen wurden drei vorwärts laufende Prüfungen genutzt. "
        "November und Dezember 2024 dienten anschließend separat zur Kalibrierung der "
        "Prüfhinweisschwelle: Für November wurde nur bis Oktober, für Dezember nur bis November "
        "trainiert.",
    )
    folds = [
        ["Fold 1", "01/2024–04/2024", "05/2024–06/2024", "Modellvergleich"],
        ["Fold 2", "01/2024–06/2024", "07/2024–08/2024", "Modellvergleich"],
        ["Fold 3", "01/2024–08/2024", "09/2024–10/2024", "Modellvergleich"],
        ["Kalibrierung", "jeweils alle früheren Monate", "11/2024–12/2024", "Schwelle festlegen"],
        ["Benchmark", "gesamtes Jahr 2024", "01/2025–12/2025", "retrospektive Prüfung"],
    ]
    fold_table = add_table_after(
        doc,
        anchor_p._p,
        ["Block", "Lernen", "Bewerten", "Zweck"],
        folds,
        [2.4, 4.8, 4.3, 4.5],
    )
    anchor_p = paragraph_after_table(
        doc,
        fold_table,
        "Für den Random Forest wurden acht Kombinationen aus maximaler Baumtiefe (8 oder "
        "unbegrenzt), minimaler Blattgröße (5 oder 20) und Merkmalsanteil je Aufteilung (0,7 oder "
        "1,0) in allen drei Folds geprüft. Dies entspricht 24 kontrollierten Modellanpassungen. "
        "Die Anzahl der Bäume blieb bei 300 und der Zufallsstartwert bei 42. Die beste Kombination "
        "war maximale Tiefe 8, mindestens 5 Fälle pro Blatt und ein Merkmalsanteil von 0,7.",
    )
    anchor_p = add_paragraph_after(
        doc,
        anchor_p._p,
        "Nach Modellwahl und Schwellenkalibrierung wurde der ausgewählte Random Forest auf allen "
        "gültigen Entwicklungsdaten aus 2024 neu trainiert. Modelltyp, Hyperparameter und "
        "Entscheidungsgrenze blieben für 2025 unverändert. Die zulässigen Historienmerkmale wurden "
        "nach jedem abgeschlossenen Monat aktualisiert. Dies bildet eine One-Step-Ahead-Anwendung "
        "nach und keine einmalige Prognose aller zwölf Monate zu Jahresbeginn.",
    )

    visual_heading = add_paragraph_after(
        doc, anchor_p._p, "4.1.6   Visualisierung", "Heading 3"
    )
    anchor_p = add_paragraph_after(
        doc,
        visual_heading._p,
        "Für die Modellentscheidung wurden der RMSE-Vergleich mit Baselines, die Streuung der "
        "Zeit-Folds, die gruppierte Permutation Importance und die sortierten "
        "Kalibrierungsfehler visualisiert. Für die fachliche Prüfung verbindet das Dashboard "
        "Istwert, Prognose, Residuum, Schwellenfaktor und betriebliche Kontextmerkmale je Fall.",
    )

    export_heading = add_paragraph_after(
        doc, anchor_p._p, "4.1.7   Export und Ergebnisübergabe", "Heading 3"
    )
    add_paragraph_after(
        doc,
        export_heading._p,
        "Die Ergebnisdatei für das Dashboard enthält pro Zähler-Monat die kWh-Prognose, den "
        "Istwert, das Residuum in kWh und VLS, den Schwellenfaktor, den Prüfhinweis sowie "
        "Datenqualitäts-, Produktions- und Wartungskontext. Damit bleibt die Modellrechnung von "
        "der fachlichen Entscheidung getrennt und zugleich nachvollziehbar.",
    )


def build_data_model(doc: Document):
    heading = find_paragraph(doc, "4.2")
    next_heading = find_paragraph(doc, "4.3")
    delete_elements_between(heading._p, next_heading._p)
    anchor = add_paragraph_after(
        doc,
        heading._p,
        "Die finale Modellierungsdatei umfasst 20 Spalten. Vollaststunden sind die interne "
        "kontinuierliche Zielvariable; verbrauch_kwh bleibt der realisierte Monatswert und die "
        "fachliche Ausgabeeinheit. Drei als ziel_rekonstruiert gekennzeichnete Werte wurden weder "
        "für das Training noch für Modellmetriken oder die Schwellenkalibrierung verwendet.",
    )
    formula = add_paragraph_after(
        doc,
        anchor._p,
        "Vollaststunden = Verbrauch in kWh ÷ Vertragsleistung in kW",
    )
    formula.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in formula.runs:
        run.bold = True
    rows = [
        [
            "Identifikation und Zeit",
            "zaehler_id, monat, jahr, split",
            "Text, Datum, Ganzzahl",
            "Eindeutige Zuordnung und zeitliche Aufteilung",
        ],
        [
            "Ziel und Ausgabe",
            "vollaststunden, verbrauch_kwh",
            "Zahl",
            "Interne Zielvariable und realisierte kWh-Ausgabe",
        ],
        [
            "Vertrag und Kunde",
            "vertragsleistung_kw, kundentyp",
            "Zahl, Text",
            "Normierung, Rückrechnung und Kundengruppe",
        ],
        [
            "Kalender und Wetter",
            "monat_idx, arbeitstage, feiertage_im_monat, heizgradtage",
            "Ganzzahl, Zahl",
            "Zum Prognosezeitpunkt bekannte Saison- und Wettermerkmale",
        ],
        [
            "Planung",
            "produktionsplan_index, wartung_aktiv",
            "Zahl, 0/1",
            "Geplante Produktion und Wartung; fehlende Werte foldweise imputiert",
        ],
        [
            "Verbrauchshistorie",
            "vormonat_vls, letzte_3_monate_vls",
            "Zahl",
            "Je Zähler mit shift(1) aus höchstens drei abgeschlossenen Monaten",
        ],
        [
            "Referenz, nicht im Hauptmodell",
            "vorjahr_vls, anomalie",
            "Zahl, 0/1",
            "Vorjahr wegen fehlender 2023-Historie; EDA-Regel ohne bestätigte Labels",
        ],
        [
            "Datenqualität",
            "unmoeglich, ziel_rekonstruiert",
            "0/1",
            "Steuert Ausschluss von Mess- beziehungsweise Zielwerten",
        ],
    ]
    add_table_after(
        doc,
        formula._p,
        ["Gruppe", "Spalten", "Typen", "Rolle im Modell"],
        rows,
        [3.4, 5.0, 2.4, 5.2],
    )


def build_results(doc: Document):
    heading = find_paragraph(doc, "4.3")
    heading.text = "4.3   Modellierung und Ergebnisse"
    quality_heading = find_paragraph(doc, "4.4")
    delete_elements_between(heading._p, quality_heading._p)

    sub = add_paragraph_after(doc, heading._p, "4.3.1   Modellwahl 2024", "Heading 3")
    anchor = add_paragraph_after(
        doc,
        sub._p,
        "Der Random Forest erzielte den niedrigsten mittleren RMSE in den drei zeitlichen "
        "Prüfungen und wurde deshalb anhand der vorab festgelegten Hauptmetrik ausgewählt. Sein "
        "Vorteil gegenüber der linearen Regression beträgt 2,7 Prozent und wird nicht als großer "
        "Leistungssprung interpretiert. Die lineare Regression bleibt eine nachvollziehbare "
        "Referenz und ein vertretbarer Fallback.",
    )
    model_rows = [
        ["Random Forest auf VLS", "13.272 kWh", "3.843 kWh", "9.188 kWh"],
        ["Lineare Regression auf VLS", "13.643 kWh", "3.472 kWh", "9.414 kWh"],
        ["Mittel aus bis zu drei Vormonaten", "15.483 kWh", "3.070 kWh", "10.922 kWh"],
        ["Vormonat", "18.460 kWh", "6.366 kWh", "12.386 kWh"],
    ]
    table = add_table_after(
        doc,
        anchor._p,
        ["Kandidat", "CV-RMSE 2024", "Fold-Streuung", "RMSE 2025"],
        model_rows,
        [6.4, 3.2, 3.2, 3.2],
    )
    anchor = paragraph_after_table(
        doc,
        table,
        "Die Fold-Streuung ist die Standardabweichung der drei zeitlichen RMSE-Werte. Sie ist "
        "weder ein Konfidenzintervall noch die Unsicherheit einer einzelnen Prognose. Ihre Höhe "
        "zeigt, dass die Modellgüte zwischen Zeitfenstern schwankt und nicht aus einem einzelnen "
        "Monat abgeleitet werden sollte.",
    )

    sub = add_paragraph_after(doc, anchor._p, "4.3.2   Retrospektiver Benchmark 2025", "Heading 3")
    anchor = add_paragraph_after(
        doc,
        sub._p,
        "Das ausgewählte Modell erreichte auf 8.398 bewertbaren Zähler-Monaten einen RMSE von "
        "9.188 kWh und einen MAE von 3.725 kWh. Gegenüber der besten einfachen Baseline, dem "
        "Mittel aus bis zu drei Vormonaten, reduzierte sich der RMSE um 15,9 Prozent. Die "
        "Vormonatsbaseline konnte wegen zehn fehlender Vergleichswerte nur auf 8.388 Fällen "
        "bewertet werden; die ausgewiesene Verbesserung bezieht sich deshalb auf die Baseline mit "
        "vergleichbarer Fallbasis.",
    )
    anchor = add_paragraph_after(
        doc,
        anchor._p,
        "In einem ergänzenden Methodenvergleich erreichte ein direkt auf kWh trainierter Random "
        "Forest 9.800 kWh RMSE, während der VLS-Ansatz 9.188 kWh erreichte. Die Verringerung um "
        "6,2 Prozent begründet die interne Zielvariable für diesen Datensatz und diesen "
        "Versuchsaufbau. Sie beweist keine allgemeine Überlegenheit von Vollaststunden.",
    )

    sub = add_paragraph_after(doc, anchor._p, "4.3.3   Einfluss der Merkmalsgruppen", "Heading 3")
    anchor = add_paragraph_after(
        doc,
        sub._p,
        "Die gruppierte Permutation Importance misst, wie stark der RMSE steigt, wenn die Werte "
        "einer Informationsgruppe zufällig vertauscht werden. Die Verbrauchshistorie besitzt mit "
        "einem RMSE-Anstieg von rund 7.349 kWh den größten Prognosebeitrag. Produktionsplan und "
        "geplante Wartung folgen mit rund 717 beziehungsweise 578 kWh. Kalender, Wetter und "
        "Jahreszeit liefern kleinere zusätzliche Beiträge; der Kundentyp liegt mit rund 5 kWh "
        "nahe null.",
    )
    anchor = add_paragraph_after(
        doc,
        anchor._p,
        "Die Auswertung beschreibt den Prognosenutzen innerhalb des Modells, aber keine "
        "Kausalität. Korrelierte Merkmale können sich Bedeutung teilen. Die Vertragsleistung "
        "wurde nicht in diese Rangfolge aufgenommen, weil sie zugleich Modellmerkmal und "
        "struktureller Faktor der Rückrechnung in kWh ist; ein zufälliges Vertauschen würde beide "
        "Rollen vermischen.",
    )

    sub = add_paragraph_after(doc, anchor._p, "4.3.4   Residuenbasierte Prüfhinweise", "Heading 3")
    anchor = add_paragraph_after(
        doc,
        sub._p,
        "Nach Eingang des Istwerts wird das Residuum als Differenz zwischen tatsächlichen und "
        "prognostizierten Vollaststunden berechnet. Für die Kalibrierung wurden 1.397 absolute "
        "Prognosefehler aus November und Dezember 2024 aufsteigend sortiert. Das 99. Perzentil "
        "liegt bei 144,4 VLS-Stunden. Es bezeichnet eine Position in dieser Fehlerverteilung und "
        "weder einen Anteil der Gesamtleistung noch eine Defektwahrscheinlichkeit.",
    )
    anchor = add_paragraph_after(
        doc,
        anchor._p,
        "Ein Prüfhinweis entsteht, wenn der absolute VLS-Fehler mindestens 144,4 Stunden beträgt. "
        "Die zugehörige kWh-Grenze ist je Zähler unterschiedlich und ergibt sich aus 144,4 "
        "VLS-Stunden multipliziert mit der Vertragsleistung. Das 99. Perzentil ist eine "
        "konservative Pilotannahme mit überschaubarem Prüfvolumen, nicht ein nachgewiesener "
        "optimaler Sweet Spot.",
    )
    threshold_rows = [
        ["95,0 %", "68,0 VLS-h", "400", "33,3"],
        ["97,5 %", "82,4 VLS-h", "256", "21,3"],
        ["99,0 %", "144,4 VLS-h", "114", "9,5"],
        ["99,5 %", "196,3 VLS-h", "69", "5,8"],
    ]
    threshold_table = add_table_after(
        doc,
        anchor._p,
        ["Perzentil", "Schwelle", "Prüfhinweise 2025", "je Monat"],
        threshold_rows,
        [3.4, 4.2, 4.8, 3.6],
    )
    paragraph_after_table(
        doc,
        threshold_table,
        "Mit der gewählten 99-Prozent-Regel entstehen 114 Prüfhinweise auf 107 verschiedenen "
        "Zählern. Dies entspricht 1,36 Prozent der 8.398 bewertbaren Zähler-Monate. 68 Hinweise "
        "betreffen ungewöhnlich hohen, 46 ungewöhnlich niedrigen Verbrauch. Diese Zahlen "
        "quantifizieren das Arbeitsvolumen der Regel und nicht die Anzahl tatsächlicher Defekte.",
    )


def build_quality_and_dashboard(doc: Document):
    quality = find_paragraph(doc, "4.4")
    dashboard = find_paragraph(doc, "4.5")
    delete_elements_between(quality._p, dashboard._p)
    anchor = add_paragraph_after(
        doc,
        quality._p,
        "Die Modellierung ist methodisch plausibel, weil sie die zeitliche Reihenfolge einhält, "
        "einfache Baselines einbezieht, die Vorverarbeitung nur auf Trainingsdaten lernt und die "
        "Baumtiefe des gewählten Random Forest begrenzt. Das Modell bleibt auch im späteren Jahr "
        "vor der linearen Referenz und den historischen Baselines. Ein eindeutiges "
        "Overfitting-Signal ist damit nicht erkennbar.",
    )
    anchor = add_paragraph_after(
        doc,
        anchor._p,
        "Produktionsreife lässt sich daraus noch nicht ableiten. Der Datensatz umfasst nur zwei "
        "Jahre, und die Schwelle basiert ausschließlich auf zwei winterlichen Monaten. Das Jahr "
        "2025 wurde im Projekt bereits ausgewertet und ist deshalb ein retrospektiver Benchmark, "
        "kein unangesehener Blindtest. Ein prospektiver Schattenbetrieb muss die Stabilität unter "
        "zukünftigen Bedingungen erst bestätigen.",
    )
    add_paragraph_after(
        doc,
        anchor._p,
        "Außerdem fehlen fachlich bestätigte Anomalielabel. Precision, Recall, Fehlalarmquote und "
        "die Zahl übersehener relevanter Fälle können deshalb noch nicht seriös berechnet werden. "
        "Vor einem Produktiveinsatz sind eine längere Kalibrierungsbasis, dokumentierte "
        "Fachentscheidungen und stichprobenartige Kontrollen nicht markierter Fälle erforderlich.",
    )

    literature = find_paragraph_with_style(doc, "Literaturverzeichnis", "Heading 1")
    delete_elements_between(dashboard._p, literature._p)
    anchor = add_paragraph_after(
        doc,
        dashboard._p,
        "Das Dashboard überführt die Modellergebnisse in einen kontrollierten Arbeitsprozess. "
        "Der Perzentil-Regler ändert nicht das trainierte Modell, sondern ausschließlich die "
        "Entscheidungsgrenze. Kennzahlen, farbige Punkte, Fallliste und Detailansicht werden "
        "gemeinsam aktualisiert, damit der Zusammenhang zwischen Sensitivität und Prüfaufwand "
        "sichtbar bleibt.",
    )
    for item in [
        "Prüfhinweise nach Schwellenfaktor und kWh-Auswirkung priorisieren",
        "einen Zähler-Monat mit Istwert, Prognose und Residuum öffnen",
        "Datenqualität, Produktionsplan, Wartung und betrieblichen Kontext prüfen",
        "fachliche Entscheidung und Ursache dokumentieren",
        "bestätigte und verworfene Hinweise als Feedback für das Monitoring sichern",
    ]:
        anchor = add_number_after(doc, anchor._p, item)
    anchor = add_paragraph_after(
        doc,
        anchor._p,
        "Diese Rückmeldungen können in einer späteren Projektstufe als Zielvariable für ein "
        "Klassifikationsmodell dienen. Die Regression würde weiterhin den erwarteten Verbrauch "
        "bestimmen; das zusätzliche Modell könnte lernen, welche Prüfhinweise tatsächlich relevant "
        "waren. Erst mit solchen Labels lassen sich Precision und Recall berechnen und "
        "Schwellenwert beziehungsweise Priorisierung belastbar optimieren.",
    )
    page_break = add_paragraph_after(doc, anchor._p)
    page_break.add_run().add_break(WD_BREAK.PAGE)


def update_abbreviations(doc: Document):
    table = doc.tables[1]
    rows = [
        ["CRM", "Customer Relationship Management"],
        ["CV", "Cross Validation (Kreuzvalidierung)"],
        ["EDA", "Exploratory Data Analysis"],
        ["ETL", "Extract, Transform, Load"],
        ["KPI", "Key Performance Indicator"],
        ["MAE", "Mean Absolute Error (mittlerer absoluter Fehler)"],
        ["ML", "Machine Learning"],
        ["RMSE", "Root Mean Squared Error (Wurzel des mittleren quadratischen Fehlers)"],
        ["VLS", "Vollaststunden"],
    ]
    parent = table._tbl.getparent()
    index = parent.index(table._tbl)
    delete_table(table)
    replacement = doc.add_table(rows=1, cols=2)
    replacement.rows[0].cells[0].text = "Abkürzung"
    replacement.rows[0].cells[1].text = "Bedeutung"
    for values in rows:
        cells = replacement.add_row().cells
        cells[0].text, cells[1].text = values
    style_table(replacement, [4.0, 12.0])
    parent.insert(index, replacement._tbl)

    note = find_paragraph(doc, "Listen Sie alle Abkürzungen")
    set_paragraph(
        note,
        "Die im Bericht verwendeten Abkürzungen sind alphabetisch aufgeführt.",
    )


def set_document_defaults(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    normal = doc.styles["normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for style_name, size in [("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)]:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
        style.paragraph_format.line_spacing = 1.5


def mark_fields_for_update(doc: Document):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def validate(doc: Document):
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    required = [
        "2.3   Kohärenz des Canvas",
        "layerorientiert organisiert",
        "13.272 kWh",
        "9.188 kWh",
        "144,4 VLS-Stunden",
        "114 Prüfhinweise",
        "retrospektiver Benchmark",
        "Precision, Recall",
    ]
    for phrase in required:
        if phrase not in text:
            raise AssertionError(f"Pflichtinhalt fehlt: {phrase}")
    forbidden = [
        "Smart Meter",
        "Smart-Meter",
        "99,7",
        "99.7",
        "Prognosegüte, Anomalie-Erkennung und Datenqualität",
        "Dadurch sinkt die Streuung zwischen den Zählern von 92 % auf 56 %",
        "Prediction / Vorhersage: [Ihre Antwort.]",
        "[Beispielsatz: Auf dem Testdatensatz",
    ]
    for phrase in forbidden:
        if phrase in text:
            raise AssertionError(f"Veralteter oder unvollständiger Inhalt gefunden: {phrase}")
    if len(doc.sections) != 1:
        raise AssertionError("Unerwartete Abschnittsanzahl")
    section = doc.sections[0]
    for margin in (
        section.top_margin,
        section.bottom_margin,
        section.left_margin,
        section.right_margin,
    ):
        if abs(margin.cm - 2.5) > 0.02:
            raise AssertionError("Seitenrand ist nicht 2,5 cm")
    if not any(
        table.rows[0].cells[0].text == "Gruppe" and len(table.rows) == 9
        for table in doc.tables
    ):
        raise AssertionError("Das gruppierte Datenmodell fehlt")


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    doc = Document(SOURCE)
    set_document_defaults(doc)
    compact_cover(doc)
    apply_consistency_edits(doc)
    update_abbreviations(doc)
    build_canvas(doc)
    build_project_organization(doc)
    build_sprint_three(doc)
    build_workflow(doc)
    build_data_model(doc)
    build_results(doc)
    build_quality_and_dashboard(doc)
    update_static_toc(doc)
    mark_fields_for_update(doc)
    validate(doc)
    doc.core_properties.title = (
        "Prognose des monatlichen Energieverbrauchs und Erkennung von Anomalien"
    )
    doc.core_properties.subject = "IHK-Projektbericht mit ausgearbeitetem Machine-Learning-Layer"
    doc.save(OUTPUT)
    reopened = Document(OUTPUT)
    validate(reopened)
    print(f"Erstellt: {OUTPUT}")
    print(f"Absätze: {len(reopened.paragraphs)} | Tabellen: {len(reopened.tables)}")


if __name__ == "__main__":
    main()
