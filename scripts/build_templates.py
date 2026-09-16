"""Build editable Office templates, local TTFs, and a runnable notebook starter."""

from __future__ import annotations

import json
import shutil
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt as DocPt, RGBColor as DocColor
from fontTools.ttLib import TTFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DS = ROOT / "brand/design-system"
DEST = ROOT / "brand/templates"
DATA = json.loads((DS / "dist/tokens.resolved.json").read_text(encoding="utf-8"))
T = DATA["tokens"]
STAMP = datetime(2026, 9, 16, tzinfo=timezone.utc)


def font_name(token):
    return T[token].split(",")[0].strip('"')


SANS = font_name("font-sans")
MONO = font_name("font-mono")


def color(token):
    return RGBColor.from_string(T[token].lstrip("#"))


def build_fonts():
    for slug, weights in [("ibm-plex-sans", [400, 500, 600, 700]), ("geist-mono", [400, 500, 600])]:
        for weight in weights:
            source = DS / f"dist/fonts/{slug}-latin-{weight}-normal.woff2"
            target = DS / f"dist/fonts/{slug}-{weight}.ttf"
            font = TTFont(source, recalcTimestamp=False)
            font.flavor = None  # Lossless container conversion; glyphs and naming remain original.
            font.save(target)
            font.close()
            package = ROOT / "src/energy_analytics/visualization/fonts"
            package.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(target, package / target.name)
        shutil.copyfile(DS / f"dist/licenses/{slug}-LICENSE.txt", package / f"{slug}-LICENSE.txt")


def stable_zip(path):
    """Normalise ZIP timestamps so generated Office documents are reproducible."""
    with zipfile.ZipFile(path) as archive:
        contents = {name: archive.read(name) for name in archive.namelist()}
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in sorted(contents.items()):
            info = zipfile.ZipInfo(name, (2026, 9, 16, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, content)


def build_presentation():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = "SWW Projektvorlage"
    prs.core_properties.author = "StadtWerke Westhafen · Gruppe 6"
    prs.core_properties.created = prs.core_properties.modified = STAMP
    # Set Office theme fonts/colours, so newly inserted text and charts inherit the CI.
    from lxml import etree

    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    for part in prs.part.package.iter_parts():
        if str(part.partname).startswith("/ppt/theme/"):
            xml = etree.fromstring(part.blob)
            for node in xml.xpath(
                "//a:fontScheme/a:majorFont/a:latin | //a:fontScheme/a:minorFont/a:latin",
                namespaces=ns,
            ):
                node.set("typeface", SANS)
            mapping = {
                "dk1": "text-primary",
                "lt1": "surface-card",
                "dk2": "navy-800",
                "lt2": "surface-page",
                "accent1": "data-actual",
                "accent2": "teal-500",
                "accent3": "chart-kommunal",
                "accent4": "data-forecast",
                "accent5": "data-threshold",
                "accent6": "data-anomaly",
                "hlink": "text-link",
                "folHlink": "navy-700",
            }
            for name, token in mapping.items():
                slot = xml.find(f".//a:clrScheme/a:{name}", ns)
                for child in list(slot):
                    slot.remove(child)
                node = etree.SubElement(slot, "{" + ns["a"] + "}srgbClr")
                node.set("val", T[token].lstrip("#"))
            part._blob = etree.tostring(
                xml, xml_declaration=True, encoding="UTF-8", standalone=True
            )

    def box(slide, x, y, w, h, fill="surface-card"):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = color(fill)
        shape.line.color.rgb = color("border-default")
        return shape

    def text(slide, value, x, y, w, h, size=22, ink="text-primary", bold=False, mono=False):
        tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = 0
        tf.margin_top = tf.margin_bottom = 0
        for i, line in enumerate(value.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = MONO if mono else SANS
            p.font.size = Pt(size)
            p.font.bold = bold
            p.font.color.rgb = color(ink)
            p.space_after = Pt(12)
        return tf

    def slide(title, eyebrow, dark=False):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = color("navy-800" if dark else "surface-card")
        text(
            s,
            eyebrow.upper(),
            0.65,
            0.45,
            12,
            0.35,
            12,
            "teal-300" if dark else "text-accent",
            True,
        )
        text(s, title, 0.65, 1.0, 12, 1.25, 32, "text-inverse" if dark else "text-primary", True)
        text(
            s,
            "SWW · [Projekttitel] · [Datum]",
            0.65,
            7.02,
            10,
            0.22,
            10,
            "navy-200" if dark else "text-muted",
        )
        text(s, str(len(prs.slides)), 12, 7.02, 0.5, 0.22, 10, "navy-200" if dark else "text-muted")
        s.notes_slide.notes_text_frame.text = "Vorlage: Platzhalter und Demowerte durch belegte Projektergebnisse ersetzen. Schriften aus brand/design-system/dist/fonts bei Bedarf lokal installieren."
        return s

    s = slide("Verbrauchsprognose und\nAnomalie-Frühwarnung", "Datenanalyse · Gruppe 6", True)
    box(s, 0.65, 2.5, 4.05, 1.2)
    s.shapes.add_picture(
        str(DS / "assets/logo-sww-wordmark.png"), Inches(0.85), Inches(2.7), width=Inches(3.65)
    )
    text(s, "[Kernaussage des Projekts in einem Satz]", 0.65, 4.3, 11, 1, 26, "text-inverse")
    text(s, "[Team] · [Prüfungstermin] · [Version]", 0.65, 5.6, 11, 0.6, 18, "navy-200")
    s = slide("Der Weg von den Daten zur Entscheidung", "Agenda")
    for i, label in enumerate(
        [
            "Ausgangslage und Datenqualität",
            "Fragestellung und Modellierung",
            "Ergebnisse und Anomalien",
            "Empfehlungen und nächste Schritte",
        ]
    ):
        text(s, f"{i + 1:02d}", 0.75, 2.25 + i * 0.92, 1, 0.5, 26, "text-accent", True, True)
        text(s, label, 1.65, 2.25 + i * 0.92, 10, 0.7, 24)
    s = slide("[Kapitelüberschrift]", "02 · Abschnitt", True)
    text(
        s, "[Welche Frage wird in diesem Kapitel beantwortet?]", 0.65, 3.05, 11, 1.5, 28, "navy-200"
    )
    s = slide("[Eine klare Aussage zu den Kennzahlen]", "Ergebnisse")
    for i, (label, value, unit) in enumerate(
        [
            ("MAE", "[Wert]", "kWh"),
            ("RMSE", "[Wert]", "kWh"),
            ("R²", "[Wert]", ""),
            ("Auffälligkeiten", "[Wert]", "Zähler"),
        ]
    ):
        x = 0.65 + i * 3.05
        box(s, x, 2.35, 2.85, 2)
        text(s, label, x + 0.2, 2.6, 2.45, 0.4, 16, "text-muted")
        text(s, value, x + 0.2, 3.15, 2.45, 0.6, 26, "text-primary", True, True)
        text(s, unit, x + 0.2, 3.85, 2.45, 0.3, 14, "text-muted")
    text(
        s,
        "Bezugszeitraum: [MM/JJJJ–MM/JJJJ] · Testverfahren: [zeitlicher Split]",
        0.65,
        5.15,
        12,
        0.8,
        18,
        "text-secondary",
    )
    s = slide("Ist und Prognose gemeinsam beurteilen", "Diagramm · Synthetisches Beispiel")
    cd = CategoryChartData()
    cd.categories = ["Jan", "Feb", "Mär", "Apr", "Mai", "Jun"]
    cd.add_series("Ist", [14500, 13900, 12800, 11600, 12300, 13500])
    cd.add_series("Prognose", [14300, 14100, 12600, 11900, 12500, 13200])
    chart = s.shapes.add_chart(
        XL_CHART_TYPE.LINE, Inches(0.8), Inches(2.1), Inches(11.7), Inches(4.1), cd
    ).chart
    chart.has_legend = True
    chart.font.name = SANS
    chart.font.size = Pt(14)
    for series, token in zip(chart.series, ["data-actual", "data-forecast"]):
        series.format.line.color.rgb = color(token)
        series.format.line.width = Pt(2.25)
    from pptx.enum.dml import MSO_LINE_DASH_STYLE

    chart.series[1].format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    chart.value_axis.has_title = True
    chart.value_axis.axis_title.text_frame.text = "Verbrauch (kWh)"
    text(
        s,
        "Demodaten · Januar–Juni 2025 · Vor Präsentation durch echte Ergebnisse ersetzen.",
        0.8,
        6.42,
        12,
        0.35,
        13,
        "text-muted",
    )
    s = slide("[Baseline und Modell fair vergleichen]", "Vergleich")
    for x, title in [(0.65, "Baseline"), (6.85, "Modell")]:
        box(s, x, 2.25, 5.8, 3.7)
        text(s, title, x + 0.25, 2.55, 5.3, 0.5, 24, "text-brand", True)
        text(
            s,
            "[Methode]\n[MAE / RMSE im gleichen Testzeitraum]\n[Stärken und Grenzen]",
            x + 0.25,
            3.35,
            5.2,
            2.2,
            20,
        )
    s = slide("Die Modellentscheidung nachvollziehbar machen", "Methodik")
    for i, (title, body) in enumerate(
        [
            ("Daten", "[Quellen, Zeitraum, Bereinigung]"),
            ("Zielvariable", "[Definition und Einheit]"),
            ("Validierung", "[Train / Validierung / Test]"),
            ("Schwellwert", "[Kalibrierung auf Trainings-/Validierungsdaten]"),
            ("Bewertung", "[Baseline und Metriken]"),
            ("Betrieb", "[Prüfprozess, Grenzen und Monitoring]"),
        ]
    ):
        x = 0.65 + (i % 2) * 6.2
        y = 2.2 + (i // 2) * 1.4
        text(s, title, x, y, 5.8, 0.4, 20, "text-brand", True)
        text(s, body, x, y + 0.48, 5.8, 0.75, 17, "text-secondary")
    s = slide("[Befunde mit Quelle und Maßnahme]", "Tabelle")
    table = s.shapes.add_table(4, 3, Inches(0.65), Inches(2.25), Inches(12), Inches(3.45)).table
    rows = [
        ["Befund", "Beleg / Zeitraum", "Maßnahme"],
        ["[Befund 1]", "[Quelle]", "[Nächster Schritt]"],
        ["[Befund 2]", "[Quelle]", "[Nächster Schritt]"],
        ["[Befund 3]", "[Quelle]", "[Nächster Schritt]"],
    ]
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = color("navy-700" if r == 0 else "surface-page")
            for p in cell.text_frame.paragraphs:
                p.font.name = SANS
                p.font.size = Pt(18)
                p.font.color.rgb = color("text-inverse" if r == 0 else "text-primary")
    s = slide("[Die wichtigste Erkenntnis als Aussage]", "Kernaussage", True)
    text(
        s,
        "[Beleg, Beobachtung oder begründete Schlussfolgerung in zwei bis drei Sätzen.]",
        0.8,
        2.75,
        11.5,
        2.4,
        30,
        "text-inverse",
    )
    text(s, "Quelle: [Dokument / Auswertung / Zeitraum]", 0.8, 5.5, 11.5, 0.6, 16, "navy-200")
    s = slide("Aus der Analyse folgt ein klarer nächster Schritt", "Empfehlungen")
    text(
        s,
        "01  [Empfehlung und Verantwortliche]\n02  [Prüfschritt und Termin]\n03  [Offene Frage und Entscheidung]",
        0.8,
        2.35,
        11.5,
        3,
        24,
    )
    text(
        s,
        "Das Modell ersetzt keine Abrechnungsentscheidung — es flaggt nur Untersuchungswürdiges.",
        0.8,
        5.9,
        11.5,
        0.8,
        17,
        "text-secondary",
    )
    out = DEST / "presentations"
    out.mkdir(parents=True, exist_ok=True)
    pptx = out / "sww-project-template.pptx"
    prs.save(pptx)
    stable_zip(pptx)
    with (
        zipfile.ZipFile(pptx) as source,
        zipfile.ZipFile(out / "sww-project-template.potx", "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "[Content_Types].xml":
                content = content.replace(
                    b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
                    b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml",
                )
            target.writestr(entry, content)


def build_document():
    doc = Document()
    from lxml import etree

    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    for part in doc.part.package.parts:
        if "/theme/" in str(part.partname):
            xml = etree.fromstring(part.blob)
            for node in xml.xpath("//a:majorFont/a:latin | //a:minorFont/a:latin", namespaces=ns):
                node.set("typeface", SANS)
            for i, token in enumerate(
                [
                    "navy-700",
                    "teal-500",
                    "chart-kommunal",
                    "data-forecast",
                    "data-threshold",
                    "data-anomaly",
                ],
                1,
            ):
                slot = xml.find(f".//a:clrScheme/a:accent{i}", ns)
                for child in list(slot):
                    slot.remove(child)
                node = etree.SubElement(slot, "{" + ns["a"] + "}srgbClr")
                node.set("val", T[token].lstrip("#"))
            part._blob = etree.tostring(
                xml, xml_declaration=True, encoding="UTF-8", standalone=True
            )
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.4)
    sec.left_margin = sec.right_margin = Cm(2.5)
    for name, size, token in [
        ("Normal", 11, "text-primary"),
        ("Title", 28, "text-brand"),
        ("Heading 1", 20, "text-brand"),
        ("Heading 2", 15, "text-brand"),
        ("Heading 3", 12, "text-accent"),
    ]:
        style = doc.styles[name]
        style.font.name = SANS
        style.font.size = DocPt(size)
        style.font.color.rgb = DocColor.from_string(T[token].lstrip("#"))
    doc.styles["Normal"].paragraph_format.space_after = DocPt(8)
    sec.header.paragraphs[0].add_run().add_picture(
        str(DS / "assets/logo-sww-wordmark.png"), width=Cm(4.5)
    )
    footer = sec.footer.paragraphs[0]
    footer.text = "SWW · [Projekttitel] · [Version]"
    doc.add_heading("[Projekttitel]", 0)
    doc.add_paragraph("Verbrauchsprognose und Anomalie-Frühwarnung")
    doc.add_paragraph("[Team / Autorinnen und Autoren]\n[Datum] · [Version]")
    doc.add_paragraph(
        "Arbeitsvorlage im SWW-Design. Verbindliche IHK-Vorgaben und die mitgelieferte IHK-Berichtsvorlage haben Vorrang."
    )
    doc.add_page_break()
    for heading, body in [
        (
            "1 Ausgangslage und Ziel",
            "[Geschäftlicher Kontext, Fragestellung, Abgrenzung und Erfolgskriterien.]",
        ),
        (
            "2 Daten und Qualität",
            "[Quelle, Bezugszeitraum, Einheiten, Befunde und nachvollziehbare Bereinigungsschritte.]",
        ),
        (
            "3 Methodik",
            "[Baseline, Merkmale, zeitlicher Split, Verfahren und Kalibrierung des Anomalieschwellwerts.]",
        ),
        (
            "4 Ergebnisse",
            "[Metriken und Diagramme mit Einheiten, Zeitraum und Quellen. Demowerte vollständig ersetzen.]",
        ),
        (
            "5 Bewertung und Ausblick",
            "[Grenzen, Verantwortlichkeiten und empfohlene nächste Schritte.]",
        ),
        (
            "6 Quellen und Anhang",
            "[Datenherkunft, Literatur, Softwareversionen und reproduzierbare Ausführung.]",
        ),
    ]:
        doc.add_heading(heading, 1)
        doc.add_paragraph(body)
        if heading.startswith("4"):
            table = doc.add_table(rows=1, cols=3)
            table.style = "Light Shading Accent 1"
            for cell, label in zip(
                table.rows[0].cells, ["Metrik", "Wert / Einheit", "Testzeitraum"]
            ):
                cell.text = label
            for label in ["MAE", "RMSE", "R²"]:
                for cell, value in zip(table.add_row().cells, [label, "[Wert]", "[Zeitraum]"]):
                    cell.text = value
    doc.add_paragraph(
        "Das Modell ersetzt keine Abrechnungsentscheidung — es flaggt nur Untersuchungswürdiges."
    )
    doc.core_properties.author = "StadtWerke Westhafen · Gruppe 6"
    doc.core_properties.created = doc.core_properties.modified = STAMP
    out = DEST / "documents"
    out.mkdir(parents=True, exist_ok=True)
    target = out / "sww-report-template.docx"
    doc.save(target)
    stable_zip(target)


def build_notebook():
    cells = []

    def cell(kind, source):
        c = dict(
            cell_type=kind,
            id=f"sww-{len(cells):02d}",
            metadata={},
            source=source.splitlines(keepends=True),
        )
        if kind == "code":
            c.update(execution_count=None, outputs=[])
        cells.append(c)

    cell(
        "markdown",
        '# SWW · Analysevorlage\n\nProjekt: [Titel] · Autor: [Name] · Zeitraum: [MM/JJJJ–MM/JJJJ]\n\nDie Beispiele verwenden **synthetische Demodaten**, keine Projektergebnisse. Einrichtung vom Repository-Stamm: `python -m pip install -e ".[notebooks,export]"`. Anschließend denselben Python-Interpreter als Notebook-Kernel auswählen.\n',
    )
    cell(
        "code",
        "from IPython.display import HTML, display\nfrom energy_analytics.visualization import eda, theme\n\neda.setup()\ndisplay(HTML(theme.notebook_css()))\n",
    )
    cell(
        "markdown",
        "## Fragestellung und Daten\n\n[Fragestellung, Quelle, Bezugszeitraum und Einheit dokumentieren.]\n",
    )
    cell(
        "code",
        'monate = ["01/2025", "02/2025", "03/2025", "04/2025", "05/2025", "06/2025"]\nist = [14500, 13900, 12800, 11600, 12300, 13500]\nprognose = [14300, 14100, 12600, 11900, 12500, 13200]\nfig = eda.timeseries_forecast(monate, ist, prognose, title="Demodaten · Januar–Juni 2025", y_title="Verbrauch (kWh)")\nfig.show(renderer="notebook")  # Plotly wird offline eingebettet.\n',
    )
    cell(
        "code",
        'residuen = [actual - predicted for actual, predicted in zip(ist, prognose)]\neda.residual_bars(monate, residuen, schwelle=250, title="Demodaten · Residuen Januar–Juni 2025").show(renderer="notebook")\n',
    )
    cell(
        "markdown",
        '## Export und Bewertung\n\nPNG für Folien: `eda.save_for_slide(fig, "prognose.png")` (Kaleido und Chrome erforderlich). Für identische Exporttypografie die mitgelieferten TTF-Schriften lokal installieren.\n\n[Ergebnis, Einschränkungen und nächsten Prüfschritt festhalten.]\n\nDas Modell ersetzt keine Abrechnungsentscheidung — es flaggt nur Untersuchungswürdiges.\n',
    )
    notebook = dict(
        cells=cells,
        metadata=dict(
            kernelspec=dict(display_name="Python 3 (SWW)", language="python", name="python3"),
            language_info=dict(name="python", version="3.11"),
        ),
        nbformat=4,
        nbformat_minor=5,
    )
    for target in [
        DEST / "notebooks/sww-analysis-template.ipynb",
        ROOT / "notebooks/00_design_system.ipynb",
    ]:
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(
            json.dumps(notebook, ensure_ascii=False, indent=1) + "\n", encoding="utf-8"
        )


if __name__ == "__main__":
    build_fonts()
    build_presentation()
    build_document()
    build_notebook()
    print("Built local TTF fonts, editable PPTX/POTX, DOCX and notebook templates.")
