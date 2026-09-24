from __future__ import annotations

import math
import zipfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "IHK_Bericht_Gruppe_6_Patrick_ergaenzt.docx"
OUTPUT = ROOT / "docs" / "IHK_Bericht_Gruppe_6_final.docx"
ASSET_DIR = ROOT / ".build" / "ihk_report_final_assets"

NAVY = "17365D"
PALE_BLUE = "EAF2F8"
WHITE = "FFFFFF"
BLACK = "000000"
GRAY = "5B6575"
LIGHT_GRAY = "D9D9D9"


# These values are filled after the first render. Changing only the digits does not
# affect pagination, so a second deterministic build can materialize the final TOC.
TOC_PAGES = {
    "Abkürzungsverzeichnis": "1",
    "1   Ausgangssituation und Ist-Zustand": "2",
    "2   Machine Learning Canvas": "3",
    "3   Projektplanung und -durchführung": "4",
    "4   Workflow, Datenmodell und Ergebnisse": "10",
    "Literaturverzeichnis": "17",
    "Anhang A   Machine Learning Canvas": "18",
    "Anhang B   Datenmodell und Formeln": "19",
    "Anhang C   Datenverständnis und EDA": "22",
    "Anhang D   Modellierung und Evaluation": "25",
    "Anhang E   Dashboard-Prototyp": "29",
}


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = tr_pr.find(qn("w:tblHeader"))
    if node is None:
        node = OxmlElement("w:tblHeader")
        tr_pr.append(node)
    node.set(qn("w:val"), "true")


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = LIGHT_GRAY) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_borders_none(table) -> None:
    """Remove all visible borders, used for equation-number layouts."""
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def set_run_font(run, size: float = 12, *, bold: bool | None = None, italic: bool | None = None,
                 color: str = BLACK, name: str = "Arial") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_lang_de(run) -> None:
    r_pr = run._element.get_or_add_rPr()
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "de-DE")


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.27)
    section.footer_distance = Cm(1.27)
    section.different_first_page_header_footer = True

    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    # The cover is intentionally unnumbered; the visible numbering begins with 1.
    pg_num_type.set(qn("w:start"), "0")

    first_footer = section.first_page_footer
    for paragraph in first_footer.paragraphs:
        paragraph.clear()

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, value, end):
        run._r.append(element)
    set_run_font(run, 12)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True

    heading_specs = {
        "Heading 1": (18, 16, 8),
        "Heading 2": (14, 13, 6),
        "Heading 3": (12, 10, 4),
        "Heading 4": (12, 8, 3),
    }
    for name, (size, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Caption" in styles:
        caption = styles["Caption"]
        caption.font.name = "Arial"
        caption._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        caption.font.size = Pt(10)
        caption.font.italic = True
        caption.font.color.rgb = RGBColor.from_string(BLACK)
        caption.paragraph_format.space_before = Pt(3)
        caption.paragraph_format.space_after = Pt(4)


def truncate_after_cover(document: Document) -> None:
    toc = next(p for p in document.paragraphs if p.text.strip() == "Inhaltsverzeichnis")
    body = document._body._body
    started = False
    for element in list(body):
        if element is toc._p:
            started = True
        if started and element.tag != qn("w:sectPr"):
            body.remove(element)


def add_paragraph(document: Document, text: str = "", *, bold_lead: str | None = None,
                  style: str | None = None, space_after: float | None = None,
                  keep_with_next: bool = False, first_line: bool = False) -> object:
    paragraph = document.add_paragraph(style=style)
    if bold_lead:
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, 12, bold=True)
        set_lang_de(lead)
    if text:
        run = paragraph.add_run(text)
        set_run_font(run, 12)
        set_lang_de(run)
    paragraph.paragraph_format.keep_with_next = keep_with_next
    if first_line:
        paragraph.paragraph_format.first_line_indent = Cm(0.6)
    if space_after is not None:
        paragraph.paragraph_format.space_after = Pt(space_after)
    return paragraph


def add_heading(document: Document, text: str, level: int, *, page_break: bool = False):
    paragraph = document.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.page_break_before = page_break
    for run in paragraph.runs:
        set_run_font(run, {1: 18, 2: 14, 3: 12}.get(level, 12), bold=True)
        set_lang_de(run)
    return paragraph


def add_labeled_paragraph(document: Document, items: list[tuple[str, str]], *, space_after: float = 6):
    paragraph = document.add_paragraph()
    for index, (label, text) in enumerate(items):
        if index:
            spacer = paragraph.add_run("  ")
            set_run_font(spacer, 12)
        lead = paragraph.add_run(label)
        set_run_font(lead, 12, bold=True)
        set_lang_de(lead)
        run = paragraph.add_run(text)
        set_run_font(run, 12)
        set_lang_de(run)
    paragraph.paragraph_format.space_after = Pt(space_after)
    return paragraph


def add_bullet(document: Document, text: str):
    paragraph = add_paragraph(document, text)
    paragraph.paragraph_format.left_indent = Cm(0.7)
    paragraph.paragraph_format.first_line_indent = Cm(-0.45)
    paragraph.runs[0].text = "•\t" + paragraph.runs[0].text
    return paragraph


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths_cm: list[float],
              *, font_size: float = 11, line_spacing: float = 1.15, header_fill: str = NAVY):
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = False
    set_table_borders(table)
    for index, value in enumerate(headers):
        table.cell(0, index).text = value
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value

    for row_index, row in enumerate(table.rows):
        keep_row_together(row)
        if row_index == 0:
            set_repeat_table_header(row)
        for col_index, cell in enumerate(row.cells):
            cell.width = Cm(widths_cm[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, header_fill if row_index == 0 else (PALE_BLUE if row_index % 2 == 0 else WHITE))
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = line_spacing
                for run in paragraph.runs:
                    set_run_font(
                        run,
                        font_size,
                        bold=row_index == 0,
                        color=WHITE if row_index == 0 else BLACK,
                    )
                    set_lang_de(run)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_table_caption(document: Document, text: str):
    # The German source template does not necessarily contain Word's English
    # built-in "Caption" style, so apply the caption formatting explicitly.
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, 10, italic=True)
    set_lang_de(run)
    return paragraph


def add_toc(document: Document) -> None:
    heading = add_heading(document, "Inhaltsverzeichnis", 1, page_break=False)
    heading.paragraph_format.space_before = Pt(0)
    entries = [
        ("Abkürzungsverzeichnis", 0),
        ("1   Ausgangssituation und Ist-Zustand", 0),
        ("2   Machine Learning Canvas", 0),
        ("3   Projektplanung und -durchführung", 0),
        ("4   Workflow, Datenmodell und Ergebnisse", 0),
        ("Literaturverzeichnis", 0),
        ("Anhang A   Machine Learning Canvas", 0),
        ("Anhang B   Datenmodell und Formeln", 0),
        ("Anhang C   Datenverständnis und EDA", 0),
        ("Anhang D   Modellierung und Evaluation", 0),
        ("Anhang E   Dashboard-Prototyp", 0),
    ]
    table = document.add_table(rows=0, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(14.2)
    table.columns[1].width = Cm(1.8)
    for title, _level in entries:
        row = table.add_row()
        row.cells[0].width = Cm(14.2)
        row.cells[1].width = Cm(1.8)
        row.cells[0].text = title
        row.cells[1].text = TOC_PAGES.get(title, "")
        row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for cell in row.cells:
            set_cell_margins(cell, top=25, start=0, bottom=25, end=0)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.05
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_run_font(run, 11.5, bold=title.startswith(tuple("1234")))
                    set_lang_de(run)
    # Invisible borders for a restrained TOC layout.
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    tbl_pr.append(borders)

    add_heading(document, "Abkürzungsverzeichnis", 2)
    add_table(
        document,
        ["Abkürzung", "Bedeutung"],
        [
            ["CV", "Cross-Validation (Kreuzvalidierung)"],
            ["EDA", "Explorative Datenanalyse"],
            ["KPI", "Key Performance Indicator (Leistungskennzahl)"],
            ["MAE", "Mean Absolute Error (mittlerer absoluter Fehler)"],
            ["ML", "Machine Learning (maschinelles Lernen)"],
            ["RMSE", "Root Mean Squared Error (Wurzel des mittleren quadratischen Fehlers)"],
            ["VLS", "Vollaststunden"],
        ],
        [3.2, 12.8],
        font_size=10.5,
        line_spacing=1.0,
    )


def render_formula(filename: str, latex: str, *, width: float = 8.0, height: float = 0.72,
                   fontsize: int = 18) -> Path:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    path = ASSET_DIR / filename
    fig = plt.figure(figsize=(width, height), dpi=300)
    fig.patch.set_alpha(0)
    fig.text(0.5, 0.5, latex, ha="center", va="center", fontsize=fontsize, color="#111111")
    fig.savefig(path, transparent=True, bbox_inches="tight", pad_inches=0.08)
    plt.close(fig)
    return path


def add_formula(document: Document, image_path: Path, number: int, explanation: str,
                *, width_cm: float = 12.5):
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    set_table_borders_none(table)
    widths = (14.7, 1.3)
    for index, cell_width in enumerate(widths):
        table.columns[index].width = Cm(cell_width)
    for cell, cell_width in zip(table.rows[0].cells, widths):
        cell.width = Cm(cell_width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
        cell.paragraphs[0].paragraph_format.space_before = Pt(0)
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        cell.paragraphs[0].paragraph_format.keep_with_next = True
    formula_paragraph = table.cell(0, 0).paragraphs[0]
    formula_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    formula_paragraph.add_run().add_picture(str(image_path), width=Cm(min(width_cm, 13.8)))
    number_paragraph = table.cell(0, 1).paragraphs[0]
    number_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    number_run = number_paragraph.add_run(f"({number})")
    set_run_font(number_run, 11)
    explanation_paragraph = add_paragraph(document, explanation, space_after=5)
    explanation_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return table


def extract_embedded_images() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    wanted = {
        "word/media/image3.png": "eda_verteilung.png",
        "word/media/image4.png": "eda_kundentyp.png",
        "word/media/image10.png": "eda_physik.png",
        "word/media/image13.png": "eda_korrelation.png",
        "word/media/image5.png": "eda_produktion.png",
        "word/media/image1.png": "eda_wartung.png",
    }
    result: dict[str, Path] = {}
    with zipfile.ZipFile(SOURCE) as archive:
        for member, target_name in wanted.items():
            target = ASSET_DIR / target_name
            target.write_bytes(archive.read(member))
            result[target_name] = target
    return result


def add_figure(document: Document, image_path: Path, caption: str, source: str,
               *, max_width_cm: float = 15.8, max_height_cm: float = 8.4):
    with Image.open(image_path) as image:
        width_px, height_px = image.size
    aspect = width_px / height_px
    width_cm = min(max_width_cm, max_height_cm * aspect)
    height_cm = width_cm / aspect
    if height_cm > max_height_cm:
        height_cm = max_height_cm
        width_cm = height_cm * aspect

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_picture(str(image_path), width=Cm(width_cm), height=Cm(height_cm))

    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.keep_with_next = True
    run = caption_paragraph.add_run(caption)
    set_run_font(run, 10, italic=True)
    set_lang_de(run)

    source_paragraph = document.add_paragraph()
    source_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source_paragraph.paragraph_format.space_after = Pt(5)
    run = source_paragraph.add_run(source)
    set_run_font(run, 8.5, italic=True, color=GRAY)
    set_lang_de(run)


def add_source_note(document: Document, text: str):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(7)
    run = paragraph.add_run(text)
    set_run_font(run, 9, italic=True, color=GRAY)
    set_lang_de(run)
    return paragraph


def build_main_report(document: Document, formulas: dict[str, Path]) -> None:
    add_toc(document)

    add_heading(document, "1   Ausgangssituation und Ist-Zustand", 1, page_break=True)
    add_heading(document, "1.1   Vorstellung der Organisation", 2)
    add_paragraph(
        document,
        "Die StadtWerke Westhafen GmbH versorgt im Hamburger Hafengebiet 700 gewerbliche, industrielle und "
        "kommunale Großkunden. Der Jahresumsatz liegt bei rund 180 Mio. EUR. Für Energiebeschaffung und "
        "Netzbetrieb werden verlässliche Prognosen und eine zeitnahe Prüfung auffälliger Zählerwerte benötigt.",
    )

    add_heading(document, "1.2   Aktuelle Situation und Problemstellung", 2)
    add_paragraph(
        document,
        "Die Beschaffungsplanung beruht bisher weitgehend auf manuellen Fortschreibungen. Abweichungen können "
        "kurzfristige Käufe oder Verkäufe am Spotmarkt erforderlich machen. Auffällige Monatsverbräuche fallen "
        "oft erst bei der quartalsweisen Auswertung auf. Messfehler, Abrechnungsprobleme oder technische "
        "Störungen werden dadurch spät geprüft.",
    )
    add_paragraph(
        document,
        "Das Projekt prognostiziert deshalb je Zähler den Verbrauch des nächsten Monats. Sobald der Istwert "
        "vorliegt, wird die Abweichung berechnet. Eine ungewöhnlich große Differenz erzeugt einen Prüfhinweis. "
        "Dieser bestätigt keine Störung, sondern fordert eine Daten-, Kontext- und Fachprüfung aus.",
    )

    add_heading(document, "1.3   Bereits genutzte Analytics-Prozesse", 2)
    add_paragraph(
        document,
        "Die monatlichen Auswertungen unterstützen Energiebeschaffung und Netzmanagement, sind aber überwiegend "
        "rückblickend. Ein Folgemonatsmodell mit vorab festgelegter Fehlerschwelle fehlt. Das Projekt ergänzt den "
        "Ablauf um eine Prognose und eine begrenzte Prüfliste. Die fachliche Entscheidung bleibt bei den "
        "zuständigen Bereichen.",
    )

    add_heading(document, "1.4   Projektziel und Abgrenzung", 2)
    add_paragraph(
        document,
        "Ziel ist ein reproduzierbarer Python-Workflow für Folgemonatsprognose, Baselinevergleich und "
        "Prüfhinweise. Bewertet wird rückblickend mit Daten aus 2024 und 2025. Nicht untersucht werden eine "
        "Produktivanbindung, automatische Maßnahmen oder Euro-Einsparungen. Dafür wäre ein späterer Pilotbetrieb "
        "nötig.",
    )

    add_heading(document, "2   Machine Learning Canvas", 1, page_break=True)
    add_heading(document, "2.1   Überblick", 2)
    add_paragraph(
        document,
        "Der Machine Learning Canvas verbindet fachliche Aufgabe, Daten, Modell und spätere Nutzung. Die zehn "
        "Festlegungen werden nachfolgend beschrieben und in Tabelle A1 im Anhang zusammengefasst.",
    )

    add_heading(document, "2.2   Die zehn Felder", 2)
    add_labeled_paragraph(document, [
        ("1 Mehrwert. ", "Die Lösung soll die Folgemonatsplanung verbessern und auffällige Verbräuche monatlich sichtbar machen. "),
        ("2 Datenquellen. ", "Verwendet werden Zählerverbrauch, Vertragsleistung, Kundentyp, Kalender, Heizgradtage, Produktionsplan und geplante Wartung."),
    ])
    add_labeled_paragraph(document, [
        ("3 Vorhersage. ", "Der Verbrauch wird je Zähler und Monat in Vollaststunden prognostiziert und für die Bewertung in kWh zurückgerechnet. Nach Eingang des Istwerts wird die Abweichung geprüft. "),
        ("4 Merkmale. ", "Das Modell nutzt Historie, Saison, Kalender, Wetterprognose, Produktionsplan, Wartung und Kundentyp. Alle Informationen müssen zum Prognosezeitpunkt bekannt sein."),
    ])
    add_labeled_paragraph(document, [
        ("5 Lernansatz. ", "Die Aufgabe ist eine überwachte Regression. Verglichen werden eine lineare Regression und ein Random Forest. "),
        ("6 Evaluation. ", "Hauptmetrik ist der RMSE in kWh; MAE und R² ergänzen ihn. Baselines sind der Vormonatswert und das Mittel aus bis zu drei Vormonaten."),
    ])
    add_labeled_paragraph(document, [
        ("7 Entscheidung. ", "Die Prognose unterstützt die Energiebeschaffung. Große Abweichungen werden nach Monatsabschluss zur menschlichen Prüfung weitergegeben. "),
        ("8 Auswirkung. ", "Im Pilot sollen Prognosefehler, Hinweiszahl, Prüfzeit und Bestätigungsquote gemessen werden. Euro-Einsparungen sind nicht belegt. Für Klassifikationsmetriken fehlen bestätigte Labels."),
    ])
    add_labeled_paragraph(document, [
        ("9 Zeitpunkt. ", "Die Prognose entsteht vor Monatsbeginn; die Residuenprüfung folgt nach Eingang des Istwerts. "),
        ("10 Monitoring und Wartung. ", "Fehler, mögliche Drift und Hinweisvolumen werden monatlich kontrolliert. Neu trainiert wird erst bei dokumentiertem Bedarf."),
    ])
    add_heading(document, "3   Projektplanung und -durchführung", 1, page_break=True)
    add_heading(document, "3.1   Projektorganisation", 2)
    add_paragraph(
        document,
        "Die Gruppe bestand aus drei Mitgliedern, deren Aufgaben zu Projektbeginn nach fachlichen Schwerpunkten "
        "verteilt wurden. Iana Kraievska übernahm Datenmanagement und Datenqualität. Patrick Olmo Hederer "
        "bearbeitete die explorative Datenanalyse und die Visualisierung. Kiko Ramon Lukas war für Machine "
        "Learning und Evaluation verantwortlich. Diese Aufteilung blieb über alle drei Sprints bestehen. "
        "Entscheidungen, die mehrere Arbeitspakete betrafen, wurden gemeinsam getroffen. Tabelle 1 zeigt die "
        "Aufgaben und Beiträge im Überblick.",
    )
    add_table_caption(document, "Tabelle 1: Rollen und individuelle Beiträge")
    add_table(
        document,
        ["Person", "Verantwortung", "Beitrag im Projekt"],
        [
            ["Iana Kraievska", "Datenmanagement und Datenqualität", "Bereinigung, Datentypen, Einheiten, Fehlwerte und Plausibilitätsregeln"],
            ["Patrick Olmo Hederer", "Explorative Datenanalyse und Visualisierung", "EDA, Treiberanalyse, Diagramme und visuelle Aufbereitung"],
            ["Kiko Ramon Lukas", "Machine Learning und Evaluation", "ML Canvas, Zielgröße, Modellwahl, zeitliche Validierung, Kalibrierung und Anomalieprüfung"],
        ],
        [3.6, 5.4, 7.0],
        font_size=10.5,
        line_spacing=1.15,
    )

    add_heading(document, "3.2   Projektziel", 2)
    add_paragraph(
        document,
        "Ziel des Projekts war es, einen nachvollziehbaren Workflow von der Rohdatei bis zum Export der "
        "Ergebnisse für das Dashboard zu entwickeln. Der Folgemonatsverbrauch sollte für jeden Zähler "
        "prognostiziert und nach Eingang des Istwerts mit der Prognose verglichen werden. Für die Bewertung "
        "wurden drei Anforderungen festgelegt: ein zeitlich korrekter Vergleich mit zwei einfachen Baselines, "
        "eine begründete Modellwahl und eine vor dem Benchmarkjahr festgelegte Schwelle für Prüfhinweise. "
        "Dadurch lässt sich das Vorgehen überprüfen und später in einem Pilotbetrieb erproben.",
    )

    add_heading(document, "3.3   Agile Projektsteuerung mit SCRUM", 2)
    add_paragraph(
        document,
        "Die Projektarbeit wurde mit einem an die kurze Laufzeit angepassten SCRUM-Ansatz organisiert. Der "
        "Zeitraum von drei Wochen wurde in drei Sprints von jeweils einer Woche aufgeteilt. Für jeden Sprint "
        "legte die Gruppe ein Ziel, die anstehenden Aufgaben und das erwartete Ergebnis fest. Am Ende wurden "
        "der Arbeitsstand und notwendige Anpassungen in einer kurzen Retrospektive besprochen und im "
        "Sprint-Log dokumentiert. Offene Punkte gingen in den folgenden Sprint ein.",
    )

    add_heading(document, "3.3.1   Sprint 1 Datenbasis", 3)
    add_paragraph(
        document,
        "Ziel des ersten Sprints war es, aus den 16.830 Rohzeilen mit 700 Zählern und 24 Monaten eine "
        "konsistente Grundlage für die weitere Analyse und Modellierung zu erstellen. Zunächst wurden "
        "Kategorien und Datumsformate vereinheitlicht. 672 in MWh gespeicherte Verbrauchswerte wurden in kWh "
        "umgerechnet, 30 doppelte Zähler-Monat-Kombinationen entfernt und drei fehlende Zielwerte anhand der "
        "vorhandenen Zeitreihe nachvollziehbar rekonstruiert.",
    )
    add_paragraph(
        document,
        "Nach der Bereinigung umfasste die gemeinsame Datenbasis 16.800 Zähler-Monate. Die schwierigste Frage "
        "war der Umgang mit auffälligen Verbrauchswerten, weil plausible Lastspitzen nicht als Messfehler "
        "entfernt werden durften. Als Maßstab diente deshalb die technisch mögliche Monatsenergie aus "
        "Vertragsleistung und Monatsstunden. Zehn physikalisch unmögliche Werte aus 2024 wurden vom Training "
        "ausgeschlossen. Die zehn entsprechenden Fälle aus 2025 blieben im Benchmark enthalten und wurden "
        "vom Verfahren sämtlich als Prüfhinweis markiert.",
    )

    add_heading(document, "3.3.2   Sprint 2 EDA und Canvas", 3)
    add_paragraph(
        document,
        "Ziel des zweiten Sprints war es, die bereinigten Daten explorativ zu untersuchen und daraus die "
        "Zielgröße sowie die für das Modell verfügbaren Merkmale abzuleiten. Die Analyse zeigte, dass 92 % der "
        "Streuung des Rohverbrauchs zwischen den Zählern lagen und die Vertragsleistung das Verbrauchsniveau "
        "stark bestimmte. Nach der Normierung lagen die Mediane der Kundentypen nur noch zwischen 160 und 173 "
        "VLS-Stunden. Deshalb wurden Vollaststunden als interne Zielgröße festgelegt.",
    )
    add_paragraph(
        document,
        "Der Produktionsplan und die geplante Wartung erwiesen sich als fachlich relevante Informationen. Die "
        "Heizgradtage wurden anstelle der nahezu redundanten Temperaturvariable verwendet. Der Vorjahreswert "
        "wurde nicht als Modellmerkmal aufgenommen, weil für 2024 keine Historie aus dem Jahr 2023 vorlag. "
        "Scheinbare Wetter- und Feiertagseffekte wurden durch den ebenfalls saisonalen Produktionsplan "
        "überlagert. Die Zusammenhänge wurden deshalb innerhalb vergleichbarer Produktionsniveaus geprüft und "
        "nicht vorschnell als Ursache interpretiert.",
    )
    add_paragraph(document, "Die wichtigsten Nachweise aus der EDA sind in den Abbildungen C1 bis C6 im Anhang dokumentiert.")

    add_heading(document, "3.3.3   Sprint 3 Modellierung und Bericht", 3)
    add_paragraph(
        document,
        "Ziel des dritten Sprints war es, ein verständliches Regressionsmodell auszuwählen, zeitlich zu "
        "validieren und die Schwelle für spätere Prüfhinweise getrennt zu bestimmen. Eine lineare Regression "
        "und ein Random Forest wurden in drei vorwärts laufenden Folds aus 2024 mit dem Vormonatswert und dem "
        "Mittel aus bis zu drei Vormonaten verglichen. Für den Random Forest wurden acht Kombinationen von "
        "Hyperparametern ausschließlich innerhalb dieser Folds geprüft.",
    )
    add_paragraph(
        document,
        "Der Random Forest erzielte mit 13.272 kWh den niedrigsten mittleren CV-RMSE und wurde deshalb als "
        "finales Modell ausgewählt. Im retrospektiven Benchmark 2025 erreichte er einen RMSE von 9.188 kWh. "
        "Damit lag sein Fehler 15,9 % unter dem der besten einfachen Baseline. Für die Kalibrierung wurden "
        "1.397 rollierend erzeugte Prognosefehler aus November und Dezember 2024 verwendet. Das 99. Perzentil "
        "lag bei 144,4 VLS-Stunden.",
    )
    add_paragraph(
        document,
        "Modellauswahl, Schwellenkalibrierung und Benchmark wurden zeitlich voneinander getrennt. Außerdem "
        "wurden die historischen Merkmale für jeden Zähler ausschließlich aus vergangenen Monaten berechnet, "
        "damit keine Informationen aus der Zukunft in das Modell gelangen. Diese Trennung verhindert Data "
        "Leakage. Die Zeitaufteilung, die Modellergebnisse und die Wirkung der Schwelle sind in den Abbildungen "
        "D1 bis D7 dokumentiert.",
    )

    add_heading(document, "3.3.4   Finalisierungsphase", 3)
    add_paragraph(
        document,
        "In der Finalisierungsphase wurden Bericht, Präsentation und Live-Demo auf denselben Stand gebracht. "
        "Kennzahlen und Begriffe wurden noch einmal abgeglichen, damit in allen drei Formaten dieselben "
        "Ergebnisse erklärt werden. Bei der Überarbeitung zeigte sich, dass nicht alle Analysen aus den drei "
        "Sprints in der Präsentation verständlich dargestellt werden konnten.",
    )
    add_paragraph(
        document,
        "Die Gruppe konzentrierte sich deshalb auf zwei Ergebnisse. Im retrospektiven Benchmark 2025 lag der "
        "RMSE des Random Forest 15,9 % unter dem der besten einfachen Baseline. Mit der gewählten Schwelle "
        "entstanden durchschnittlich 9,5 Prüfhinweise pro Monat. Die Präsentation folgt dem tatsächlichen "
        "Projektverlauf vom Geschäftsproblem über Datenbereinigung und EDA bis zur Modellbewertung und zur "
        "Prüfempfehlung.",
    )
    add_paragraph(
        document,
        "Da das Fachgespräch einzeln geführt wird, bereitete sich jedes Gruppenmitglied auch auf Fragen "
        "außerhalb des eigenen Schwerpunkts vor. Gemeinsam geübt wurden insbesondere die zeitliche "
        "Validierung, die Bedeutung des 99. Perzentils und der Unterschied zwischen einem statistischen "
        "Prüfhinweis und einer fachlich bestätigten Anomalie.",
    )

    add_heading(document, "3.4   User Stories", 2)
    add_labeled_paragraph(document, [("Stefan Lechtenberg – Bereichsleiter Energiebeschaffung. ", "Als Bereichsleiter möchte ich vor Beginn des nächsten Monats eine verlässliche Verbrauchsprognose erhalten, damit die benötigte Energiemenge besser geplant werden kann.")])
    add_labeled_paragraph(document, [("Anke Bürger – Leiterin Netzmanagement. ", "Als Leiterin möchte ich nach dem Monatsabschluss eine priorisierte Liste auffälliger Zähler erhalten, damit mein Team die möglichen Ursachen gezielt prüfen kann.")])
    add_labeled_paragraph(document, [("Henrik Maaß – Senior-Datenanalyst. ", "Als Senior-Datenanalyst möchte ich einen reproduzierbaren Daten- und ML-Workflow mit einer dokumentierten Schwelle entwickeln, damit die Ergebnisse überprüft und bei Bedarf erneut erzeugt werden können.")])

    add_heading(document, "3.5   Meilensteine und Retrospektiven", 2)
    add_table_caption(document, "Tabelle 2: Meilensteine, Erkenntnisse und Anpassungen")
    add_table(
        document,
        ["Zeitpunkt", "Meilenstein", "Retrospektive und Anpassung"],
        [
            ["Ende Sprint 1", "Bereinigte gemeinsame Datenbasis", "Die parallel geführten Zwischenstände wichen voneinander ab. Deshalb wurden eine verbindliche Datei und eine gemeinsame Plausibilitätsregel festgelegt."],
            ["Ende Sprint 2", "EDA-Vorgaben und ML Canvas", "Der Rohverbrauch gab großen Zählern bei der Bewertung zu viel Gewicht. Deshalb wurde VLS als normierte Zielgröße gewählt."],
            ["Ende Sprint 3", "Gewähltes Modell und kalibrierte Schwelle", "Das gelieferte Drei-Monats-Mittel enthielt Informationen, die zum jeweiligen Zeitpunkt noch nicht verfügbar waren. Es wurde deshalb je Zähler nur aus vergangenen Werten neu berechnet."],
            ["Finalisierung", "Bericht, Präsentation und Demo", "Kennzahlen und Begriffe wurden in Bericht, Präsentation und Demo vereinheitlicht. Die fachlichen Grenzen wurden in allen drei Formaten benannt."],
        ],
        [3.1, 4.7, 8.2],
        font_size=10.5,
        line_spacing=1.15,
    )

    add_heading(document, "3.6   Persönliches Fazit zu den Lernerfahrungen", 2)
    add_labeled_paragraph(document, [("Iana Kraievska – Datenmanagement und Datenqualität. ", "Ich habe gelernt, dass Datenbereinigung keine rein technische Routine ist. Viele Entscheidungen lassen sich erst treffen, wenn die Werte fachlich eingeordnet werden. Das zeigte sich besonders bei den auffälligen Verbrauchswerten. Ein automatisches Entfernen nach einer statistischen Regel hätte auch reale industrielle Lastspitzen erfassen können. Erst der Vergleich mit der technisch möglichen Monatsenergie aus Vertragsleistung und Monatsstunden erlaubte eine begründete Trennung zwischen Messfehlern und plausiblen Verbrauchswerten. Wichtig war für mich außerdem, dass fehlende Historie nicht automatisch imputiert werden darf.")])
    add_labeled_paragraph(document, [("Patrick Olmo Hederer – Explorative Datenanalyse und Visualisierung. ", "Ich habe gelernt, dass eine auffällige Korrelation in der EDA noch keinen Verbrauchstreiber beweist. Besonders deutlich wurde das bei den Wetter- und Feiertagsmerkmalen. Erst die Prüfung innerhalb vergleichbarer Produktionsniveaus zeigte, dass ein Teil der Zusammenhänge durch die saisonale Produktion überlagert war. Aus der EDA entstand außerdem die Entscheidung, den Verbrauch auf Vollaststunden zu normieren, weil 92 % der Streuung des Rohverbrauchs zwischen den Zählern lagen. Die Visualisierung war deshalb nicht nur die Darstellung fertiger Ergebnisse, sondern ein Teil der eigentlichen Analyse.")])
    add_labeled_paragraph(document, [("Kiko Ramon Lukas – Machine Learning und Evaluation. ", "Ich habe gelernt, dass bei Zeitreihendaten die Validierung wichtiger ist als die Wahl des komplexesten Modells. Besonders herausfordernd war die Vermeidung von Data Leakage. Das gelieferte Drei-Monats-Mittel musste für jeden Zähler ausschließlich aus vergangenen Werten neu berechnet werden. Auch die Modellauswahl und die Kalibrierung der Schwelle mussten zeitlich vom Benchmark getrennt bleiben. Baselines und lineare Regression waren wichtige Vergleichspunkte. Nur so ließ sich zeigen, dass der Random Forest 2025 einen um 15,9 % niedrigeren RMSE als die beste einfache Baseline erreichte.")])

    add_heading(document, "3.7   Ausblick", 2)
    add_paragraph(
        document,
        "Für die StadtWerke Westhafen GmbH empfiehlt sich zunächst ein prospektiver Schattenbetrieb mit einer "
        "ausgewählten Gruppe von Zählern. Die Modellprognosen würden dabei parallel zum bisherigen Prozess "
        "erstellt, ohne direkt in die Beschaffung oder das Netzmanagement einzugreifen. Während dieser Phase "
        "sollten Prognosefehler, Anzahl der Prüfhinweise, Bearbeitungszeit, Datenqualität und die fachlich "
        "bestätigten Ursachen dokumentiert werden. Erst wenn die Ergebnisse über mehrere neue Monate stabil "
        "sind, sollte die Lösung in die bestehenden Prozesse integriert und auf weitere Zähler ausgeweitet werden.",
    )
    add_labeled_paragraph(document, [
        ("Kiko Ramon Lukas – persönlicher Ausblick. ", "Mein nächster Schritt ist die Entwicklung eines strukturierten Feedbackprozesses im Dashboard. Nach jeder Prüfung soll gespeichert werden, ob sich ein Hinweis fachlich bestätigt hat und welche Ursache festgestellt wurde. Sobald genügend belastbare Ja/Nein-Labels vorliegen, könnte das Regressionsmodell um ein Klassifikationsmodell ergänzt werden. Das Regressionsmodell würde weiterhin den erwarteten Verbrauch liefern. Das Klassifikationsmodell würde anschließend die Wahrscheinlichkeit einer fachlichen Bestätigung schätzen. Eine solche Prozentangabe wäre erst nach einer eigenen Validierung und Wahrscheinlichkeitskalibrierung belastbar. Bewertet werden müsste das Klassifikationsmodell unter anderem mit Precision und Recall."),
    ])

    add_heading(document, "4   Workflow, Datenmodell und Ergebnisse", 1, page_break=True)
    add_heading(document, "4.1   Workflow-Dokumentation", 2)
    workflow_steps = [
        ("(1) Quelle. ", "Grundlage ist die Datei verbrauch.csv mit 16.830 Monatswerten von 700 Zählern aus dem Zeitraum Januar 2024 bis Dezember 2025. Neben dem Verbrauch enthält sie Angaben zur Vertragsleistung, zum Kundentyp, zu Kalender und Wetter sowie zu Produktionsplanung und Wartung."),
        ("(2) Import und Prüfung. ", "Die Datei wurde mit pandas in Python eingelesen. Danach wurden Spalten, Datentypen und Datumswerte geprüft. Zusätzlich wurde kontrolliert, ob jede Kombination aus Zähler und Monat eindeutig ist und ob für alle Zähler die erwarteten Monate vorliegen."),
        ("(3) Bereinigung. ", "Bei der Bereinigung wurden Kundentypen und Datumsformate vereinheitlicht, Werte in MWh nach kWh umgerechnet und 30 doppelte Zähler-Monat-Kombinationen entfernt. Fehlende Werte wurden nach ihrer Ursache behandelt und nicht pauschal ersetzt. Die zehn physikalisch unmöglichen Werte aus 2024 wurden vom Training ausgeschlossen. Hohe, aber technisch mögliche Verbrauchswerte blieben erhalten."),
        ("(4) Transformation und EDA. ", "Für die Modellierung wurden aus Verbrauch und Vertragsleistung die Vollaststunden berechnet. Die historischen Merkmale wurden für jeden Zähler zeitlich verschoben. Dadurch enthält eine Zeile nur Verbrauchswerte aus bereits vergangenen Monaten. Die explorative Datenanalyse diente anschließend dazu, das endgültige Merkmalset festzulegen."),
        ("(5) Modellierung. ", "Die Modellauswahl erfolgte mit drei zeitlichen Folds aus dem Jahr 2024. Danach wurden die Prognosefehler für November und Dezember 2024 verwendet, um die Schwelle für Prüfhinweise festzulegen. Das ausgewählte Modell wurde anschließend mit allen gültigen Daten aus 2024 neu trainiert. Modellkonfiguration und Schwelle blieben für den Benchmark 2025 unverändert. Die historischen Merkmale wurden im Verlauf des Jahres jeweils mit den zu diesem Zeitpunkt bekannten Vormonatswerten aktualisiert."),
        ("(6) Visualisierung. ", "Für die Ergebnisse wurden einheitliche Diagramme im Corporate Design erstellt. Sie zeigen die zeitliche Aufteilung, den Modellvergleich, die gewählte Schwelle und einzelne Prüffälle. Die ausführlichen Darstellungen befinden sich in den Anhängen C und D."),
        ("(7) Export. ", "Prognosen, Istwerte, Residuen, Schwellenfaktor, Abweichungsrichtung und Kennzahlen werden als JSON-Datei an den lokalen Dashboard-Prototyp übergeben. Vor dem Export prüft das Skript, ob alle Pflichtfelder vorhanden und die Schlüssel eindeutig sind."),
    ]
    for label, text in workflow_steps:
        add_labeled_paragraph(document, [(label, text)], space_after=4)

    add_heading(document, "4.2   Datenmodell und Feature Engineering", 2)
    add_paragraph(
        document,
        "Die Modellierungsbasis enthält für jeden Zähler und Monat genau eine Zeile. Tabelle B1 beschreibt "
        "alle 20 Spalten. Für das Training wurden neun Merkmale aus fünf Informationsgruppen verwendet. Die "
        "Vertragsleistung wurde nicht als Merkmal trainiert. Sie wird zur Berechnung der Vollaststunden und "
        "zur späteren Rückrechnung in kWh benötigt.",
    )
    add_table_caption(document, "Tabelle 3: Tatsächlich verwendete Modellmerkmale")
    add_table(
        document,
        ["Gruppe", "Merkmale", "Begründung"],
        [
            ["Historie", "vormonat_vls, letzte_3_monate_vls", "Zuletzt bekannte Auslastung und geglätteter Verlauf"],
            ["Kalender", "monat_idx, arbeitstage, feiertage_im_monat", "Saisonale Lage und nutzbare Arbeitstage"],
            ["Wetter", "heizgradtage", "Vor Monatsbeginn als Prognosewert bereitstellbarer Wärmebedarf"],
            ["Planung", "produktionsplan_index, wartung_aktiv", "Erwartete Aktivität und geplanter Stillstand"],
            ["Kundengruppe", "kundentyp", "Unterschiede zwischen Gewerbe, Industrie und Kommunal"],
        ],
        [2.6, 6.1, 7.3],
        font_size=10.5,
        line_spacing=1.1,
    )
    add_formula(
        document,
        formulas["vls"],
        1,
        "Beispiel: Der Zähler ZL-00000 verbrauchte im Januar 2024 insgesamt 8.179 kWh bei einer "
        "Vertragsleistung von 74 kW. Daraus ergeben sich 8.179 kWh ÷ 74 kW = 110,5 Vollaststunden. Wird der "
        "ungerundete Wert wieder mit 74 kW multipliziert, erhält man erneut 8.179 kWh. Die Vollaststunden "
        "dienen hier als Vergleichsgröße und sind keine gemessene Laufzeit.",
    )
    add_paragraph(
        document,
        "Fehlende numerische Merkmale wurden innerhalb jedes Trainingsfolds durch den Median der jeweiligen "
        "Trainingsdaten ersetzt. Zusätzlich erhielt das Modell ein Kennzeichen dafür, dass der ursprüngliche "
        "Wert fehlte. Der Kundentyp wurde per One-Hot-Kodierung umgewandelt. Damit wurden für die "
        "Vorverarbeitung eines Bewertungszeitraums nur Informationen aus dem zugehörigen Training verwendet.",
    )

    add_heading(document, "4.3   Ergebnisse", 2)
    add_heading(document, "4.3.1   Zielvariable, Validierung und Hauptmetrik", 3)
    add_paragraph(
        document,
        "Da die Daten zeitlich geordnet sind, wurde keine zufällige Kreuzvalidierung verwendet. Fold 1 lernte "
        "aus Januar bis April 2024 und bewertete Mai und Juni. Fold 2 lernte aus Januar bis Juni und bewertete "
        "Juli und August. Fold 3 lernte aus Januar bis August und bewertete September und Oktober. So liegt "
        "jeder Bewertungsmonat zeitlich nach den Monaten, aus denen das Modell gelernt hat. Abbildung D1 zeigt "
        "diese Aufteilung sowie die anschließende Kalibrierung und den Benchmark 2025.",
    )
    add_formula(
        document,
        formulas["rmse"],
        2,
        "Für jeden bewerteten Zähler-Monat wurde die Differenz zwischen tatsächlichem und prognostiziertem "
        "Verbrauch berechnet. Vorher wurden die VLS-Prognosen wieder in kWh umgerechnet. Beim RMSE werden die "
        "Fehler quadriert, wodurch große Abweichungen stärker in das Ergebnis eingehen. Für die 8.398 "
        "auswertbaren Zähler-Monate des Jahres 2025 beträgt der RMSE 9.188 kWh. MAE und R² werden ergänzend "
        "angegeben.",
    )

    add_heading(document, "4.3.2   Modellwahl im Jahr 2024", 3)
    add_paragraph(
        document,
        "Trainiert wurden eine lineare Regression und ein Random Forest. Der Vormonatswert und das Mittel aus "
        "bis zu drei Vormonaten dienten als Baselines und wurden nicht trainiert. Für den Random Forest wurden "
        "mit ParameterGrid acht vorher festgelegte Kombinationen systematisch geprüft. Die Auswahl beruhte "
        "ausschließlich auf den drei zeitlichen Folds aus 2024. Die Daten aus 2025 wurden dafür nicht verwendet. "
        "Die beste Kombination verwendet 300 Bäume, eine maximale Tiefe von 8, mindestens 5 Beobachtungen pro "
        "Blatt, 70 % der Merkmale je Aufteilung und den Zufallsstartwert 42.",
    )
    add_table_caption(document, "Tabelle 4: Zeitliche Modellwahl 2024")
    add_table(
        document,
        ["Kandidat", "Mittlerer CV-RMSE", "Fold-Streuung"],
        [
            ["Random Forest", "13.272 kWh", "3.843 kWh"],
            ["Lineare Regression", "13.643 kWh", "3.472 kWh"],
            ["Bis-zu-3-Monats-Mittel", "15.483 kWh", "3.070 kWh"],
            ["Vormonat", "18.460 kWh", "6.366 kWh"],
        ],
        [7.0, 4.5, 4.5],
        font_size=11,
        line_spacing=1.1,
    )
    add_paragraph(
        document,
        "Im Durchschnitt über die drei Folds erreichte der Random Forest mit 13.272 kWh den niedrigsten RMSE. "
        "Die lineare Regression folgte mit 13.643 kWh. Der Abstand von 2,7 % ist gering, beide Modelle lagen "
        "jedoch vor der besten Baseline. Die angegebene Fold-Streuung ist die Standardabweichung der drei "
        "Einzelergebnisse. Die Whisker in Abbildung D2 zeigen diese Streuung und kein Konfidenzintervall.",
    )

    add_heading(document, "4.3.3   Retrospektiver Benchmark 2025", 3)
    add_paragraph(
        document,
        "Mit der ausgewählten Konfiguration wurde anschließend das Jahr 2025 bewertet. Für jeden Monat standen "
        "nur die Informationen zur Verfügung, die zu diesem Zeitpunkt bekannt gewesen wären. Insgesamt "
        "konnten 8.398 Zähler-Monate ausgewertet werden. Der Random Forest erreichte einen RMSE von 9.188 kWh, "
        "einen MAE von 3.725 kWh und ein R² von 0,901. Die lineare Regression erreichte 9.414 kWh RMSE, das "
        "Mittel aus bis zu drei Vormonaten 10.922 kWh und der Vormonatswert 12.386 kWh. Gegenüber der besten "
        "Baseline war der RMSE des Random Forest damit 15,9 % niedriger. Abbildung D3 zeigt den Vergleich.",
    )
    add_paragraph(
        document,
        "Um den Einfluss der Zielgröße zu prüfen, wurde derselbe Modelltyp zusätzlich direkt auf kWh trainiert. "
        "Diese Variante erreichte 9.800 kWh RMSE. Das über Vollaststunden trainierte Modell lag nach der "
        "Rückrechnung mit 9.188 kWh um 6,2 % niedriger. Dieses Ergebnis gilt für den verwendeten Datensatz und "
        "Versuchsaufbau. Eine allgemeine Überlegenheit der Vollaststunden lässt sich daraus nicht ableiten "
        "(Abbildung D4).",
    )

    add_heading(document, "4.3.4   Einfluss der Informationsgruppen", 3)
    add_paragraph(
        document,
        "Für die Permutationsanalyse wurden die Werte einer Informationsgruppe zufällig vertauscht, während "
        "das trainierte Modell unverändert blieb. Steigt der RMSE danach deutlich, hat das Modell die "
        "Informationen dieser Gruppe für seine Prognosen genutzt. Für die Auswertung wurden die fünf "
        "Trainingsgruppen weiter unterteilt. Produktionsplan und Wartung wurden getrennt betrachtet, ebenso "
        "Jahreszeit und die übrigen Kalendermerkmale.",
    )
    add_paragraph(
        document,
        "Das Mischen der Verbrauchshistorie erhöhte den RMSE um 7.256 kWh. Für den Produktionsplan betrug der "
        "Anstieg 798 kWh, für Wartung 591 kWh, für Kalender 276 kWh, für Wetter 154 kWh, für Jahreszeit 121 "
        "kWh und für den Kundentyp rund 3 kWh. Das Modell nutzt damit vor allem die Verbrauchshistorie. Die "
        "Werte zeigen jedoch keine Ursache-Wirkungs-Beziehung. Bei korrelierten Merkmalen kann sich die "
        "gemessene Bedeutung außerdem auf mehrere Gruppen verteilen. Tabelle D1 enthält die Ergebnisse.",
    )

    add_heading(document, "4.3.5   Kalibrierung und Prüfhinweise", 3)
    add_paragraph(
        document,
        "Nach der Modellwahl wurde die Schwelle für Prüfhinweise getrennt bestimmt. Dafür entstanden 1.397 "
        "Prognosefehler aus November und Dezember 2024. Die Novemberwerte wurden mit einem Modell prognostiziert, "
        "das aus Januar bis Oktober gelernt hatte. Für Dezember kamen die Daten von Januar bis November zum "
        "Einsatz. Das 99. Perzentil der absoluten VLS-Fehler liegt bei 144,4 Stunden. Rund 99 % der "
        "Kalibrierungsfehler sind damit höchstens so groß. Abbildung D5 zeigt die sortierten Fehler und die Lage "
        "dieser Schwelle.",
    )
    add_formula(
        document,
        formulas["alert"],
        3,
        "Beim Zähler ZL-00147 lag der tatsächliche Wert im August 2025 bei 480,9 VLS-Stunden, die Prognose bei "
        "149,1 VLS-Stunden. Das absolute Residuum beträgt somit 331,8 VLS-Stunden. Geteilt durch die Schwelle "
        "von 144,4 ergibt sich ein Schwellenfaktor von 2,30. Da dieser Wert mindestens 1 beträgt, wird der Fall "
        "als Prüfhinweis markiert. Der Faktor ist keine Wahrscheinlichkeit. Die entsprechende Grenze in kWh "
        "hängt von der Vertragsleistung des jeweiligen Zählers ab.",
        width_cm=14.7,
    )
    add_paragraph(
        document,
        "Die Schwelle gehört nicht zu den Hyperparametern des Random Forest. Sie wird erst nach der Modellwahl "
        "aus separaten Prognosefehlern berechnet und verändert das Training nicht. Das 99. Perzentil bedeutet "
        "außerdem nicht, dass in späteren Daten genau 1 % aller Fälle einen Prüfhinweis erhalten.",
    )
    add_paragraph(
        document,
        "Mit dieser Schwelle entstanden im Jahr 2025 insgesamt 114 Prüfhinweise für 107 Zähler. Das entspricht "
        "1,36 % der 8.398 auswertbaren Fälle und durchschnittlich 9,5 Hinweisen pro Monat. Bei 68 Fällen lag "
        "der Istwert über der Prognose, bei 46 Fällen darunter. Zehn Hinweise betrafen gleichzeitig einen "
        "gekennzeichneten Datenqualitätsfall. Die Wahl des Perzentils verändert den Arbeitsaufwand deutlich: "
        "Bei 95 %, 97,5 %, 99 % und 99,5 % ergeben sich im Mittel 33,3, 21,3, 9,5 beziehungsweise 5,8 Hinweise "
        "pro Monat. Das 99. Perzentil wurde für den Pilot als handhabbarer Ausgangswert gewählt. Es ist keine "
        "mathematisch einzig richtige Grenze. Abbildung D6 zeigt den Vergleich; Abbildung D7 stellt Istwert und "
        "Prognose aller Fälle gegenüber.",
    )

    add_heading(document, "4.4   Bewertung der Ergebnisqualität", 2)
    add_heading(document, "4.4.1   Aussagekraft", 3)
    add_paragraph(
        document,
        "Der Random Forest erreichte in der zeitlichen Validierung 2024 und im Benchmark 2025 einen niedrigeren "
        "Fehler als die Baselines. Ein deutliches Zeichen für Overfitting ist damit nicht zu erkennen. Bei der "
        "kurzen Zeitreihe lässt es sich jedoch nicht vollständig ausschließen. Die Ergebnisse aus 2024 und 2025 "
        "sind außerdem nicht wie ein gewöhnlicher Trainings- und Testfehler direkt miteinander vergleichbar, "
        "da die beiden Jahre unterschiedliche Lastsituationen enthalten.",
    )
    add_paragraph(
        document,
        "Das R² von 0,901 beschreibt, wie viel Streuung das Modell gegenüber einer einfachen "
        "Mittelwertprognose erklärt. Es ist keine Trefferquote und keine Wahrscheinlichkeit für eine Anomalie. "
        "Der RMSE liegt deutlich über dem MAE, weil einzelne große Fehler beim RMSE stärker gewichtet werden.",
    )
    add_heading(document, "4.4.2   Grenzen", 3)
    add_paragraph(
        document,
        "Die Daten aus 2025 waren zum Zeitpunkt der Projektauswertung bereits bekannt. Der Benchmark ist deshalb "
        "retrospektiv und kein künftig ungesehener Blindtest. Die Modellkonfiguration blieb nach der Auswahl mit "
        "den Daten aus 2024 unverändert. Die Lag-Merkmale wurden jedoch für jeden Monat mit den jeweils bekannten "
        "Vormonatswerten aktualisiert, wie es auch bei einer späteren Anwendung vorgesehen wäre.",
    )
    add_paragraph(
        document,
        "Die Schwelle wurde nur aus den Fehlern für November und Dezember 2024 berechnet. Ob sie auch in anderen "
        "Jahreszeiten gleich gut passt, muss mit neuen Daten geprüft werden. Außerdem fehlen vollständige "
        "fachliche Rückmeldungen dazu, welche Hinweise tatsächlich relevante Anomalien waren. Precision, Recall "
        "und eine Defektwahrscheinlichkeit können deshalb noch nicht belastbar berechnet werden.",
    )
    add_heading(document, "4.4.3   Eignung für den Pilotbetrieb", 3)
    add_paragraph(
        document,
        "Für einen Pilotbetrieb kann das Verfahren als Priorisierungshilfe eingesetzt werden. Es trifft keine "
        "automatische Fachentscheidung, sondern zeigt, welche Zähler-Monate zuerst geprüft werden sollten. Im "
        "Pilot sollten RMSE und MAE, Veränderungen der Merkmale, die Zahl der Prüfhinweise, die Bearbeitungszeit "
        "und die fachliche Bestätigungsquote monatlich beobachtet werden. Ein erneutes Training oder eine "
        "Änderung der Schwelle sollte erst erfolgen, wenn sich eine Verschlechterung über mehrere Zeiträume "
        "zeigt und die neue Einstellung wieder getrennt geprüft wurde.",
    )

    add_heading(document, "4.5   Visualisierung und Ergebnisübergabe", 2)
    add_paragraph(
        document,
        "Der lokale Dashboard-Prototyp zeigt für jeden Prüffall die Prognose, den Istwert, den Schwellenfaktor, "
        "die Richtung der Abweichung und den Bearbeitungsstatus. Das Perzentil kann als Szenario verändert "
        "werden. Dadurch wird direkt sichtbar, wie sich die Zahl der Prüfhinweise und die markierten Zähler "
        "ändern. Der Prototyp verwendet den retrospektiven Export aus 2025 und besitzt keine Live-Datenanbindung. "
        "Nach einer fachlichen Prüfung könnte im Dashboard zusätzlich gespeichert werden, ob ein Hinweis "
        "bestätigt oder verworfen wurde. Diese Rückmeldungen wären die Datengrundlage für das in Abschnitt 3.7 "
        "beschriebene Klassifikationsmodell.",
    )


def build_literature(document: Document) -> None:
    add_heading(document, "Literaturverzeichnis", 1, page_break=True)
    add_paragraph(
        document,
        "DIHK-Bildungs-GmbH (2026). Data Analyst (IHK) – Online-Zertifikatslehrgang. Verfügbar unter: "
        "https://www.dihk-bildungs-gmbh.de/weiterbildung/ihk-zertifikate/zertifikatslehrgaenge/"
        "data-analyst-ihk [abgerufen September 2026].",
    )
    add_paragraph(
        document,
        "Lemanzyk, S. (2020). Data Analyst (IHK) – Konzept Plus. DIHK-Bildungs-GmbH, Bonn. K181/1/1.",
    )


def build_appendix_a(document: Document) -> None:
    add_heading(document, "Anhang A   Machine Learning Canvas", 1, page_break=True)
    add_paragraph(
        document,
        "Tabelle A1 fasst die zehn Felder des in Kapitel 2 beschriebenen Machine Learning Canvas zusammen.",
    )
    add_table_caption(document, "Tabelle A1: Machine Learning Canvas der Verbrauchsprognose und Anomalieprüfung")
    add_table(
        document,
        ["Nr.", "Feld", "Festlegung im Projekt"],
        [
            ["1", "Mehrwert", "Genauere Folgemonatsplanung und monatliche Prüfung ungewöhnlicher Zählerwerte."],
            ["2", "Datenquellen", "Zählerverbräuche, Verträge, Kundentyp, Kalender, Wetter, Produktionsplan und Wartung."],
            ["3", "Vorhersage", "VLS je Zähler-Monat; Rückrechnung in kWh; große Residuen werden Prüfhinweise."],
            ["4", "Merkmale", "Historie, Saison, Kalender, Wetterprognose, Produktionsplan, Wartung und Kundentyp."],
            ["5", "Lernansatz", "Überwachte Regression; lineare Regression als Referenz, Random Forest als nichtlinearer Vergleich."],
            ["6", "Evaluation", "RMSE in kWh als Hauptmetrik; MAE und R² ergänzend; Vergleich mit zwei Baselines."],
            ["7", "Entscheidung", "Prognose unterstützt Beschaffung; Hinweis führt zur menschlichen Prüfung."],
            ["8", "Auswirkung", "Pilot-KPIs: Fehler, Hinweisvolumen, Prüfzeit und Bestätigungsquote."],
            ["9", "Zeitpunkt", "Prognose vor Monatsbeginn; Residuenprüfung nach Eintreffen des Monatswerts."],
            ["10", "Monitoring", "Monatliche Fehler-, Drift- und Hinweiskontrolle; Retraining nur bei belegter Verschlechterung."],
        ],
        [1.0, 3.7, 11.3],
        font_size=9.5,
        line_spacing=1.0,
    )


def build_appendix_b(document: Document, formulas: dict[str, Path]) -> None:
    add_heading(document, "Anhang B   Datenmodell und Formeln", 1, page_break=True)
    add_heading(document, "B.1   Datenwörterbuch der Modellierungsbasis", 2)
    add_table_caption(document, "Tabelle B1: Spalten der Datei modellierung_basis_bis_3_monate.csv")
    add_table(
        document,
        ["Spalte", "Typ", "Rolle und Herkunft"],
        [
            ["zaehler_id", "String", "Eindeutige Zählerkennung; Originalfeld"],
            ["monat", "Datum", "Beobachtungsmonat; bereinigt"],
            ["jahr", "Integer", "Aus monat abgeleitet"],
            ["split", "Kategorie", "train, test oder ausschluss; Steuerung der zeitlichen Nutzung"],
            ["vollaststunden", "Float [h]", "Interne Trainingszielgröße: verbrauch_kwh / vertragsleistung_kw"],
            ["verbrauch_kwh", "Float [kWh]", "Realisierter Monatsverbrauch; Ausgabe- und Bewertungseinheit"],
            ["vertragsleistung_kw", "Float [kW]", "Originalfeld; Normierung und Rückrechnung, kein gelerntes Merkmal"],
            ["kundentyp", "Kategorie", "Bereinigtes Originalfeld"],
            ["monat_idx", "Integer", "Kalendermonat 1 bis 12"],
            ["arbeitstage", "Integer", "Kalendermerkmal"],
            ["feiertage_im_monat", "Integer", "Kalendermerkmal"],
            ["heizgradtage", "Float", "Wettermerkmal"],
            ["produktionsplan_index", "Float", "Vorausplanbare Produktionsaktivität"],
            ["wartung_aktiv", "Bool", "Geplante Wartung im Monat"],
            ["vormonat_vls", "Float [h]", "VLS des unmittelbar vorherigen verfügbaren Monats"],
            ["letzte_3_monate_vls", "Float [h]", "Mittel aus bis zu drei gültigen Vormonaten"],
            ["vorjahr_vls", "Float [h]", "Vergleichswert aus dem Vorjahresmonat; nicht trainiert"],
            ["anomalie", "Bool", "Bereitgestelltes Analysekennzeichen; kein Trainingslabel"],
            ["unmoeglich", "Bool", "Technisches Plausibilitätskennzeichen"],
            ["ziel_rekonstruiert", "Bool", "Kennzeichnet rekonstruierte Zielwerte"],
        ],
        [5.0, 3.1, 7.9],
        font_size=9.5,
        line_spacing=1.0,
    )

    add_heading(document, "B.2   Modelle und Vorverarbeitung", 2)
    add_table_caption(document, "Tabelle B2: Trainierte Modelle und ausgewählte Konfiguration")
    add_table(
        document,
        ["Modell", "Rolle", "Konfiguration"],
        [
            ["Lineare Regression", "Erklärbare lineare Referenz", "Standardisierung numerischer Merkmale; One-Hot-Kodierung des Kundentyps"],
            ["Random Forest Regressor", "Finales nichtlineares Modell", "300 Bäume; Tiefe 8; min. 5 Fälle je Blatt; max_features 0,7; random_state 42"],
        ],
        [4.2, 5.2, 6.6],
        font_size=10.5,
        line_spacing=1.1,
    )
    add_source_note(document, "Der Vormonatswert und das Mittel aus bis zu drei Vormonaten wurden nur als Baselines berechnet. Sie sind keine trainierten Modelle.")

    add_heading(document, "B.3   Formelsammlung und Rechenbeispiel", 2)
    add_formula(
        document,
        formulas["mae"],
        4,
        "Der MAE gibt die durchschnittliche absolute Abweichung in kWh an. Für die 8.398 auswertbaren "
        "Zähler-Monate des Jahres 2025 beträgt sie 3.724,6 kWh, gerundet 3.725 kWh.",
    )
    add_formula(
        document,
        formulas["r2"],
        5,
        "Im Benchmark 2025 beträgt R² 0,901. Der Wert beschreibt die erklärte Streuung im Vergleich zu einer "
        "Mittelwertprognose. Er bedeutet nicht, dass das Modell in 90,1 % der Fälle richtig liegt. R² hat keine "
        "Einheit und kann bei einem ungeeigneten Modell auch negativ sein.",
    )
    add_formula(
        document,
        formulas["advantage"],
        6,
        "Die beste Baseline erreicht 10.922,441 kWh RMSE, der Random Forest 9.188,397 kWh. Eingesetzt in die "
        "Formel ergibt sich eine Verbesserung von 15,9 %.",
    )
    add_formula(
        document,
        formulas["baselines"],
        7,
        "Beispiel für ZL-00000 im März 2024: Die Vormonats-Baseline verwendet den Februarwert von 101,7 "
        "VLS-Stunden. Für das Mittel aus bis zu drei Vormonaten stehen zu diesem Zeitpunkt nur Januar und "
        "Februar zur Verfügung. Daraus ergeben sich (110,5 + 101,7) ÷ 2 = 106,1 VLS-Stunden. Die Variable k "
        "gibt an, wie viele gültige Vormonate verfügbar sind, höchstens jedoch drei.",
    )
    add_formula(
        document,
        formulas["folds"],
        8,
        "Für den Random Forest ergaben sich 17.537,588 kWh in Fold 1, 8.223,484 kWh in Fold 2 und "
        "14.056,150 kWh in Fold 3. Der Mittelwert beträgt 13.272,407 kWh, die "
        "Populations-Standardabweichung 3.842,640 kWh. Die Whisker in Abbildung D2 zeigen diese Streuung und "
        "kein Konfidenzintervall.",
    )
    add_formula(
        document,
        formulas["percentile"],
        9,
        "Die 1.397 absoluten Kalibrierungsfehler wurden vom kleinsten bis zum größten Wert sortiert. Für das "
        "99. Perzentil ergibt sich in nullbasierter Zählung die Position 1.382,04. Diese liegt zwischen dem "
        "1.383. Wert mit 144,2215 VLS-Stunden und dem 1.384. Wert mit 147,5526 VLS-Stunden. Durch lineare "
        "Interpolation ergibt sich die Schwelle von 144,3547 VLS-Stunden, gerundet 144,4. Insgesamt liegen "
        "14 Kalibrierungsfehler darüber.",
        width_cm=14.7,
    )


def build_appendix_c(document: Document, embedded: dict[str, Path]) -> None:
    add_heading(document, "Anhang C   Datenverständnis und EDA", 1, page_break=True)
    source = "Quelle: eigene Darstellung auf Basis der bereinigten Projektdaten."
    figures = [
        (embedded["eda_verteilung.png"], "Abbildung C1: Verteilung von Monatsverbrauch und Vollaststunden", "Nach der Division durch die Vertragsleistung ist die Verteilung weniger stark von der Größe der Anschlüsse geprägt. Die Zähler lassen sich dadurch besser miteinander vergleichen."),
        (embedded["eda_kundentyp.png"], "Abbildung C2: Verbrauch je Kundentyp, roh und normiert", "Beim Rohverbrauch unterscheiden sich die Kundentypen deutlich. In Vollaststunden liegen ihre Medianwerte wesentlich näher beieinander."),
        (embedded["eda_physik.png"], "Abbildung C3: Vertragsleistung und physikalisch auffällige Monatswerte", "Werte oberhalb der technisch möglichen Monatsenergie sind physikalisch nicht plausibel. Hohe Werte unterhalb dieser Grenze bleiben als mögliche reale Verbrauchsspitzen erhalten."),
        (embedded["eda_korrelation.png"], "Abbildung C4: Korrelationen vor und nach der VLS-Normierung", "Nach der Normierung wird der Einfluss der Anschlussgröße kleiner. Zusammenhänge mit Produktion, Wartung und weiteren Merkmalen sind dadurch deutlicher zu erkennen."),
        (embedded["eda_produktion.png"], "Abbildung C5: Produktionsplan und relativer Verbrauch", "Die Abbildung zeigt einen Zusammenhang zwischen Produktionsplan und relativem Verbrauch. Daraus allein lässt sich jedoch noch keine Ursache-Wirkungs-Beziehung ableiten."),
        (embedded["eda_wartung.png"], "Abbildung C6: Verbrauch in Monaten mit und ohne geplante Wartung", "In Monaten mit geplanter Wartung ist der Verbrauch niedriger. Die Abbildung zeigt einen Zusammenhang, beweist aber nicht, dass die Wartung allein die Ursache ist."),
    ]
    for _index, (path, caption, explanation) in enumerate(figures):
        add_figure(document, path, caption, source, max_height_cm=7.8)
        add_paragraph(document, explanation, space_after=5)


def build_appendix_d(document: Document) -> None:
    add_heading(document, "Anhang D   Modellierung und Evaluation", 1, page_break=True)
    source = "Quelle: eigene Darstellung auf Basis von data/processed/modellierung_basis_bis_3_monate.csv."
    charts = ROOT / ".build" / "presentation" / "charts"
    figures = [
        (charts / "time-split.png", "Abbildung D1: Zeitliche Trennung von Modellwahl, Kalibrierung und Benchmark"),
        (charts / "cv-comparison.png", "Abbildung D2: Mittlerer CV-RMSE und Streuung der drei zeitlichen Folds 2024"),
        (charts / "benchmark-2025.png", "Abbildung D3: Modell- und Baselinevergleich im retrospektiven Benchmark 2025"),
        (charts / "vls-vs-kwh.png", "Abbildung D4: Random Forest mit direkter kWh-Zielgröße gegenüber VLS mit Rückrechnung"),
        (charts / "ranked-errors-q99.png", "Abbildung D5: Sortierte absolute Kalibrierungsfehler und 99. Perzentil"),
        (charts / "threshold-workload.png", "Abbildung D6: Perzentilwahl und monatlicher Prüfaufwand 2025"),
        (ROOT / ".build" / "ihk-actual-vs-predicted-final.png", "Abbildung D7: Ist gegen Prognose für 8.398 Zähler-Monate 2025"),
    ]
    for index, (path, caption) in enumerate(figures):
        if not path.is_file():
            raise FileNotFoundError(path)
        add_figure(document, path, caption, source)
        if index in (1, 3, 5):
            document.add_page_break()

    add_table_caption(document, "Tabelle D1: Gruppierte Permutationswichtigkeit im Benchmark 2025")
    add_table(
        document,
        ["Informationsgruppe", "RMSE-Anstieg nach Mischen", "Interpretation"],
        [
            ["Verbrauchshistorie", "7.256 kWh", "größter Prognosebeitrag"],
            ["Produktionsplan", "798 kWh", "relevante Planinformation"],
            ["Wartung", "591 kWh", "relevante Planinformation"],
            ["Kalender", "276 kWh", "ergänzendes Signal"],
            ["Wetter", "154 kWh", "ergänzendes Signal"],
            ["Jahreszeit", "121 kWh", "kleiner zusätzlicher Beitrag"],
            ["Kundentyp", "ca. 3 kWh", "nahezu kein zusätzlicher Beitrag nach Normierung"],
        ],
        [5.0, 4.5, 6.5],
        font_size=10.5,
        line_spacing=1.1,
    )
    add_source_note(document, "Steigt der RMSE nach dem Mischen einer Gruppe, hat das Modell deren Informationen für die Prognose genutzt. Daraus folgt keine Ursache-Wirkungs-Beziehung. Bei korrelierten Merkmalen kann sich die gemessene Bedeutung auf mehrere Gruppen verteilen.")


def build_appendix_e(document: Document) -> None:
    add_heading(document, "Anhang E   Dashboard-Prototyp", 1, page_break=True)
    add_paragraph(
        document,
        "Das Dashboard ist ein lokaler Prototyp auf Grundlage der historischen Ergebnisse aus 2025. Es besitzt "
        "keine Verbindung zu einem produktiven System und löst keine Maßnahmen automatisch aus.",
    )
    dashboard = ROOT / ".build" / "presentation" / "dashboard" / "anomalie-cockpit-q99.png"
    if not dashboard.is_file():
        raise FileNotFoundError(dashboard)
    add_figure(
        document,
        dashboard,
        "Abbildung E1: Anomalieprüfung mit 99. Perzentil, Prüfaufwand und Ist-Prognose-Diagramm",
        "Quelle: eigener Oberflächenprototyp auf Basis des retrospektiven 2025-Exports.",
        max_height_cm=15.5,
    )
    add_paragraph(
        document,
        "Im vorgesehenen Ablauf wählt der Nutzer zunächst die Schwelle und erhält eine danach sortierte "
        "Prüfliste. Anschließend öffnet er einen Zählerfall und vergleicht Istwert, Prognose und "
        "Schwellenfaktor. Danach werden Datenqualität, Produktionsplanung und Wartung geprüft. Die fachliche "
        "Entscheidung und ihre Begründung können gespeichert und später als Trainingsdaten für ein ergänzendes "
        "Klassifikationsmodell verwendet werden.",
    )


def remove_highlights(document: Document) -> None:
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            r_pr = run._element.get_or_add_rPr()
            node = r_pr.find(qn("w:highlight"))
            if node is not None:
                r_pr.remove(node)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        r_pr = run._element.get_or_add_rPr()
                        node = r_pr.find(qn("w:highlight"))
                        if node is not None:
                            r_pr.remove(node)


def main() -> None:
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)

    formulas = {
        "vls": render_formula(
            "formula_vls.png",
            r"$\mathrm{VLS}_{\mathrm{Ist},i}=\frac{\mathrm{Verbrauch}_{\mathrm{Ist},i}\,[\mathrm{kWh}]}{\mathrm{Vertragsleistung}_{i}\,[\mathrm{kW}]}$" "\n"
            r"$\mathrm{Verbrauch}_{\mathrm{Prognose},i}\,[\mathrm{kWh}]=\mathrm{VLS}_{\mathrm{Prognose},i}\,[\mathrm{h}]\cdot\mathrm{Vertragsleistung}_{i}\,[\mathrm{kW}]$",
            width=13.0,
            height=1.10,
            fontsize=15,
        ),
        "rmse": render_formula(
            "formula_rmse.png",
            r"$\mathrm{RMSE}_{\mathrm{kWh}}=\sqrt{\frac{1}{N_{\mathrm{ZaehlerMonate}}}\sum_{i=1}^{N_{\mathrm{ZaehlerMonate}}}(\mathrm{Verbrauch}_{\mathrm{Ist},i}-\mathrm{Verbrauch}_{\mathrm{Prognose},i})^2}$",
            width=11.8,
            fontsize=15,
        ),
        "alert": render_formula(
            "formula_alert.png",
            r"$\mathrm{Residuum}_{\mathrm{VLS},i}=\mathrm{VLS}_{\mathrm{Ist},i}-\mathrm{VLS}_{\mathrm{Prognose},i}\qquad \mathrm{Grenze}_{99}=144{,}4\,\mathrm{h}$" "\n"
            r"$\mathrm{Schwellenfaktor}_{i}=\frac{|\mathrm{Residuum}_{\mathrm{VLS},i}|}{\mathrm{Grenze}_{99}}\geq1$",
            width=13.2,
            height=1.10,
            fontsize=15,
        ),
        "mae": render_formula(
            "formula_mae.png",
            r"$\mathrm{MAE}_{\mathrm{kWh}}=\frac{1}{N_{\mathrm{ZaehlerMonate}}}\sum_{i=1}^{N_{\mathrm{ZaehlerMonate}}}|\mathrm{Verbrauch}_{\mathrm{Ist},i}-\mathrm{Verbrauch}_{\mathrm{Prognose},i}|$",
            width=11.8,
            fontsize=15,
        ),
        "r2": render_formula(
            "formula_r2.png",
            r"$R^2=1-\frac{\sum_i(\mathrm{Verbrauch}_{\mathrm{Ist},i}-\mathrm{Verbrauch}_{\mathrm{Prognose},i})^2}{\sum_i(\mathrm{Verbrauch}_{\mathrm{Ist},i}-\overline{\mathrm{Verbrauch}}_{\mathrm{Ist}})^2}$",
            width=12.2,
            fontsize=15,
        ),
        "advantage": render_formula(
            "formula_advantage.png",
            r"$\mathrm{Vorteil}_{\mathrm{RMSE}}=\frac{\mathrm{RMSE}_{\mathrm{BisZu3MonatsMittel}}-\mathrm{RMSE}_{\mathrm{RandomForest}}}{\mathrm{RMSE}_{\mathrm{BisZu3MonatsMittel}}}\cdot100\,\%$",
            width=12.5,
            fontsize=13,
        ),
        "baselines": render_formula(
            "formula_baselines.png",
            r"$\widehat{\mathrm{VLS}}_{\mathrm{Vormonat},t}=\mathrm{VLS}_{\mathrm{Ist},t-1}\qquad \widehat{\mathrm{VLS}}_{\mathrm{BisZu3Monaten},t}=\frac{1}{k}\sum_{r=1}^{k}\mathrm{VLS}_{\mathrm{Ist},t-r},\quad k\leq3$",
            width=12.8,
            fontsize=13,
        ),
        "folds": render_formula(
            "formula_folds.png",
            r"$\overline{\mathrm{RMSE}}_{\mathrm{CV}}=\frac{\mathrm{RMSE}_{\mathrm{Fold}\,1}+\mathrm{RMSE}_{\mathrm{Fold}\,2}+\mathrm{RMSE}_{\mathrm{Fold}\,3}}{3}$" "\n"
            r"$\sigma_{\mathrm{Fold}}=\sqrt{\frac{1}{3}\sum_{j=1}^{3}(\mathrm{RMSE}_{\mathrm{Fold}\,j}-\overline{\mathrm{RMSE}}_{\mathrm{CV}})^2}$",
            width=13.2,
            height=1.10,
            fontsize=15,
        ),
        "percentile": render_formula(
            "formula_percentile.png",
            r"$\mathrm{Position}_{99}=(N_{\mathrm{Kalibrierungsfehler}}-1)\cdot0{,}99=(1397-1)\cdot0{,}99=1382{,}04$" "\n"
            r"$\mathrm{Grenze}_{99}=\mathrm{Fehler}_{\mathrm{Rang}\,1383}+0{,}04\cdot(\mathrm{Fehler}_{\mathrm{Rang}\,1384}-\mathrm{Fehler}_{\mathrm{Rang}\,1383})=144{,}3547\,\mathrm{h}$",
            width=13.2,
            height=1.15,
            fontsize=13,
        ),
    }
    embedded = extract_embedded_images()

    document = Document(SOURCE)
    truncate_after_cover(document)
    configure_document(document)
    build_main_report(document, formulas)
    build_literature(document)
    build_appendix_a(document)
    build_appendix_b(document, formulas)
    build_appendix_c(document, embedded)
    build_appendix_d(document)
    build_appendix_e(document)
    remove_highlights(document)

    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
