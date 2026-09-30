"""Endversion der gemeinsamen Präsentation von Gruppe 6 im SWW-Design.

Quelle ist die im Drive zusammengeschnittene Datei ``docs/Copy of Praesentationsvorlage_IHK(3).pptx``
(bleibt unverändert). Kikos und Patricks Folien werden übernommen; geändert werden dort nur Label,
Seitenzahl, Fußzeile und Backup-Verweise in den Notizen. Ianas Folien, Fazit, Dank und die Backup-
Folien im alten Vorlagen-Design werden mit denselben Texten in den Layouts des Designsystems neu
gesetzt; ihre drei Diagramme werden aus ``data/raw/verbrauch_bereinigt.csv`` mit den Plotly-
Bausteinen des Designsystems neu berechnet. Kikos Folien 17 und 18 werden zu einer Folie.

    .venv/Scripts/python.exe scripts/build_final_praesentation.py
"""

from __future__ import annotations

import hashlib
import io
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.enum.text import PP_ALIGN

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
import sww_layouts as L  # noqa: E402
import team_kit as kit  # noqa: E402
from sww_layouts import (GREY200, GREY500, GREY600, GREY900, NAVY100, NAVY200, NAVY700, NAVY800, TEAL300,  # noqa: E402
                         TEAL600, TEAL700, WHITE, X0, badge, card, header, panel, text)

ROOT = kit.ROOT
SRC = ROOT / "docs" / "Copy of Praesentationsvorlage_IHK(3).pptx"
OUT = ROOT / "docs" / "presentation" / "final" / "Gruppe6_Praesentation_final.pptx"
CHART_DIR = OUT.parent / "assets"
DATA = ROOT / "data" / "raw" / "verbrauch_bereinigt.csv"
SOURCE_SHA = hashlib.sha1(SRC.read_bytes()).hexdigest() if SRC.exists() else ""
FOOT = "Verbrauchsprognose & Frühwarnung"
NAMES = {"IANA": "Iana Kraievska", "KIKO": "Kiko Ramon Lukas", "PATRICK": "Patrick Olmo Hederer",
         "TEAM": "Gruppe 6"}
MONO = kit.MONO
MAIN_TOTAL = 27

# Neue Labels für Patricks Folien (Schema KAPITEL n · THEMA · NAME); Kikos folgen dem Schema bereits
PATRICK_LABELS = {
    19: "KAPITEL 5 UND 6 · PATRICK", 20: "KAPITEL 5 · MODELL-PERFORMANCE · PATRICK",
    21: "KAPITEL 5 · TOP-3-TREIBER · PATRICK", 22: "KAPITEL 5 · COCKPIT BESCHAFFUNG · PATRICK",
    23: "KAPITEL 5 · COCKPIT NETZMANAGEMENT · PATRICK", 24: "KAPITEL 5 · GESCHÄFTSNUTZEN · PATRICK",
    25: "KAPITEL 6 · EMPFEHLUNGEN · PATRICK", 26: "KAPITEL 6 · EMPFEHLUNGEN UND AUSBLICK · PATRICK",
    31: "BACKUP · EDA · PATRICK", 32: "BACKUP · GRENZEN · PATRICK", 33: "BACKUP · COCKPIT-PRÜFFALL · PATRICK",
}
KIKO_OLD = [9, 10, 11, 12, 13, 14, 15, 16, 34, 35, 36, 37]
NOTE_FIXES = {14: ("Backup B1", "Backup B5"), 15: ("Backup B3", "Backup B7"), 23: ("Backup B3", "Backup B4")}
BACKUPS = [("B1", "Datenqualität im Detail", "Iana"), ("B2", "EDA: Normierung und Treiber", "Patrick"),
           ("B3", "Was die Ergebnisse belegen", "Patrick"), ("B4", "Cockpit-Prüffall ZL-00337", "Patrick"),
           ("B5", "Hyperparameter", "Kiko"), ("B6", "Lag-Merkmale", "Kiko"),
           ("B7", "Kalibrierung der Schwelle", "Kiko"), ("B8", "ML Canvas im Wortlaut", "Kiko")]


# --------------------------------------------------------------------------- Hilfen für neue Folien


def foot(s, num: str, who: str, dark: bool = False) -> None:
    """Fußzeile wie im Designsystem, Projekttitel mit Namen der sprechenden Person."""
    L.line(s, 0, L.FOOT, 1280, L.FOOT, L.GLASS_LINE if dark else L.GREY100, 1, name="Fußlinie")
    if dark:
        L.box(s, X0, 666, 70, 30, fill=WHITE, radius=4, name="Logo-Grund")
        kit.picture(s, kit.LOGO, X0 + 6, 670, h=22, name="Logo")
        x = X0 + 86
    else:
        kit.picture(s, kit.LOGO, X0, 668, h=26, name="Logo")
        x = X0 + 92
    t = f"{FOOT} · {NAMES[who]}"
    text(s, x, 671, kit.measure(t, 16) + 12, 20, t, 16, L.NAVY300 if dark else GREY600, lh=20, wrap=False)
    if num:
        text(s, 1208 - 80, 671, 80, 20, num, 16, L.NAVY300 if dark else GREY600, align="r", lh=20, wrap=False,
             name="Seitenzahl")


def copy_notes(dst, src) -> None:
    if src.has_notes_slide and src.notes_slide.notes_text_frame.text.strip():
        dst.notes_slide.notes_text_frame.text = src.notes_slide.notes_text_frame.text


def code(t: str):
    return (t, {"font": MONO, "color": TEAL700, "size": 15})


def rings(s, clip):
    kit.ring_arc(s, clip[2], 20, 260, clip, "rgba(108,192,210,.2)", 2)
    kit.ring_arc(s, clip[2] - 10, 30, 170, clip, "rgba(108,192,210,.12)", 2)


# --------------------------------------------------------------------------- Ianas Diagramme


def _eda_data():
    df = pd.read_csv(DATA)
    df["jahr"] = pd.to_datetime(df["monat"]).dt.year
    v = df.groupby(["zaehler_id", "jahr"])["verbrauch_kwh"].agg(["median", "max"]).unstack().dropna()
    v.columns = [f"{m}_{j}" for m, j in v.columns]
    v["kundentyp"] = df.groupby("zaehler_id")["kundentyp"].first()
    v["spitze_2024"] = v["max_2024"] / v["median_2024"]
    v["spitze_2025"] = v["max_2025"] / v["median_2025"]
    return v


def eda_stats() -> dict:
    v = _eda_data()
    return {"zaehler": len(v), "max_2024": float(v["max_2024"].max())}


def eda_charts() -> dict:
    """Ianas Diagramme als PNG für die Folien (1.520 × 940 px)."""
    from energy_analytics.visualization import eda

    CHART_DIR.mkdir(parents=True, exist_ok=True)
    return {key: eda.save_for_slide(fig, CHART_DIR / f"iana_{key}.png", width=760, height=470)
            for key, fig in eda_figures().items()}


def eda_figures() -> dict:
    """Ianas drei Vergleiche 2024/2025 mit den Plotly-Bausteinen des Designsystems (energy_analytics)."""
    import plotly.graph_objects as go
    from energy_analytics.visualization import eda
    from energy_analytics.visualization import theme as T

    eda.setup()
    v = _eda_data()
    out = {}
    for key, col, label in [("niveau", "median", "Üblicher Verbrauch"), ("spitzen", "max", "Höchster Monatsverbrauch")]:
        fig = eda.parity(v[f"{col}_2025"], v[f"{col}_2024"], x_title=f"{label} 2024 (kWh)",
                         y_title=f"{label} 2025 (kWh)")
        fig.update_layout(yaxis_tickformat=",d", margin=dict(l=100, r=30, t=20, b=80))
        out[key] = fig
    # Relative Spitze nach Kundentyp, logarithmische Achsen wie im Original (Ticks 1, 2, 5 …)
    r = v[np.isfinite(v["spitze_2024"]) & np.isfinite(v["spitze_2025"]) & (v["spitze_2024"] >= 1)
          & (v["spitze_2025"] >= 1) & v["kundentyp"].notna()]
    upper = max(max(r["spitze_2024"].quantile(.99), r["spitze_2025"].quantile(.99)) * 1.15, 5)
    if max(r["spitze_2024"].max(), r["spitze_2025"].max()) <= 100:
        upper = min(upper, 100)
    ticks = [t for t in (1, 2, 5, 10, 20, 50, 100) if t <= upper]
    fig = go.Figure()
    fig.add_trace(go.Scatter(x=[1, upper], y=[1, upper], mode="lines", name="Gleiches Niveau", hoverinfo="skip",
                             line=dict(color=T.GREY_300, width=1.5, dash="4,4")))
    for kt in ("Gewerbe", "Industrie", "Kommunal"):
        g = r[r["kundentyp"] == kt]
        fig.add_trace(go.Scatter(x=g["spitze_2024"], y=g["spitze_2025"], mode="markers", name=kt,
                                 marker=dict(size=7, color=T.KUNDENTYP_COLORS[kt], opacity=.75)))
    axis = dict(type="log", range=[0, np.log10(upper)], tickvals=ticks, ticktext=[str(t) for t in ticks],
                showgrid=True, gridcolor=T.GREY_100)
    fig.update_layout(template=T.template(), xaxis=dict(axis, title="Relative Verbrauchsspitze 2024"),
                      yaxis=dict(axis, title="Relative Verbrauchsspitze 2025"), hovermode="closest",
                      legend=dict(orientation="h", yanchor="bottom", y=1.02, x=0),
                      margin=dict(l=100, r=30, t=50, b=80))
    out["relativ"] = fig
    for fig in out.values():  # save_for_slide vergrößert Ticks und Legende, die Achsentitel nicht
        fig.update_layout(xaxis_title_font_size=18, yaxis_title_font_size=18)
    return out


# --------------------------------------------------------------------------- Neue Folien


def s_titel(prs, num, src):
    s = L.slide(prs, NAVY800)
    rings(s, (0, 0, 1280, 720))
    L.box(s, X0, 56, 150, 150, fill=WHITE, radius=10, name="Logo-Grund")
    logo = next(sh for sh in src.slides[0].shapes if sh.shape_type == 13)
    kit.picture(s, io.BytesIO(logo.image.blob), X0 + 11, 67, 128, 128, name="Logo SWW")
    text(s, X0, 250, 900, 20, "Datenanalyse-Projekt · Gruppe 6", 16, TEAL300, True, lh=20, tracking=.1, upper=True,
         wrap=False, name="Label")
    text(s, X0, 280, 1136, 120, "Prognose des monatlichen Energieverbrauchs und Erkennung von Anomalien", 48,
         WHITE, True, lh=54, name="Titel")
    text(s, X0, 412, 1000, 104, "Wie können historische Verbrauchsdaten und weitere Einflussfaktoren genutzt werden, "
         "um den monatlichen Energieverbrauch zuverlässig zu prognostizieren und auffällige Abweichungen frühzeitig "
         "zu erkennen?", 22, NAVY200, lh=32)
    text(s, X0, 560, 1000, 26, "Gruppe 6 · Iana Kraievska · Patrick Olmo Hederer · Kiko Ramon Lukas", 19, WHITE,
         True, lh=26)
    text(s, X0, 592, 1000, 24, "IHK-Zertifizierung Data Analyst (IHK) · 01.10.2026", 17, NAVY200, lh=24)
    copy_notes(s, src.slides[0])
    return s


AGENDA = [("01", "Ausgangssituation und Problem", "Iana"), ("02", "Daten und Datenqualität", "Iana"),
          ("03", "ML Canvas — Überblick", "Kiko"), ("04", "Methodik und Modell", "Kiko"),
          ("05", "Ergebnisse", "Patrick"), ("06", "Empfehlungen und Ausblick", "Patrick")]


def s_agenda(prs, num, src):
    s = L.slide(prs)
    header(s, "Agenda · Gruppe 6", "Agenda")
    for i, (nr, name, who) in enumerate(AGENDA):
        x, y = X0 + (i % 3) * 385, 166 + (i // 3) * 214
        card(s, x, y, 366, 196)
        text(s, x + 26, y + 22, 100, 44, nr, 34, TEAL600, True, MONO, lh=44, wrap=False)
        text(s, x + 26, y + 80, 314, 60, name, 22, GREY900, True, lh=29)
        text(s, x + 26, y + 150, 314, 26, who, 18, GREY600, lh=26)
    foot(s, num, "TEAM")
    copy_notes(s, src.slides[1])
    return s


def s_ausgangssituation(prs, num, src):
    s = L.slide(prs)
    header(s, "Kapitel 1 · Ausgangssituation · Iana", "Ausgangssituation")
    panel(s, X0, 160, 360, 110, "Organisation", "StadtWerke Westhafen GmbH", head_size=18, body_size=20)
    panel(s, X0, 284, 360, 164, "Bisherige Analytics-Prozesse", "Ein automatisiertes Prognosesystem ist nicht "
          "vorhanden.", head_size=18, body_size=19)
    L.eyebrow(s, 452, 160, 300, "Probleme", GREY500)
    L.eyebrow(s, 852, 160, 300, "Folgen", GREY500)
    rows = [("1. Ungenaue Prognosen des Stromverbrauchs", "1. Steigende Beschaffungskosten"),
            ("2. Verspätete Erkennung von Anomalien", "2. Verzögerungen bei der Erkennung technischer Probleme")]
    for i, (problem, folge) in enumerate(rows):
        y = 188 + i * 130
        card(s, 452, y, 356, 116)
        text(s, 476, y, 316, 116, problem, 20, GREY900, True, anchor="m", lh=27)
        kit.icon(s, "arrow-right", TEAL600, 816, y + 46, 26)
        card(s, 852, y, 356, 116, fill=L.GREY50)
        text(s, 876, y, 316, 116, folge, 20, GREY900, anchor="m", lh=27)
    panel(s, X0, 460, 558, 150, "Projektziel", "ML-Modell für Verbrauchsprognose und Anomalieerkennung",
          variant="teal", head_size=18, body_size=21)
    panel(s, 650, 460, 558, 150, "Projektnutzen", "Datenbasierte Lösung der identifizierten Probleme.",
          variant="navy", head_size=18, body_size=21)
    foot(s, num, "IANA")
    copy_notes(s, src.slides[2])
    return s


def s_datenstruktur(prs, num, src):
    s = L.slide(prs)
    header(s, "Kapitel 2 · Datenstruktur · Iana", "Datenstruktur")
    old = next(sh for sh in src.slides[3].shapes if getattr(sh, "has_table", False) and sh.has_table)
    rows = [[c.text for c in r.cells] for r in old.table.rows if r.cells[1].text.strip()]
    kit.table(s, X0, 160, [320, 816], rows, [48] + [54] * (len(rows) - 1), size=19, head_size=17,
              name="Tabelle Datenstruktur")
    text(s, X0, 560, 1136, 28, [[("Datensatz vor der Bereinigung: ", {"bold": True}),
                                 ("16.830 Zeilen · 16 Spalten · 700 Zähler · 01.2024–12.2025 (24 Monate)", {})]],
         19, GREY600, lh=28)
    foot(s, num, "IANA")
    copy_notes(s, src.slides[3])
    return s


def s_datenqualitaet(prs, num, src):
    s = L.slide(prs)
    header(s, "Kapitel 2 · Datenqualität · Iana", "Datenqualität")
    rows = [["Problem", "Umfang", "Maßnahme"],
            [[("Duplikate", {"bold": True})], "30 (0,18 %)", "Entfernt"],
            [[[("Fehlende Werte", {"bold": True})], [code("vorjahr_monat_verbrauch_kwh")],
              [code("vormonat_verbrauch_kwh")]], [["8.400 (50 %)"], ["700 (4 %)"]],
             "Strukturell bedingt → nicht imputiert"],
            [[[("Inkonsistente Datumswerte", {"bold": True})], [code("monat")], [code("kundentyp")]],
             [["11.316 (67 %)"], ["505 (3 %)"]], "Datumsformat und Kategorien vereinheitlicht"],
            [[[("Ausreißer / fehlerhafte Werte", {"bold": True})], [code("verbrauch_kwh")]], "20 (0,12 %)",
             "Geprüft"]]
    kit.table(s, X0, 160, [470, 250, 416], rows, [48, 52, 96, 96, 72], size=19, head_size=17,
              name="Tabelle Datenqualität")
    text(s, X0, 560, 1136, 28, [[("Datensatz nach der Bereinigung: ", {"bold": True}),
                                 ("16.800 Zeilen · 16 Spalten · 700 Zähler · 01.2024–12.2025 (24 Monate)", {})]],
         19, GREY600, lh=28)
    foot(s, num, "IANA")
    copy_notes(s, src.slides[4])
    return s


EDA = {"niveau": ("Kapitel 2 · Verbrauchsniveau · Iana", "Typischer Stromverbrauch der Zähler",
                  "Vergleich des üblichen Verbrauchsniveaus zwischen 2024 und 2025",
                  "Üblicher Verbrauch = Median der Monatswerte eines Zählers."),
       "spitzen": ("Kapitel 2 · Verbrauchsspitzen · Iana", "Höchster monatlicher Stromverbrauch",
                   "Vergleich der stärksten Verbrauchsspitzen zwischen 2024 und 2025",
                   "Höchster Monatsverbrauch eines Zählers im jeweiligen Jahr."),
       "relativ": ("Kapitel 2 · Relative Spitzen · Iana", "Verbrauchsspitze im Verhältnis zum üblichen Verbrauch",
                   "Vergleich der relativen Verbrauchsspitzen nach Kundentyp",
                   "Relative Spitze = höchster Monatsverbrauch ÷ üblicher Verbrauch; Achsen logarithmisch.")}


def s_eda(key):
    def make(prs, num, src, charts):
        s = L.slide(prs)
        label, title, subtitle, definition = EDA[key]
        header(s, label, title)
        L.image_card(s, charts[key], 84, 172, 700, name=f"Diagramm {key}")
        panel(s, 832, 160, 376, 200, "", subtitle, variant="teal", body_size=20)
        panel(s, 832, 376, 376, 250, "So liest man das Diagramm",
              [f"Ein Punkt je Zähler. {definition}", "Gestrichelt: gleiches Niveau in beiden Jahren."],
              head_size=18, body_size=17)
        foot(s, num, "IANA")
        return s
    return make


def s_kiko_ausblick(prs, num, src):
    s = L.slide(prs, NAVY800)
    header(s, "Kapitel 4 · Ausblick · Kiko", "Vom Prüfhinweis zum Klassifikationsmodell", dark=True)
    text(s, X0, 152, 520, 30, "Heute im Cockpit", 22, WHITE, True, lh=29)
    steps = [("Fallakte", "Warum wurde der Fall markiert? Faktor, Prognose und Ist."),
             ("Bewertung", "Messwert gültig? Abweichung erklärt? Ursache."),
             ("Status", "Gespeichert in der Prüfwarteschlange.")]
    for i, (head, body) in enumerate(steps):
        y = 196 + i * 86
        kit.marker(s, X0 + 16, y + 16, i + 1, fill=L.TEAL600, d=30, size=16)
        text(s, X0 + 46, y + 2, 480, 28, head, 20, WHITE, True, lh=27)
        text(s, X0 + 46, y + 32, 480, 44, body, 17, NAVY200, lh=24)
    L.box(s, X0, 460, 540, 60, fill=L.GLASS, line=L.GLASS_LINE, radius=10)
    text(s, X0 + 18, 460, 510, 60, "Messwert gültig = Ja + Abweichung erklärt = Nein ⇒ bestätigte Anomalie", 16,
         WHITE, font=MONO, anchor="m", lh=22)
    text(s, 652, 152, 556, 30, "Nächster Schritt", 22, WHITE, True, lh=29)
    nxt = [("01", "Labels sammeln", "Jede Bewertung ist ein Ja/Nein-Label mit Ursache."),
           ("02", "Klassifikation trainieren", "Ein zweites Modell schätzt je Hinweis die Wahrscheinlichkeit in %."),
           ("03", "Workflow automatisieren", "Hinweise automatisiert vorsortieren; der Mensch entscheidet.")]
    for i, (nr, head, body) in enumerate(nxt):
        y = 196 + i * 112
        L.box(s, 652, y, 556, 104, fill=L.GLASS, line=L.GLASS_LINE, radius=10)
        text(s, 672, y + 14, 60, 34, nr, 26, TEAL300, True, MONO, lh=34, wrap=False)
        text(s, 736, y + 12, 452, 28, head, 20, WHITE, True, lh=27)
        text(s, 736, y + 42, 452, 56, body, 17, NAVY200, lh=24)
    L.box(s, X0, 538, 1136, 52, fill=L.GLASS, line=L.GLASS_LINE, radius=10)
    bw = badge(s, X0 + 16, 547, "Vorbehalt", "dark")
    text(s, X0 + 30 + bw, 538, 1136 - 50 - bw, 52, "Das Modell ersetzt keine Abrechnungsentscheidung — es flaggt nur "
         "Untersuchungswürdiges.", 17, WHITE, anchor="m", lh=23)
    text(s, X0, 600, 1136, 30, "Weiter mit Patrick: Ergebnisse 2025 und wie echte Zählerwerte das Retraining der "
         "Prognose verbessern.", 18, TEAL300, True, lh=26)
    foot(s, num, "KIKO", dark=True)
    kit.set_notes(s, [
        ("STICHWORTE", ["• Fallakte → Bewertung → Status", "• gültig = Ja + erklärt = Nein ⇒ bestätigt",
                        "• Labels → Klassifikation mit % → automatisiert vorsortieren", "• Übergabe an Patrick"]),
        ("SPRECHTEXT", [
            "Was passiert mit einem Prüfhinweis? Im Cockpit öffnet die Prüferin die Fallakte: Dort steht, warum "
            "der Fall markiert wurde, mit Faktor, Prognose und Istwert. Dann bewertet sie: Ist der Messwert gültig? "
            "Ist die Abweichung erklärt? Was war die Ursache? Die Bewertung wird gespeichert, und der Fall bekommt "
            "seinen Status in der Prüfwarteschlange.",
            "Die Regel dahinter: Ist der Messwert gültig und die Abweichung nicht erklärt, ist es eine bestätigte "
            "Anomalie. [Pause]",
            "Genau darauf baut mein nächster Schritt auf. Jede Bewertung ist ein Label. Sobald genug bestätigte "
            "Fälle vorliegen, trainiere ich ein zweites Modell, ein Klassifikationsmodell. Es schätzt für jeden "
            "Hinweis, wie wahrscheinlich er sich bestätigt, zum Beispiel 80 Prozent. Damit lassen sich die "
            "Hinweise automatisiert vorsortieren; entscheiden bleibt beim Menschen.",
            "Wie das Modell 2025 abgeschnitten hat und wie echte Zählerwerte das Retraining der Prognose "
            "verbessern, zeigt jetzt Patrick."]),
        ("SO VERSTEHST DU ES", [
            "• Links ist der heutige Prozess im Cockpit, rechts dein Plan. Die Bewertungen von links sind das "
            "Trainingsmaterial für rechts.",
            "• Regression sagt eine Zahl vorher (kWh). Klassifikation sagt eine Klasse vorher (bestätigt ja/nein) "
            "und kann eine Wahrscheinlichkeit angeben.",
            "• Belastbar wird die Prozentzahl erst, wenn man sie auf neuen Fällen geprüft und kalibriert hat."]),
        ("FALLS GEFRAGT WIRD", [
            ("Wie viele Labels braucht man?", "Das legt der Pilot fest; wichtig sind bestätigte und verworfene "
                                              "Fälle, auch aus einer Stichprobe unauffälliger Monate."),
            ("Was ist Precision@K?", "Der Anteil bestätigter Fälle unter den K obersten Hinweisen der Prüfliste.")]),
    ])
    return s


def s_fazit(prs, num, src):
    s = L.slide(prs)
    header(s, "Kapitel 6 · Fazit · Team", "Persönliches Fazit")
    quotes = [("Iana Kraievska", "„Datenbereinigung ist keine rein technische Routine: Erst der Vergleich mit der "
                                 "technisch möglichen Monatsenergie trennt Messfehler von plausiblen "
                                 "Verbrauchswerten.“"),
              ("Kiko Ramon Lukas", "„Bei Zeitreihendaten ist die Validierung wichtiger als die Wahl des komplexesten "
                                   "Modells.“"),
              ("Patrick Olmo Hederer", "„Eine auffällige Korrelation beweist noch keinen Treiber: Normierung "
                                       "verschiebt die Korrelationen und hilft dem Modell sein Training zu "
                                       "spezifizieren.“")]
    for i, (name, quote) in enumerate(quotes):
        x = X0 + i * 385
        card(s, x, 160, 366, 262)
        text(s, x + 24, 180, 318, 26, name, 18, TEAL700, True, lh=26)
        text(s, x + 24, 214, 318, 196, quote, 18, GREY900, lh=26)
    L.box(s, X0, 440, 1136, 180, fill=NAVY700, radius=10, name="Antwort")
    L.eyebrow(s, X0 + 26, 460, 700, "Antwort auf die Projektfrage", TEAL300)
    text(s, X0 + 26, 494, 1084, 116, [
        "• Historie und Plandaten liefern vor jedem Monat eine Prognose je Zähler, 15,9 % genauer als die einfache "
        "Fortschreibung.",
        "• Große Abweichungen landen nach Monatsende auf einer Prüfliste – nicht erst im Quartal."], 20, WHITE,
        lh=28, gap=8, indent=22)
    foot(s, num, "TEAM")
    copy_notes(s, src.slides[26])
    return s


def s_dank(prs, num, src):
    s = L.slide(prs, NAVY800)
    rings(s, (0, 0, 1280, 720))
    text(s, X0, 250, 900, 76, "Vielen Dank.", 64, WHITE, True, lh=76, name="Titel")
    text(s, X0, 340, 900, 34, "Wir freuen uns auf das Fachgespräch.", 26, NAVY200, lh=34)
    for i, name in enumerate(["Iana Kraievska", "Kiko Ramon Lukas", "Patrick Olmo Hederer"]):
        text(s, X0 + i * 330, 460, 310, 28, name, 20, WHITE, True, lh=28)
    foot(s, num, "TEAM", dark=True)
    copy_notes(s, src.slides[27])
    return s


def s_backup(prs, num, src):
    s = L.slide(prs, NAVY800)
    rings(s, (0, 0, 1280, 720))
    text(s, X0, 80, 900, 70, "Backup", 60, WHITE, True, lh=70, name="Titel")
    text(s, X0, 156, 900, 30, "Detail-Folien für mögliche Rückfragen.", 22, NAVY200, lh=30)
    for i, (nr, title, who) in enumerate(BACKUPS):
        x, y = X0 + (i // 4) * 578, 230 + (i % 4) * 90
        L.box(s, x, y, 558, 74, fill=L.GLASS, line=L.GLASS_LINE, radius=10)
        text(s, x + 20, y, 60, 74, nr, 22, TEAL300, True, MONO, anchor="m", lh=28, wrap=False)
        text(s, x + 84, y, 360, 74, title, 18, WHITE, True, anchor="m", lh=24)
        text(s, x + 440, y, 100, 74, who, 17, NAVY200, align="r", anchor="m", lh=22, wrap=False)
    copy_notes(s, src.slides[28])
    return s


def s_dq_detail(prs, num, src):
    s = L.slide(prs)
    header(s, "Backup · Datenqualität · Iana", "Datenqualität im Detail")
    old = next(sh for sh in src.slides[29].shapes if getattr(sh, "has_table", False) and sh.has_table)
    rows = []
    for i, r in enumerate(old.table.rows):
        cells = [c.text.strip() for c in r.cells]
        if i and " – " in cells[0] and cells[0].split(" – ")[0].islower():
            col, rest = cells[0].split(" – ", 1)
            cells[0] = [code(col), (f" – {rest}", {})]
        rows.append(cells)
    kit.table(s, X0, 150, [520, 150, 150, 316], rows, [40] + [39] * (len(rows) - 1), size=16, head_size=15,
              right_cols=(1, 2), name="Tabelle Datenqualität Detail")
    foot(s, num, "IANA")
    copy_notes(s, src.slides[29])
    return s


# --------------------------------------------------------------------------- Übernommene Folien anpassen


def _set_text(shape, value: str, align=None) -> None:
    tf = shape.text_frame
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    p = tf.paragraphs[0]
    runs = p.runs
    if runs:
        runs[0].text = value
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p.add_run().text = value
    if align is not None:
        p.alignment = align


def _size(shape) -> tuple[float, bool, str]:
    r = next((r for p in shape.text_frame.paragraphs for r in p.runs), None)
    size = r.font.size.pt / 0.75 if r is not None and r.font.size else 16.0
    return size, bool(r is not None and r.font.bold), (r.font.name if r is not None and r.font.name else kit.SANS)


def _texts(slide):
    return [sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]


def adapt_old(slide, old_nr: int, num: str) -> None:
    """Label, Seitenzahl und Fußzeile einer übernommenen Folie; Inhalt bleibt unverändert."""
    who = "PATRICK" if old_nr in PATRICK_LABELS else "KIKO"
    for sh in _texts(slide):
        t = sh.text_frame.text.strip()
        top, left = sh.top / 9525, sh.left / 9525
        if top < 70 and t == t.upper() and "·" in t and old_nr in PATRICK_LABELS:
            _set_text(sh, PATRICK_LABELS[old_nr])
            size, bold, font = _size(sh)
            sh.width = kit.E(min(1136, kit.measure(PATRICK_LABELS[old_nr], size, bold, font) * 1.2 + 12))
        elif top > 640 and left > 1000 and re.fullmatch(r"\d+( / \d+)?|B\d+", t):
            right = sh.left + sh.width
            _set_text(sh, num, PP_ALIGN.RIGHT)
            for r in sh.text_frame.paragraphs[0].runs:  # Patricks Zahlen standen in Geist Mono
                r.font.name = kit.SANS
                if r._r.rPr is not None and r._r.rPr.get("spc"):
                    del r._r.rPr.attrib["spc"]
            sh.width = kit.E(110)
            sh.left = right - sh.width
        elif top > 640 and t.startswith(FOOT):
            value = f"{FOOT} · {NAMES[who]}"
            _set_text(sh, value)
            size, bold, font = _size(sh)
            sh.width = kit.E(kit.measure(value, size, bold, font) * 1.15 + 16)
    if old_nr == 9:  # Kikos Kapitelfolie hat keine Fußzeile; Projekttitel mit Namen rechts unten ergänzen
        text(slide, 661, 671, 460, 20, f"{FOOT} · {NAMES['KIKO']}", 16, GREY600, lh=20, wrap=False)
    if old_nr in NOTE_FIXES and slide.has_notes_slide:
        old, new = NOTE_FIXES[old_nr]
        for p in slide.notes_slide.notes_text_frame.paragraphs:
            for r in p.runs:
                if old in r.text:
                    r.text = r.text.replace(old, new)


# --------------------------------------------------------------------------- Zusammenbau

NEW_SLIDES: list[int] = []


def plan():
    """Endgültige Reihenfolge: ("neu", Funktion) oder ("alt", Foliennummer der Quelle)."""
    return ([("neu", s_titel), ("neu", s_agenda), ("neu", s_ausgangssituation), ("neu", s_datenstruktur),
             ("neu", s_datenqualitaet), ("eda", "niveau"), ("eda", "spitzen"), ("eda", "relativ")]
            + [("alt", n) for n in range(9, 17)] + [("neu", s_kiko_ausblick)]
            + [("alt", n) for n in range(19, 27)] + [("neu", s_fazit), ("neu", s_dank)]
            + [("neu", s_backup), ("neu", s_dq_detail)] + [("alt", n) for n in (31, 32, 33, 34, 35, 36, 37)])


def numbers(entries) -> list[str]:
    out, b = [], 0
    for i, (kind, what) in enumerate(entries):
        if i == 0 or what is s_backup:
            out.append("")
        elif i < MAIN_TOTAL:
            out.append(f"{i + 1} / {MAIN_TOTAL}")
        else:
            b += 1
            out.append(f"B{b}")
    return out


def build(path: Path = OUT) -> Path:
    src = Presentation(SRC)
    prs = Presentation(SRC)
    charts = eda_charts()
    old_ids = list(prs.slides._sldIdLst)
    entries = plan()
    missing = kit.K["_MISSING"]
    missing.clear()
    order = []
    for (kind, what), num in zip(entries, numbers(entries)):
        if kind == "alt":
            sld = old_ids[what - 1]
            adapt_old(prs.slides[what - 1], what, num)
            order.append(sld)
            continue
        maker = s_eda(what) if kind == "eda" else what
        s = maker(prs, num, src, charts) if kind == "eda" else maker(prs, num, src)
        kit.flatten_alpha(s)
        order.append(prs.slides._sldIdLst[-1])
    if missing:  # fehlende Icon-PNGs (Pfeil) einmal rendern und neu bauen
        kit.K["render_icons"](missing)
        missing.clear()
        return build(path)
    lst = prs.slides._sldIdLst
    keep = set(id(x) for x in order)
    for sld in list(lst):
        if id(sld) not in keep:  # ersetzte Folien der Quelle entfernen
            lst.remove(sld)
            prs.part.drop_rel(sld.rId)
    for sld in order:  # neue Reihenfolge
        lst.remove(sld)
        lst.append(sld)
    NEW_SLIDES[:] = [i for i, (kind, _) in enumerate(entries) if kind != "alt"]
    prs.core_properties.title = "Prognose des monatlichen Energieverbrauchs und Erkennung von Anomalien"
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path


if __name__ == "__main__":
    out = build()
    print(f"{out.relative_to(ROOT)} ({out.stat().st_size // 1024} KB)")
