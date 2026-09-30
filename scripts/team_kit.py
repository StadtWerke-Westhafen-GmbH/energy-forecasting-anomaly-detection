"""Bausteine für die gemeinsame Präsentation von Gruppe 6 (Google-Slides-tauglich).

Lädt Kikos PPTX-Generator (``build_kiko_pptx.py``) per ``runpy`` und nutzt dessen Grundbausteine
(Textfelder, Scheiben, Karten, Zählwerk). Für Google Slides werden einzelne Funktionen in dessen
Globals ersetzt; Kikos Einzel-Deck bleibt davon unberührt, weil jedes ``runpy.run_path`` einen
eigenen Namensraum anlegt.

Google-Modus:
* Versalien werden als echte Großbuchstaben geschrieben (Google Slides kennt ``cap="all"`` nicht).
* Einzeilige Textfelder bekommen 10 % Breitenreserve, weil Google die Schrift minimal anders setzt.
* PowerPoint-Diagramme entstehen auf einer Hilfsfolie, die am Ende entfällt; auf der Folie steht
  stattdessen das PNG des Diagramms aus Kikos HTML-Deck (``export_team_charts.mjs``).
* Transparenzen werden am Ende in deckende Mischfarben mit dem darunterliegenden Grund umgerechnet.
"""

from __future__ import annotations

import json
import runpy
import subprocess
from pathlib import Path

from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
# run_path liefert nur eine Kopie der Globals; ersetzt wird im echten Namensraum der Funktionen
K = runpy.run_path(str(ROOT / "scripts" / "build_kiko_pptx.py"))["build"].__globals__
BUILD = K["BUILD"]
TEAM_DIR = ROOT / "docs" / "presentation" / "team"
CHART_DIR = TEAM_DIR / "assets" / "charts"

E, rgb, mix, text, disc, line = K["E"], K["rgb"], K["mix"], K["text"], K["disc"], K["line"]
W, H, FOOT_Y = K["W"], K["H"], K["FOOT_Y"]
INK, MUTED, SOFT, LINE, LINE2 = K["INK"], K["MUTED"], K["SOFT"], K["LINE"], K["LINE2"]
NAVY, NAVY9, NAVY7 = K["NAVY"], K["NAVY9"], K["NAVY7"]
TEAL, TEAL6, TEAL7, TEAL3 = K["TEAL"], K["TEAL6"], K["TEAL7"], K["TEAL3"]
MIST, PALE, WHITE, SANS, MONO = K["MIST"], K["PALE"], K["WHITE"], K["SANS"], K["MONO"]

# Kapitel der Agenda: Nummer, Name, Icon, sprechende Person
CHAPTERS = [
    ("01", "Ausgangssituation", "building-2", "Iana"),
    ("02", "Daten", "database", "Iana"),
    ("03", "ML Canvas", "layout-grid", "Kiko"),
    ("04", "Methodik und Modell", "trees", "Kiko"),
    ("05", "Ergebnisse", "chart-column", "Patrick"),
    ("06", "Empfehlungen", "compass", "Patrick"),
]

STATE: dict = {"scratch": None, "native": []}


# --------------------------------------------------------------------------- Google-Modus


def _caps(paras, upper: bool):
    """Versalien als echte Großbuchstaben statt ``cap="all"``."""
    if isinstance(paras, str):
        return paras.upper() if upper else paras
    out = []
    for para in paras:
        if isinstance(para, str):
            out.append(para.upper() if upper else para)
        else:
            out.append([(t.upper() if o.get("upper", upper) else t, {k: v for k, v in o.items() if k != "upper"})
                        for t, o in para])
    return out


def _google_mode(k: dict) -> None:
    orig_text, orig_place, orig_header = k["text"], k["place_chart"], k["header"]

    def g_text(s, x, y, w, h, paras, size=16, color=INK, bold=False, font=SANS, align="l", anchor="t",
               lh=None, gap=0.0, name=None, tracking=0.0, wrap=True, upper=False, indent=18):
        paras = _caps(paras, upper)
        tracking = max(tracking, 0.0)  # Google Slides kennt keine Laufweite; Titel müssen ohne sie passen
        line_h = lh or 1.3 * size
        if name == "Titel" and wrap and isinstance(paras, str) and h < 1.5 * line_h:
            width = k["measure"](paras, size, bold, font)
            if width > 0.96 * w:  # einzeiliger Titel: Schrift verkleinern statt umbrechen
                factor = 0.96 * w / width
                size, lh = size * factor, line_h * factor
        if not wrap:  # Reserve je nach Ausrichtung nach rechts, links oder beidseitig
            extra = max(4.0, 0.1 * w)
            x -= {"l": 0.0, "c": extra / 2, "r": extra}[align]
            w += extra
        return orig_text(s, x, y, w, h, paras, size, color, bold, font, align, anchor, lh, gap, name, tracking,
                         wrap, False, indent)

    def g_place_chart(s, chart_type, data, rect, pad=(0, 0, 0, 0), name=None):
        STATE["native"].append(rect)
        return orig_place(STATE["scratch"], chart_type, data, rect, pad, name)

    def g_header(s, root, dark=False):
        orig_header(s, root, dark)  # Kikos Schrittleiste entfällt, die Kapitelleiste kommt danach

    k["text"], k["place_chart"], k["header"] = g_text, g_place_chart, g_header
    k["tracker"] = lambda s, el, dark=False: None


_google_mode(K)
text = K["text"]


# --------------------------------------------------------------------------- Kapitelleiste


def chapter_bar(s, chapter: str, dark: bool = False) -> None:
    """Sechs Kapitel als Icon-Scheiben (je Person ein Paar), aktives gefüllt, Label mit Namen."""
    xs = [1000, 1030, 1076, 1106, 1152, 1182]
    active = next(i for i, c in enumerate(CHAPTERS) if c[0] == chapter)
    for i, ((nr, _, icon, _), x) in enumerate(zip(CHAPTERS, xs)):
        on = i == active
        if dark:
            tone = "light" if on else "glass"
        else:
            tone = "navy" if on else ("ring" if i < active else "off")
        disc(s, x, 42, icon, tone, 26, name=f"Kapitel {nr}{' aktiv' if on else ''}")
    for x in (1062, 1138):
        line(s, x, 55, x + 8, 55, "#3D6A8C" if dark else "#B4BFCB", 2, name="Kapitel-Lücke")
    nr, name, _, who = CHAPTERS[active]
    text(s, 688, 45, 300, 18, f"{nr} · {name} · {who}", 14, TEAL3 if dark else TEAL7, True, align="r", lh=18,
         wrap=False, name="Kapitel-Label")


# --------------------------------------------------------------------------- Diagramme als Bild


def load_charts() -> list[dict]:
    path = CHART_DIR / "charts.json"
    if not path.exists():
        subprocess.run(["node", str(ROOT / "scripts" / "export_team_charts.mjs")], check=True, cwd=ROOT)
    return json.loads(path.read_text(encoding="utf-8"))


def _box_px(el):
    off, ext = el.find(".//" + qn("a:off")), el.find(".//" + qn("a:ext"))
    if off is None or ext is None:
        return None
    x, y = int(off.get("x")) / 9525, int(off.get("y")) / 9525
    return x, y, x + int(ext.get("cx")) / 9525, y + int(ext.get("cy")) / 9525


def remove_inside(s, rect, tol: float = 6.0) -> int:
    """Alle Formen, die vollständig im Rechteck liegen, entfernen (Kartenrahmen außen bleiben)."""
    x, y, w, h = rect
    removed = 0
    for sh in list(s.shapes):
        b = _box_px(sh._element)
        if b and b[0] >= x - tol and b[1] >= y - tol and b[2] <= x + w + tol and b[3] <= y + h + tol:
            sh._element.getparent().remove(sh._element)
            removed += 1
    return removed


def replace_charts(s, slide_idx: int, spec, charts: list[dict]) -> None:
    """Diagrammbereiche mit nativem PowerPoint-Diagramm durch das PNG aus dem HTML-Deck ersetzen."""
    native, STATE["native"] = STATE["native"], []
    if not native:
        return
    svgs = list(K["parse"](spec).iter("svg"))
    for c in (c for c in charts if c["slide"] == slide_idx):
        x, y, w, h = c["x"], c["y"], c["w"], c["h"]
        if not any(x <= rx + rw / 2 <= x + w and y <= ry + rh / 2 <= y + h for rx, ry, rw, rh in native):
            continue
        vb = [float(v) for v in svgs[c["svg"]].get("viewbox").split()]
        assert abs(vb[2] / vb[3] - w / h) < 0.02 * w / h, f"SVG-Zuordnung HTML/PPTX stimmt nicht: {c['file']}"
        remove_inside(s, (x, y, w, h))
        K["picture"](s, str(CHART_DIR / c["file"]), x, y, w, h, name=f"Diagramm {c['file']}")


def drop_slide(prs, slide) -> None:
    for sld_id in list(prs.slides._sldIdLst):
        if prs.part.related_part(sld_id.rId) is slide.part:
            prs.slides._sldIdLst.remove(sld_id)
            prs.part.drop_rel(sld_id.rId)
            return


# --------------------------------------------------------------------------- Transparenzen auflösen


def _hex(el) -> str:
    return "#" + el.get("val").upper()


def flatten_alpha(slide) -> None:
    """Transparente Füllungen und Linien in Mischfarben mit dem darunterliegenden Grund umrechnen.

    Der Grund ist die oberste deckende Fläche unter dem Formmittelpunkt, sonst der Folienhintergrund.
    Gruppenkinder liegen in Folienkoordinaten (python-pptx legt Gruppen ohne Transformation an).
    """
    bg_clr = slide._element.find(".//" + qn("p:bg") + "//" + qn("a:srgbClr"))
    bg = _hex(bg_clr) if bg_clr is not None else WHITE
    painted: list[tuple[tuple, str]] = []
    tree = slide.shapes._spTree
    for el in tree.iter(qn("p:sp"), qn("p:cxnSp")):
        sppr = el.find(qn("p:spPr"))
        b = _box_px(sppr) if sppr is not None else None
        if b is None:
            continue
        cx, cy = (b[0] + b[2]) / 2, (b[1] + b[3]) / 2
        under = next((c for r, c in reversed(painted) if r[0] <= cx <= r[2] and r[1] <= cy <= r[3]), bg)
        fill = sppr.find(qn("a:solidFill"))
        for clr in sppr.iter(qn("a:srgbClr")):
            alpha = clr.find(qn("a:alpha"))
            if alpha is not None:
                clr.set("val", mix(_hex(clr), under, int(alpha.get("val")) / 100000).lstrip("#"))
                clr.remove(alpha)
        if fill is not None and fill.find(qn("a:srgbClr")) is not None:
            painted.append((b, _hex(fill.find(qn("a:srgbClr")))))


def theme_fonts(prs) -> None:
    """Theme-Schriften auf IBM Plex Sans (Notizen, Tabellen ohne eigene Schrift)."""
    import re

    for part in prs.part.package.iter_parts():
        if not str(part.partname).startswith("/ppt/theme/"):
            continue
        if hasattr(part, "_element"):  # Notizen-Master-Theme wird als XML-Teil angelegt
            for latin in part._element.iter(qn("a:latin")):
                latin.set("typeface", "IBM Plex Sans")
        else:
            xml = part.blob.decode("utf-8")
            part._blob = re.sub(r'<a:latin typeface="[^"]*"', '<a:latin typeface="IBM Plex Sans"', xml).encode()


# --------------------------------------------------------------------------- Rahmen für Team-Folien

import html as _html  # noqa: E402

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402

box, card, dcard, oval, poly, picture = K["box"], K["card"], K["dcard"], K["oval"], K["poly"], K["picture"]
measure, label, pt, ring_arc, chevron = K["measure"], K["label"], K["pt"], K["ring_arc"], K["chevron"]
zw = K["zw"]
CANVAS, ML_FIELDS = BUILD["CANVAS"], BUILD["ML_FIELDS"]
FOOTER_TITLE = _html.unescape(BUILD["FOOTER_TITLE"])
LOGO = K["LOGO"]
GLASS, GLASS_LINE = K["GLASS"], K["GLASS_LINE"]
TINT = "#E3F3F7"      # helles Teal (Merksatz, hervorgehobene Karten)
AMBER, RED, GREEN = "#A86505", "#B3261E", "#2F7A33"


def blank(prs, bg: str = WHITE):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(bg)
    return s


def icon(s, name: str, color: str, x: float, y: float, size: float):
    pic = s.shapes.add_picture(K["icon_png"](name, color), E(x), E(y), E(size), E(size))
    pic.name = f"Icon {name}"
    return pic


def header(s, question: str, qicon: str, title: str, tone: str = "teal", dark: bool = False,
           chapter: str | None = None) -> None:
    """Leitfrage mit Icon-Scheibe, Titel als Antwort, Kapitelleiste (wie Kikos Kopfzeile)."""
    qtone = {"teal": "light", "amber": "lamber", "navy": "light"}[tone] if dark else tone
    disc(s, 72, 38, qicon, qtone, 32, name="Leitfrage-Scheibe")
    text(s, 116, 40.5, 560, 27, question, 20, TEAL3 if dark else TEAL6, lh=27, wrap=False, name="Leitfrage")
    text(s, 72, 80, 1136, 41, title, 36, WHITE if dark else INK, True, lh=41, tracking=-0.02, name="Titel")
    if chapter:
        chapter_bar(s, chapter, dark)


def footer(s, refs=(), speaker: str = "", num: str = "", dark: bool = False) -> None:
    """Fußzeile wie bei Kiko: Logo, Projekttitel, Canvas-Bezüge als Mini-Scheiben, Person, Seitenzahl."""
    border = GLASS_LINE if dark else LINE2
    line(s, 0, FOOT_Y, W, FOOT_Y, border, 1, name="Fußlinie")
    if dark:
        box(s, 72, 672, 64, 26, fill=WHITE, radius=4, name="Logo-Grund")
        picture(s, LOGO, 78, 676, h=18, name="Logo")
        x = 152
    else:
        picture(s, LOGO, 72, 672, h=26, name="Logo")
        x = 163.2
    text(s, x, 675, measure(FOOTER_TITLE, 16) + 12, 20, FOOTER_TITLE, 16, "#7FB4D8" if dark else "#6B7887", lh=20,
         wrap=False)
    x += measure(FOOTER_TITLE, 16) + 16
    if refs:
        x += 14
        line(s, x, 674, x, 696, border, 1)
        x += 17
        for r in refs:
            ml = r in ML_FIELDS
            disc(s, x, 674, CANVAS[r][1], ("light" if dark else "teal") if ml else ("glass" if dark else "ring"), 22)
            x += 28
            text(s, x, 676, 30, 18, r, 14, TEAL3 if dark else TEAL6, True, MONO, lh=18, wrap=False)
            x += measure(r, 14, True, MONO) + 6
            text(s, x, 676, measure(CANVAS[r][0], 14) + 10, 18, CANVAS[r][0], 14, PALE if dark else MUTED, lh=18,
                 wrap=False)
            x += measure(CANVAS[r][0], 14) + 12
    num_w = max(64, measure(num, 16, font=MONO))
    if speaker:
        text(s, 1208 - num_w - 16 - 120, 675, 120, 20, speaker, 16, PALE if dark else MUTED, True, align="r", lh=20,
             wrap=False, name="Sprecher")
    if num:
        text(s, 1208 - num_w, 674.5, num_w, 21, num, 16, "#7FB4D8" if dark else "#6B7887", font=MONO, align="r",
             lh=21, wrap=False, name="Seitenzahl")


def set_notes(slide, sections) -> None:
    """Notizen in Abschnitten: [(KOPF, [Zeile | (fetter Anfang, Rest)])]."""
    tf = slide.notes_slide.notes_text_frame
    tf.text = ""
    first = True
    for n, (head, items) in enumerate(sections):
        rows = [[(head, True)]]
        rows += [[(it, False)] if isinstance(it, str) else [(it[0] + " ", True), (it[1], False)] for it in items]
        if n < len(sections) - 1:
            rows.append([("", False)])
        for runs in rows:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            for chunk, bold in runs:
                r = p.add_run()
                r.text = chunk
                r.font.bold = bold


def draft_notes(slide, source: str, speech, keywords=(), questions=(), todo=()) -> None:
    """Entwurfsnotizen für Iana und Patrick: Herkunft, Stichworte, Sprechtext-Vorschlag, offene Punkte."""
    sections = [("ENTWURF", [f"Entwurf aus Bericht ({source}) – bitte prüfen und in eigenen Worten sprechen."]
                 + [f"• {t}" for t in todo])]
    if keywords:
        sections.append(("STICHWORTE", [f"• {k}" for k in keywords]))
    sections.append(("SPRECHTEXT", list(speech)))
    if questions:
        sections.append(("FALLS GEFRAGT WIRD", list(questions)))
    set_notes(slide, sections)


def icard(s, x, y, w, h, icon_name: str, tone: str, head: str, body, psize: float = 15, fill: str = WHITE,
          border: str = LINE, dark: bool = False, d: float = 34) -> None:
    """Karte mit Icon-Scheibe, fettem Kopf und Text (Kikos icard)."""
    if dark:
        dcard(s, x, y, w, h)
    else:
        card(s, x, y, w, h, fill=fill, border=border)
    disc(s, x + 17, y + 15, icon_name, tone, d)
    tx = x + 17 + d + 14
    tw = x + w - 17 - tx
    text(s, tx, y + 16, tw, 22, head, 16, WHITE if dark else NAVY7, True, lh=21)
    text(s, tx, y + 40, tw, h - 46, body, psize, PALE if dark else MUTED, lh=psize * 1.4)


def merk(s, x, y, w, h, msg: str, dark: bool = False) -> None:
    """Merksatz-Band mit Glühbirnen-Scheibe."""
    if dark:
        dcard(s, x, y, w, h)
    else:
        box(s, x, y, w, h, fill=TINT, radius=10)
    disc(s, x + 12, y + (h - 30) / 2, "lightbulb", "lamber" if dark else "amber", 30)
    text(s, x + 56, y, w - 76, h, msg, 18, WHITE if dark else TEAL7, anchor="m", lh=24.3, name="Merksatz")


def marker(s, cx: float, cy: float, n, fill: str = NAVY7, d: float = 30) -> None:
    """Nummerierter Kreis für Screenshot-Hinweise und Schrittfolgen."""
    grp = s.shapes.add_group_shape()
    grp.name = f"Hinweis {n}"
    oval(grp, cx - d / 2, cy - d / 2, d, fill=fill, line=WHITE, lw=2)
    text(grp, cx - d / 2 - 4, cy - d / 2, d + 8, d, str(n), 15, WHITE, True, MONO, "c", "m", lh=d, wrap=False)


def table(s, x, y, widths, rows, heights, *, size: float = 15, head_size: float = 14.5, name: str = "Tabelle",
          right_cols=(), tint_rows=()) -> None:
    """Native Tabelle im Stil von Kikos B2: Kopf Navy, feine Linien.

    Zelle: Text, [(Text, Optionen)] als ein Absatz oder [[…], […]] als mehrere Absätze.
    """
    shape = s.shapes.add_table(len(rows), len(widths), E(x), E(y), E(sum(widths)), E(sum(heights)))
    shape.name = name
    tbl = shape.table
    tbl.first_row = False
    tbl.horz_banding = False
    for j, wv in enumerate(widths):
        tbl.columns[j].width = E(wv)
    for i, (row, hv) in enumerate(zip(rows, heights)):
        tbl.rows[i].height = E(hv)
        for j, value in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = ""
            multi = isinstance(value, list) and value and isinstance(value[0], list)
            paras = value if multi else [value]
            fs = head_size if i == 0 else size
            for k, para in enumerate(paras):
                p = cell.text_frame.paragraphs[0] if k == 0 else cell.text_frame.add_paragraph()
                p.alignment = PP_ALIGN.RIGHT if (j in right_cols and i) else PP_ALIGN.LEFT
                p.line_spacing = pt(fs * 1.3)
                runs = [(para, {})] if isinstance(para, str) else [(t, {}) if isinstance(t, str) else t for t in para]
                for chunk, opts in runs:
                    r = p.add_run()
                    r.text = chunk
                    r.font.name = opts.get("font", SANS)
                    r.font.size = pt(opts.get("size", fs))
                    r.font.bold = opts.get("bold", i == 0)
                    r.font.italic = False
                    r.font.color.rgb = rgb(WHITE if i == 0 else opts.get("color", INK))
            K["_cell_borders"](cell, None if i == 0 else LINE2)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(NAVY9 if i == 0 else (TINT if i in tint_rows else WHITE))
            cell.margin_left = cell.margin_right = E(12)
            cell.margin_top = cell.margin_bottom = E(4)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            edges = [t for t, on in (("a:lnL", j == 0), ("a:lnR", j == len(widths) - 1), ("a:lnT", i == 0),
                                     ("a:lnB", i == len(rows) - 1)) if on]
            for tag in edges:  # Außenrand als Zellrand: wächst mit, wenn Zeilen eingefügt werden
                _cell_edge(cell, tag, LINE)


def _cell_edge(cell, tag: str, color: str) -> None:
    from pptx.oxml.xmlchemy import OxmlElement

    ln = cell._tc.get_or_add_tcPr().find(qn(tag))
    for child in list(ln):
        ln.remove(child)
    ln.set("w", str(int(0.75 * 12700)))
    fill = OxmlElement("a:solidFill")
    clr = OxmlElement("a:srgbClr")
    clr.set("val", color.lstrip("#"))
    fill.append(clr)
    ln.append(fill)


def code(t: str) -> tuple[str, dict]:
    """Spaltenname in Geist Mono, wie im Bericht abgesetzt."""
    return (t, {"font": MONO, "color": TEAL7, "size": 14})
