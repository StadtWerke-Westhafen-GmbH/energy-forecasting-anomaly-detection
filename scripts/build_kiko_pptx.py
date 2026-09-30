"""Kikos Präsentationsteil als bearbeitbare PowerPoint-Datei.

Baut dieselben 17 Folien (11 Hauptfolien, 6 Backups) wie ``build_kiko_praesentation.py``
(HTML, maßgeblich) nativ nach: Texte als Textfelder, Diagramme als PowerPoint-Diagramme mit
hinterlegten Daten, Lucide-Icons als PNG in nativen Kreisen (Farbe bleibt editierbar),
Zählwerk-Ziffern als einzelne Zellen, Screenshots als Bilder; Stichworte, Sprechtext,
Lesehilfe und Rückfragen stehen in den Notizen.

Eine Quelle für beide Formate: Texte, Icons, Scheiben-Töne und Canvas-Bezüge liest das Skript
aus dem HTML der Slide-Objekte (``build_slides``), Diagrammdaten aus denselben Fakten
(Notebook 13, Bericht). Beschriftungen und Legenden der Diagramme kommen aus den Inline-SVGs.

Koordinaten sind CSS-Pixel der 1280 × 720-Folie (1 px = 9525 EMU); Schriftgrößen werden mit
0,75 in Punkt umgerechnet, Zeilenhöhen exakt wie im CSS gesetzt. Schriften: IBM Plex Sans und
Geist Mono (TTF unter brand/design-system/dist/fonts). Fehlende Icon-PNGs rendert
``scripts/render_icons.mjs`` (Node + Playwright) nach docs/presentation/kiko/assets/icons/.
"""

from __future__ import annotations

import base64
import copy
import html
import io
import json
import math
import os
import re
import runpy
import subprocess
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

import lxml.html
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parents[1]
BUILD = runpy.run_path(str(ROOT / "scripts" / "build_kiko_praesentation.py"))
OUT = ROOT / "docs" / "presentation" / "kiko" / "Kiko_ML_Canvas_Methodik_Prueffall.pptx"
ASSETS = BUILD["ASSET_DIR"]
ICON_PNG_DIR = ASSETS / "icons"
FONT_DIR = BUILD["FONT_DIR"]
C, TONES, BERICHT = BUILD["C"], BUILD["TONES"], BUILD["BERICHT"]
MONTHS, STEPS = BUILD["MONTHS"], BUILD["STEPS"]
FOLD_COLORS, FOLD_LABELS = BUILD["FOLD_COLORS"], BUILD["FOLD_LABELS"]
de, signed_de = BUILD["de"], BUILD["signed_de"]

SANS, MONO = "IBM Plex Sans", "Geist Mono"
W, H = 1280, 720
M = 72  # Seitenrand
FOOT_Y = 653

INK, MUTED, SOFT = "#141A21", "#4E5A68", "#657383"
LINE, LINE2 = "#D2DAE2", "#E4E9EF"
NAVY, NAVY9, NAVY7 = "#063659", "#04263F", "#084878"
TEAL, TEAL6, TEAL7, TEAL3 = "#0080A0", "#00718E", "#005E77", "#6CC0D2"
MIST, PALE, WHITE = "#F1F4F7", "#BBD7EA", "#FFFFFF"
GLASS = "rgba(255,255,255,.07)"      # Karten auf dunklem Grund
GLASS_LINE = "rgba(255,255,255,.16)"


# --------------------------------------------------------------------------- Grundbausteine


def E(v: float) -> Emu:
    return Emu(int(round(v * 9525)))


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#"))


def pt(px: float) -> Pt:
    return Pt(px * 0.75)


def _rgba(value: str) -> tuple[str, float | None]:
    """'#RRGGBB' oder 'rgba(r,g,b,a)' -> ('#RRGGBB', Deckkraft oder None)."""
    value = value.strip()
    if value.startswith("rgba"):
        r, g, b, a = (float(v) for v in re.findall(r"[\d.]+", value))
        return f"#{int(r):02X}{int(g):02X}{int(b):02X}", a
    if re.fullmatch(r"#[0-9a-fA-F]{3}", value):  # Kurzform #fff
        value = "#" + "".join(ch * 2 for ch in value[1:])
    return value.upper(), None


def mix(fg: str, bg: str, alpha: float) -> str:
    """Deckende Mischfarbe (für Textfarben mit CSS-opacity)."""
    f, b = rgb(fg), rgb(bg)
    return "#" + "".join(f"{round(alpha * f[i] + (1 - alpha) * b[i]):02X}" for i in range(3))


def _no_shadow(shape):
    """Keine Theme-Effekte: leere Effektliste und kein Stilverweis (effectRef zieht sonst Schatten)."""
    shape.shadow.inherit = False
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def _set_alpha(solid_fill, alpha: float):
    clr = solid_fill[0]
    for old in clr.findall(qn("a:alpha")):
        clr.remove(old)
    el = OxmlElement("a:alpha")
    el.set("val", str(int(alpha * 100000)))
    clr.append(el)


def _fill(fill_format, color):
    if not color or color == "none":
        fill_format.background()
        return
    hex_, alpha = _rgba(color)
    fill_format.solid()
    fill_format.fore_color.rgb = rgb(hex_)
    if alpha is not None:
        _set_alpha(fill_format._xPr.find(qn("a:solidFill")), alpha)


def _stroke(line_format, color, width=1.0, dash=None):
    if not color or color == "none":
        line_format.fill.background()
        return
    hex_, alpha = _rgba(color)
    line_format.color.rgb = rgb(hex_)
    line_format.width = pt(width)
    if dash:
        line_format.dash_style = dash
    if alpha is not None:
        _set_alpha(line_format._get_or_add_ln().find(qn("a:solidFill")), alpha)


def box(s, x, y, w, h, fill=None, line=None, radius=0, lw=1.0, dash=None, name=None, shape=None):
    """Rechteck, abgerundetes Rechteck oder andere Grundform; Farben als Hex oder rgba()."""
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    sh = s.shapes.add_shape(kind, E(x), E(y), E(w), E(h))
    if radius and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, radius / min(w, h))
    _fill(sh.fill, fill)
    _stroke(sh.line, line, lw, dash)
    _no_shadow(sh)
    if name:
        sh.name = name
    return sh


def oval(s, x, y, d, fill=None, line=None, lw=1.0, name=None):
    return box(s, x, y, d, d, fill, line, lw=lw, name=name, shape=MSO_SHAPE.OVAL)


def line(s, x1, y1, x2, y2, color=LINE, width=1.0, dash=None, name=None):
    conn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    _stroke(conn.line, color, width, dash)
    if name:
        conn.name = name
    return conn


def poly(s, pts, stroke=None, width=1.0, fill=None, closed=False, dash=None, miter=False, name=None):
    """Freiform aus Punkten (Pfade, Rauten, Toleranzbänder, Chevrons)."""
    fb = s.shapes.build_freeform(E(pts[0][0]), E(pts[0][1]), scale=1.0)
    fb.add_line_segments([(E(x), E(y)) for x, y in pts[1:]], close=closed)
    sh = fb.convert_to_shape()
    _fill(sh.fill, fill)
    _stroke(sh.line, stroke, width, dash)
    if stroke and miter:
        sh.line._get_or_add_ln().append(OxmlElement("a:miter"))
    _no_shadow(sh)
    if name:
        sh.name = name
    return sh


def text(s, x, y, w, h, paras, size=16, color=INK, bold=False, font=SANS, align="l", anchor="t",
         lh=None, gap=0.0, name=None, tracking=0.0, wrap=True, upper=False, indent=18):
    """paras: str | [str | [(text, opts)]]; opts: size, color, bold, font, tracking, upper.

    Zeilenhöhe ``lh`` in px wird exakt gesetzt (wie line-height im CSS); ``gap`` ist der
    Abstand vor jedem weiteren Absatz. Absätze mit "• " bekommen einen hängenden Einzug.
    """
    tb = s.shapes.add_textbox(E(x), E(y), E(w), E(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        runs = [(para, {})] if isinstance(para, str) else list(para)
        p.line_spacing = pt(lh or 1.3 * max([size] + [o.get("size", size) for _, o in runs]))
        if i and gap:
            p.space_before = pt(gap)
        if runs and runs[0][0].startswith("• "):  # hängender Einzug statt Textpunkt
            runs[0] = (runs[0][0][2:], runs[0][1])
            ppr = p._p.get_or_add_pPr()
            ppr.set("marL", str(E(indent)))
            ppr.set("indent", str(-E(indent)))
            bu = OxmlElement("a:buChar")
            bu.set("char", "•")
            ppr.append(bu)
        for chunk, opts in runs:
            run = p.add_run()
            run.text = chunk
            f = run.font
            f.name = opts.get("font", font)
            f.size = pt(opts.get("size", size))
            f.bold = opts.get("bold", bold)
            f.italic = False
            f.color.rgb = rgb(opts.get("color", color))
            rpr = run._r.get_or_add_rPr()
            rpr.set("lang", "de-DE")
            if opts.get("upper", upper):  # Versalien wie text-transform im CSS; der Text bleibt editierbar
                rpr.set("cap", "all")
            spc = opts.get("tracking", tracking)
            if spc:  # Laufweite in em, wie letter-spacing im CSS
                rpr.set("spc", str(int(round(spc * opts.get("size", size) * 0.75 * 100))))
    return tb


def picture(s, data, x, y, w=None, h=None, name=None):
    stream = io.BytesIO(data) if isinstance(data, bytes) else data
    pic = s.shapes.add_picture(stream, E(x), E(y), E(w) if w else None, E(h) if h else None)
    if name:
        pic.name = name
    return pic


def card(s, x, y, w, h, fill=WHITE, border=LINE, radius=10, lw=1.0, name=None):
    return box(s, x, y, w, h, fill=fill, line=border, radius=radius, lw=lw, name=name)


def dcard(s, x, y, w, h, radius=10, fill=GLASS, border=GLASS_LINE):
    """Karte auf dunklem Grund: Weiß mit 7 % Füllung und 16 % Rand."""
    return box(s, x, y, w, h, fill=fill, line=border, radius=radius)


# --------------------------------------------------------------------------- Textmaße


_FONT_FILES = {
    (SANS, False): "ibm-plex-sans-400.ttf", (SANS, True): "ibm-plex-sans-700.ttf",
    (MONO, False): "geist-mono-400.ttf", (MONO, True): "geist-mono-600.ttf",
}


@lru_cache(None)
def _pil_font(font: str, bold: bool):
    return ImageFont.truetype(str(FONT_DIR / _FONT_FILES[(font, bool(bold))]), 100)


@lru_cache(None)
def _cmap(font: str, bold: bool) -> frozenset:
    from fontTools.ttLib import TTFont

    return frozenset(TTFont(FONT_DIR / _FONT_FILES[(font, bool(bold))]).getBestCmap())


@lru_cache(None)
def _fallback_font(bold: bool):
    """Pfeile und Relationszeichen fehlen in den Latin-Subsets; Ersatz wie im HTML (DejaVu)."""
    import matplotlib

    name = "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"
    return ImageFont.truetype(str(Path(matplotlib.__file__).parent / "mpl-data" / "fonts" / "ttf" / name), 100)


def measure(t: str, size: float, bold=False, font=SANS, tracking=0.0) -> float:
    """Laufweite in px (Markenschrift, ohne Kerning), wie sie PowerPoint setzt."""
    cmap, width = _cmap(font, bold), 0.0
    for ch in t:
        f = _pil_font(font, bold) if ord(ch) in cmap else _fallback_font(bold)
        width += f.getlength(ch)
    return width * size / 100 + tracking * size * len(t)


def wrap_lines(t: str, width: float, size: float, bold=False, font=SANS) -> list[str]:
    lines, cur = [], ""
    for word in t.split():
        trial = f"{cur} {word}".strip()
        if cur and measure(trial, size, bold, font) > width:
            lines.append(cur)
            cur = word
        else:
            cur = trial
    return lines + ([cur] if cur else [])


# --------------------------------------------------------------------------- HTML der Folien lesen


def parse(spec):
    return lxml.html.fragment_fromstring(spec.html.strip())


def _xpath(sel: str) -> str:
    out, sep = ".", "//"
    for tok in sel.split():
        if tok == ">":
            sep = "/"
            continue
        tag, *classes = tok.split(".")
        cond = "".join(f"[contains(concat(' ', normalize-space(@class), ' '), ' {c} ')]" for c in classes)
        out += f"{sep}{tag or '*'}{cond}"
        sep = "//"
    return out


def qa(el, sel: str) -> list:
    return el.xpath(_xpath(sel))


def q(el, sel: str):
    found = qa(el, sel)
    if not found:
        raise LookupError(f"'{sel}' fehlt im HTML der Folie")
    return found[0]


def classes(el) -> list[str]:
    return (el.get("class") or "").split()


def chart_svg(el):
    """Das Diagramm-SVG (role="img") eines Containers, nicht die Icons davor."""
    return el.xpath(".//svg[@role='img']")[0]


def kid(el, tag: str):
    """Erstes direktes Kind ohne Klasse (z. B. der Textteil neben einer Scheibe)."""
    return next(c for c in el if c.tag == tag and not classes(c))


def plain(el) -> str:
    return re.sub(r"\s+", " ", el.text_content()).strip()


def _match(node, sel: str) -> bool:
    tag, *cls = sel.split(".")
    return (not tag or node.tag == tag) and all(c in classes(node) for c in cls)


def _squash(runs):
    out, space = [], True  # Leerraum wie im Browser zusammenfassen
    for t, o in runs:
        t = re.sub(r"\s+", " ", t)
        if space and t.startswith(" "):
            t = t[1:]
        if not t:
            continue
        out.append([t, o])
        space = t.endswith(" ")
    if out:
        out[-1][0] = out[-1][0].rstrip()
    return [(t, o) for t, o in out if t]


def rich(el, rules=(), base=None) -> list[list[tuple[str, dict]]]:
    """Element -> Absätze aus Läufen; ``rules`` ordnet Auszeichnungen (b, .mono …) Formaten zu."""
    paras = [[]]

    def visit(node, opts):
        if not isinstance(node.tag, str) or node.tag in ("svg", "title"):
            return
        if node.tag == "br":
            paras.append([])
            return
        o = dict(opts)
        for sel, extra in rules:
            if _match(node, sel):
                o.update(extra)
        if node.text:
            paras[-1].append((node.text, o))
        for child in node:
            visit(child, o)
            if child.tail:
                paras[-1].append((child.tail, o))

    visit(el, dict(base or {}))
    return [p for p in (_squash(p) for p in paras) if p]


def _style_px(el, prop: str) -> float:
    return float(re.search(rf"(?:^|;){prop}:([\d.]+)px", el.get("style") or "").group(1))


def heading(spec) -> tuple[str, str]:
    """Leitfrage und Titel aus dem HTML-Foliensatz (eine Quelle für beide Formate)."""
    root = parse(spec)
    question, title = qa(root, "div.lq p"), qa(root, "h2")
    return (plain(question[0]) if question else ""), (plain(title[0]) if title else "")


# --------------------------------------------------------------------------- Icons und Scheiben


def _sig(svg) -> tuple:
    return tuple((c.tag, tuple(sorted(c.attrib.items()))) for c in svg.iterdescendants() if isinstance(c.tag, str))


@lru_cache(None)
def _icon_index() -> dict:
    """Lucide-Geometrie -> Iconname, damit Icons aus dem Inline-SVG erkannt werden."""
    source = (ROOT / "scripts" / "build_kiko_praesentation.py").read_text(encoding="utf-8")
    index = {}
    for f in sorted(BUILD["ICON_DIR"].glob("*.svg")):
        sig = _sig(lxml.html.fragment_fromstring(BUILD["icon"](f.stem)))
        if sig not in index or f'"{f.stem}"' in source:  # Aliasnamen: der im HTML-Generator gewinnt
            index[sig] = f.stem
    return index


def icon_of(svg) -> tuple[str, str, float]:
    """Inline-SVG eines Icons -> (Name, Farbe, Größe in px)."""
    return _icon_index()[_sig(svg)], svg.get("stroke").upper(), float(svg.get("width"))


@dataclass
class Disc:
    name: str
    tone: str
    d: float


def disc_of(span) -> Disc:
    """span.disc aus dem HTML -> Icon, Ton (TONES) und Durchmesser."""
    style = span.get("style")
    bg = re.search(r"background:([^;]+);", style).group(1).strip()
    ring = re.search(r"box-shadow:inset 0 0 0 [\d.]+px ([^;]+);", style)
    ring = ring.group(1).strip() if ring else None
    name, fg, _ = icon_of(span.find("svg"))
    tone = next(t for t, (b, f, r) in TONES.items() if b == bg and f.upper() == fg and r == ring)
    return Disc(name, tone, _style_px(span, "width"))


_MISSING: set[tuple[str, str]] = set()


@lru_cache(None)
def _blank_png() -> bytes:
    buf = io.BytesIO()
    Image.new("RGBA", (8, 8), (0, 0, 0, 0)).save(buf, format="PNG")
    return buf.getvalue()


def icon_png(name: str, color: str):
    """Pfad zum Icon-PNG; fehlende werden gesammelt und nach dem ersten Durchlauf gerendert."""
    color = color.upper()
    path = ICON_PNG_DIR / f"{name}__{color.lstrip('#')}.png"
    if path.exists():
        return str(path)
    _MISSING.add((name, color))
    return io.BytesIO(_blank_png())


def render_icons(items) -> None:
    ICON_PNG_DIR.mkdir(parents=True, exist_ok=True)
    spec = json.dumps([{"name": n, "color": c} for n, c in sorted(items)])
    subprocess.run(["node", str(ROOT / "scripts" / "render_icons.mjs"), str(ICON_PNG_DIR), spec, "192"],
                   check=True, cwd=ROOT)


def disc(s, x, y, spec: Disc | str, tone: str = "navy", d: float = 32, name: str | None = None):
    """Icon-Scheibe: native Kreisform (editierbare Farbe) mit Icon-PNG (52 % des Durchmessers)."""
    if isinstance(spec, str):
        spec = Disc(spec, tone, d)
    bg, fg, ring = TONES[spec.tone]
    d = spec.d
    grp = s.shapes.add_group_shape()
    grp.name = name or f"Scheibe {spec.name}"
    if ring:  # Ring liegt im CSS innen (inset); PowerPoint zeichnet mittig auf der Kontur
        oval(grp, x + 0.75, y + 0.75, d - 1.5, fill=bg, line=ring, lw=1.5)
    else:
        oval(grp, x, y, d, fill=bg)
    grp.shapes.add_picture(icon_png(spec.name, fg), E(x + d * .24), E(y + d * .24), E(d * .52), E(d * .52))
    return grp


def glyph(s, x, y, svg, size=None):
    """Einzelnes Icon ohne Scheibe (Pfeile, Zitatzeichen …) aus einem Inline-SVG."""
    name, color, px = icon_of(svg)
    size = size or px
    pic = s.shapes.add_picture(icon_png(name, color), E(x), E(y), E(size), E(size))
    pic.name = f"Icon {name}"
    return pic


# --------------------------------------------------------------------------- Zählwerk

ZW_TONES = {  # Zelle, Ring, Ziffer, Trennzeichen
    "navy": (NAVY9, None, WHITE, TEAL6),
    "glass": ("rgba(255,255,255,.10)", "rgba(255,255,255,.22)", WHITE, TEAL3),
    "red": ("#B3261E", None, WHITE, "#B3261E"),
}


def zw(s, x, y, value: str, tone: str, size: float, right: float | None = None) -> float:
    """Zählwerk: jede Ziffer in eigener abgerundeter Zelle, Trennzeichen ohne Zelle."""
    cw, ch, gap, sep = .74 * size, 1.16 * size, .06 * size, .34 * size
    widths = [cw if c.isdigit() else sep for c in value]
    total = sum(widths) + gap * (len(value) - 1)
    if right is not None:
        x = right - total
    cell, ring, ink, sep_ink = ZW_TONES[tone]
    grp = s.shapes.add_group_shape()
    grp.name = f"Zählwerk {value}"
    cx = x
    for c, w in zip(value, widths):
        if c.isdigit():
            box(grp, cx, y, w, ch, fill=cell, line=ring, radius=.12 * size)
            text(grp, cx - 6, y, w + 12, ch, c, size, ink, True, MONO, "c", "m", lh=size, wrap=False)
        else:
            text(grp, cx - .4 * size, y, w + .8 * size, ch, c, size, sep_ink, True, MONO, "c", "m", lh=size,
                 wrap=False)
        cx += w + gap
    return x + total


def zw_from(s, el, x, y, right=None) -> float:
    tone = next(c[3:] for c in classes(el) if c.startswith("zw-"))
    size = float(re.search(r"font-size:([\d.]+)px", el.get("style")).group(1))
    return zw(s, x, y, plain(el), tone, size, right)


# --------------------------------------------------------------------------- Inline-SVG übertragen


class SvgMap:
    """SVG-Koordinaten (viewBox) -> Folienkoordinaten."""

    def __init__(self, svg, ox: float, oy: float, width: float):
        vb = [float(v) for v in svg.get("viewbox").split()]
        self.ox, self.oy, self.k = ox, oy, width / vb[2]

    def x(self, v):
        return self.ox + v * self.k

    def y(self, v):
        return self.oy + v * self.k

    def rect(self, x, y, w, h):
        return self.x(x), self.y(y), w * self.k, h * self.k


def _path_points(d: str):
    """Pfad (M/L/H/V/Z, auch relativ; Bögen als Sehne) -> Punkte und ob geschlossen."""
    toks = re.findall(r"[-+]?(?:\d+\.?\d*|\.\d+)(?:e[-+]?\d+)?|[A-Za-z]", d)
    pts, closed, x, y, cmd, i = [], False, 0.0, 0.0, None, 0

    def nums(k):
        nonlocal i
        vals = [float(t) for t in toks[i:i + k]]
        i += k
        return vals

    while i < len(toks):
        if toks[i].isalpha():
            cmd = toks[i]
            i += 1
            if cmd in "Zz":
                closed = True
                continue
        if cmd in "Mm":
            dx, dy = nums(2)
            x, y = (dx, dy) if cmd == "M" else (x + dx, y + dy)
            cmd = "L" if cmd == "M" else "l"
        elif cmd in "Ll":
            dx, dy = nums(2)
            x, y = (dx, dy) if cmd == "L" else (x + dx, y + dy)
        elif cmd in "Hh":
            (dx,) = nums(1)
            x = dx if cmd == "H" else x + dx
        elif cmd in "Vv":
            (dy,) = nums(1)
            y = dy if cmd == "V" else y + dy
        elif cmd in "Aa":
            *_, ex, ey = nums(7)
            x, y = (ex, ey) if cmd == "A" else (x + ex, y + ey)
        else:
            raise ValueError(f"Pfadbefehl {cmd} nicht unterstützt")
        pts.append((x, y))
    return pts, closed


def _dash(spec):
    if not spec:
        return None
    first = float(spec.split()[0])
    return MSO_LINE_DASH_STYLE.ROUND_DOT if first <= 2 else (
        MSO_LINE_DASH_STYLE.SQUARE_DOT if first <= 3 else MSO_LINE_DASH_STYLE.DASH)


def is_tick(el) -> bool:
    """Achsenbeschriftung im SVG (grau, Zahl oder Monat): übernimmt das native Diagramm."""
    t = (el.text or "").strip()
    return el.tag == "text" and (el.get("fill") or "").upper() == C["grey500"] and (
        re.fullmatch(r"[−-]?[\d.,]+( %)?", t) is not None or t in MONTHS)


def svg_draw(s, svg, m: SvgMap, plot=None, skip=lambda el: False):
    """Beschriftungen, Legenden und einfache Formen eines Inline-SVG als native Formen.

    Formen, deren Mitte in ``plot`` (SVG-Koordinaten x0, y0, x1, y1) liegt, zeichnet das
    native Diagramm; Texte werden immer übertragen (außer ``skip``).
    """
    def in_plot(pts):
        if not plot:
            return False
        xs, ys = [p[0] for p in pts], [p[1] for p in pts]
        cx, cy = (min(xs) + max(xs)) / 2, (min(ys) + max(ys)) / 2
        return plot[0] <= cx <= plot[2] and plot[1] <= cy <= plot[3]

    def num(el, attr, default=0.0):
        v = el.get(attr)
        return float(v) if v not in (None, "") else default

    def walk(node):
        for el in node:
            if not isinstance(el.tag, str) or el.tag == "title" or skip(el):
                continue
            tag, k = el.tag, m.k
            stroke, sw, dash = el.get("stroke"), num(el, "stroke-width", 1) * k, _dash(el.get("stroke-dasharray"))
            fill = el.get("fill")
            if tag == "svg":  # eingebettetes Icon (z. B. Schloss)
                glyph(s, m.x(num(el, "x")), m.y(num(el, "y")), el, num(el, "width") * k)
            elif tag == "g":
                walk(el)
            elif tag == "text":
                content = (el.text or "").strip()
                if not content:
                    continue
                size = num(el, "font-size", 16) * k
                mono = "Geist Mono" in (el.get("font-family") or "")
                bold = num(el, "font-weight", 400) >= 600
                font = MONO if mono else SANS
                tw = measure(content, size, bold, font) + 10
                top = m.y(num(el, "y")) - (1.005 if mono else 1.025) * size
                x = m.x(num(el, "x"))
                anchor = el.get("text-anchor") or "start"
                bx, align = {"start": (x, "l"), "middle": (x - tw / 2, "c"), "end": (x - tw, "r")}[anchor]
                text(s, bx, top, tw, 1.3 * size, content, size, _rgba(fill or INK)[0], bold, font, align,
                     lh=1.3 * size, wrap=False)
            elif tag == "rect":
                x, y, w, h = num(el, "x"), num(el, "y"), num(el, "width"), num(el, "height")
                if in_plot([(x, y), (x + w, y + h)]):
                    continue
                box(s, *m.rect(x, y, w, h), fill=fill, line=stroke, radius=num(el, "rx") * k, lw=sw, dash=dash)
            elif tag == "line":
                x1, y1, x2, y2 = (num(el, a) for a in ("x1", "y1", "x2", "y2"))
                if in_plot([(x1, y1), (x2, y2)]):
                    continue
                line(s, m.x(x1), m.y(y1), m.x(x2), m.y(y2), stroke, sw, dash)
            elif tag == "circle":
                cx, cy, r = num(el, "cx"), num(el, "cy"), num(el, "r")
                if in_plot([(cx - r, cy - r), (cx + r, cy + r)]):
                    continue
                oval(s, m.x(cx - r), m.y(cy - r), 2 * r * k, fill=fill, line=stroke, lw=sw)
            elif tag in ("path", "polyline", "polygon"):
                if tag == "path":
                    pts, closed = _path_points(el.get("d"))
                else:
                    raw = [float(v) for v in re.findall(r"[-\d.]+", el.get("points"))]
                    pts, closed = list(zip(raw[::2], raw[1::2])), tag == "polygon"
                if len(pts) < 2 or in_plot(pts):
                    continue
                poly(s, [(m.x(px), m.y(py)) for px, py in pts], stroke=stroke, width=sw,
                     fill=fill if (closed and fill != "none") else None, closed=closed, dash=dash)

    walk(svg)


# --------------------------------------------------------------------------- Folienrahmen


def page(prs, spec, bg=WHITE):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(bg)
    notes(s, spec)
    return s, parse(spec)


def tracker(s, el, dark=False):
    """Roter Faden als sechs Icon-Scheiben (aktiv gefüllt) mit Label der aktiven Schritte."""
    xs = [998, 1030, 1062, 1118, 1150, 1182]
    for i, (step, x) in enumerate(zip(qa(el, "span.tstep"), xs), start=1):
        d = disc_of(q(step, "span.disc"))
        on = d.tone in ("navy", "teal", "light")
        disc(s, x, 41, d, name=f"Tracker {i} {STEPS[i - 1]}{' aktiv' if on else ''}")
    for x in (1094, 1106):
        line(s, x, 54, x + 6, 54, "rgba(255,255,255,.3)" if dark else "#B4BFCB", 2, name="Tracker-Lücke")
    text(s, 686, 45, 300, 18, plain(q(el, "span.tlabel")), 14, TEAL3 if dark else TEAL7, True, align="r",
         lh=18, wrap=False, name="Tracker-Label")


def cdots(s, el):
    """Canvas-Folie: „10 von 10 Feldern“ und zehn Nummernpunkte (ML-Felder gefüllt)."""
    dots = qa(el, "span.cdot")
    x0 = 1208 - 22 * len(dots) - 4 * (len(dots) - 1)
    text(s, x0 - 310, 45, 300, 18, plain(el.find("span")), 14, SOFT, align="r", lh=18, wrap=False)
    for i, dot in enumerate(dots):
        ml = "ml" in classes(dot)
        grp = s.shapes.add_group_shape()
        grp.name = f"Canvas-Feld {plain(dot)}"
        x = x0 + 26 * i
        if ml:
            oval(grp, x, 43, 22, fill=TEAL6)
        else:
            oval(grp, x + .75, 43.75, 20.5, fill=WHITE, line="#B4BFCB", lw=1.5)
        text(grp, x - 4, 43, 30, 22, plain(dot), 11, WHITE if ml else SOFT, True, MONO, "c", "m", lh=13,
             wrap=False)


def header(s, root, dark=False):
    lq = q(root, "div.lq")
    disc(s, 72, 38, disc_of(q(lq, "span.disc")), name="Leitfrage-Scheibe")
    text(s, 116, 40.5, 700, 27, plain(q(lq, "p")), 20, TEAL3 if dark else TEAL6, lh=27, wrap=False,
         name="Leitfrage")
    text(s, 72, 80, 1136, 41, plain(q(root, "header h2")), 36, WHITE if dark else INK, True, lh=41,
         tracking=-0.02, name="Titel")
    for tr in qa(root, "div.tracker"):
        tracker(s, tr, dark)
    for cd in qa(root, "div.cdots"):
        cdots(s, cd)


LOGO = BUILD["_small_logo"]()


def footer(s, root, num, dark=False):
    """Fußzeile mit Logo, Projekttitel, Canvas-Bezügen (Mini-Scheiben), Sprecher und Seitenzahl."""
    ft = q(root, "footer")
    border = GLASS_LINE if dark else LINE2
    line(s, 0, FOOT_Y, W, FOOT_Y, border, 1, name="Fußlinie")
    if dark:
        box(s, 72, 672, 64, 26, fill=WHITE, radius=4, name="Logo-Grund")
        picture(s, LOGO, 78, 676, h=18, name="Logo")
        x = 152
    else:
        picture(s, LOGO, 72, 672, h=26, name="Logo")
        x = 163.2
    title = plain(next(sp for sp in ft.findall("span") if not classes(sp)))
    text(s, x, 675, measure(title, 16) + 12, 20, title, 16, "#7FB4D8" if dark else "#6B7887", lh=20, wrap=False)
    x += measure(title, 16) + 16
    for crefs in qa(ft, "span.crefs"):
        x += 14
        line(s, x, 674, x, 696, border, 1)
        x += 17
        for cref in qa(crefs, "span.cref"):
            disc(s, x, 674, disc_of(q(cref, "span.disc")))
            nr = plain(q(cref, "b"))
            name = plain(cref)[len(nr):].strip()
            x += 28
            text(s, x, 676, 30, 18, nr, 14, TEAL3 if dark else TEAL6, True, MONO, lh=18, wrap=False)
            x += measure(nr, 14, True, MONO) + 6
            text(s, x, 676, measure(name, 14) + 10, 18, name, 14, PALE if dark else MUTED, lh=18, wrap=False)
            x += measure(name, 14) + 12
    for sponsor in qa(ft, "span.sponsor"):
        x += 14
        line(s, x, 674, x, 696, border, 1)
        text(s, x + 17, 676, 560, 18, plain(sponsor), 14, SOFT, lh=18, wrap=False)
    num_w = max(64, measure(num, 16, font=MONO))
    text(s, 1208 - num_w - 16 - 120, 675, 120, 20, BUILD["SPEAKER"], 16, PALE if dark else MUTED, True,
         align="r", lh=20, wrap=False)
    text(s, 1208 - num_w, 674.5, num_w, 21, num, 16, "#7FB4D8" if dark else "#6B7887", font=MONO, align="r",
         lh=21, wrap=False, name="Seitenzahl")


def notes(slide, spec):
    """Notizen wie im Sprechzettel: Stichworte, Sprechtext, Lesehilfe, Rückfragen."""
    sections = []
    if spec.keywords:
        sections.append(("STICHWORTE", [[("• " + _plain(k), False)] for k in spec.keywords]))
    if spec.speech:
        sections.append(("SPRECHTEXT", [[(_plain(t), False)] for t in spec.speech]))
    if spec.reading:
        sections.append(("SO LIEST DU DIE FOLIE", [[("• " + _plain(t), False)] for t in spec.reading]))
    if spec.questions:
        sections.append(("FALLS GEFRAGT WIRD",
                         [[("• " + _plain(qu) + " ", True), (_plain(a), False)] for qu, a in spec.questions]))
    tf = slide.notes_slide.notes_text_frame
    tf.text = ""
    first = True
    for n, (head, items) in enumerate(sections):
        for runs in ([[(head, True)]] + items + ([[("", False)]] if n < len(sections) - 1 else [])):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            for chunk, bold in runs:
                r = p.add_run()
                r.text = chunk
                r.font.bold = bold


def _plain(s: str) -> str:
    return html.unescape(re.sub(r"<[^>]+>", "", s))


# --------------------------------------------------------------------------- Bausteine


def icard(s, el, x, y, w, h, pad=(13, 16), psize=15.0):
    """Karte mit Icon-Scheibe, fettem Kopf und Text (ersetzt die früheren Akzentstreifen)."""
    ok = "ok" in classes(el)
    card(s, x, y, w, h, fill="#F5FAF3" if ok else WHITE, border="#CFE7C8" if ok else LINE)
    py, px = pad
    disc(s, x + 1 + px, y + 1 + py, disc_of(q(el, "span.disc")))
    tx = x + 1 + px + 34 + 14
    tw = x + w - 1 - px - tx
    text(s, tx, y + 2 + py, tw, 20, plain(q(el, "strong")), 16, NAVY7, True, lh=20)
    text(s, tx, y + 25 + py, tw, h - py - 30, rich(q(el, "p"), [("b", {"bold": True})]), psize, MUTED,
         lh=psize * 1.4)


def chip_note(s, el, x, y, w, h):
    box(s, x, y, w, h, fill=WHITE, line=LINE, radius=h / 2)
    disc(s, x + 7, y + (h - 30) / 2, disc_of(q(el, "span.disc")))
    text(s, x + 49, y, w - 68, h, plain(q(el, "p")), 15, INK, anchor="m", lh=20.25)


def merk(s, el, x, y, w, h, dark=False):
    """Merksatz-Band mit Glühbirnen-Scheibe."""
    if dark:
        dcard(s, x, y, w, h)
    else:
        box(s, x, y, w, h, fill="#E3F3F7", radius=10)
    b = 1 if dark else 0
    disc(s, x + 12 + b, y + (h - 30) / 2, disc_of(q(el, "span.disc")))
    text(s, x + 56 + b, y, w - 76 - b, h, plain(q(el, "p")), 18, WHITE if dark else TEAL7, anchor="m", lh=24.3,
         name="Merksatz")


def label(s, x, y, w, t, color=TEAL6, size=13):
    """Versal-Label mit Laufweite (fk, flabel, dnr, uk …)."""
    return text(s, x, y, w, 17, t, size, color, True, lh=17, tracking=.08, upper=True, wrap=False)


def fexpr(s, el, x, y, size=20.0) -> float:
    """Formel in Geist Mono: Operatoren türkis, Brüche mit Bruchstrich; gibt die Höhe zurück."""
    items = []
    if el.text and el.text.strip():
        items.append(("t", el.text.strip()))
    for ch in el:
        if "op" in classes(ch):
            items.append(("op", plain(ch)))
        elif "frac" in classes(ch):
            items.append(("frac", tuple(plain(c) for c in ch)))
        if ch.tail and ch.tail.strip():
            items.append(("t", ch.tail.strip()))
    fs = size * .8
    num_h = fs * 1.3 + 5
    frac_h = num_h + fs * 1.3 + 3
    total_h = frac_h if any(k == "frac" for k, _ in items) else size * 1.3
    cy = y + total_h / 2
    for kind, val in items:
        if kind == "frac":
            num_, den = val
            nw, dw = measure(num_, fs, True, MONO) + 8, measure(den, fs, True, MONO) + 8
            fw = max(nw, dw)
            top = cy - frac_h / 2
            text(s, x - 4, top, fw + 8, fs * 1.3, num_, fs, NAVY7, True, MONO, "c", lh=fs * 1.3, wrap=False)
            line(s, x + (fw - nw) / 2, top + num_h - 1, x + (fw + nw) / 2, top + num_h - 1, NAVY7, 2)
            text(s, x - 4, top + num_h + 3, fw + 8, fs * 1.3, den, fs, NAVY7, True, MONO, "c", lh=fs * 1.3,
                 wrap=False)
            x += fw + 8
        else:
            tw = measure(val, size, True, MONO)
            text(s, x, cy - size * .65, tw + 10, size * 1.3, val, size, TEAL6 if kind == "op" else NAVY7, True,
                 MONO, lh=size * 1.3, wrap=False)
            x += tw + 8
    return total_h


# --------------------------------------------------------------------------- Diagramm-Helfer


def _plot_layout(chart, x, y, w, h):
    """Innere Plotfläche fest setzen (Anteile der Diagrammfläche), damit Beschriftungen passen."""
    plot_area = chart._chartSpace.chart.plotArea
    old = plot_area.find(qn("c:layout"))
    if old is not None:
        plot_area.remove(old)
    layout = OxmlElement("c:layout")
    manual = OxmlElement("c:manualLayout")
    for tag, val in [("c:layoutTarget", "inner"), ("c:xMode", "edge"), ("c:yMode", "edge"),
                     ("c:x", x), ("c:y", y), ("c:w", w), ("c:h", h)]:
        el = OxmlElement(tag)
        el.set("val", str(val))
        manual.append(el)
    layout.append(manual)
    plot_area.insert(0, layout)


def _base(chart, size=15):
    chart.has_legend = False
    chart.has_title = False
    chart.font.name = SANS
    chart.font.size = pt(size)
    chart.font.color.rgb = rgb(SOFT)
    sp = chart._chartSpace.find(qn("c:spPr"))  # keine Rahmen/Füllung der Diagrammfläche
    if sp is None:
        sp = OxmlElement("c:spPr")
        chart._chartSpace.append(sp)
    sp.append(OxmlElement("a:noFill"))
    ln = OxmlElement("a:ln")
    ln.append(OxmlElement("a:noFill"))
    sp.append(ln)


def place_chart(s, chart_type, data, rect, pad=(0, 0, 0, 0), name=None):
    """Diagramm so setzen, dass die innere Plotfläche genau ``rect`` (Folien-px) belegt."""
    x, y, w, h = rect
    left, top, right, bottom = pad
    fw, fh = w + left + right, h + top + bottom
    gf = s.shapes.add_chart(chart_type, E(x - left), E(y - top), E(fw), E(fh), data)
    if name:
        gf.name = name
    chart = gf.chart
    _base(chart)
    _plot_layout(chart, left / fw, top / fh, w / fw, h / fh)
    return gf, chart


def _axis(ax, lo=None, hi=None, unit=None, fmt=None, grid=True, visible=True, line_visible=False, size=15,
          font=MONO, color=SOFT):
    if lo is not None:
        ax.minimum_scale = lo
    if hi is not None:
        ax.maximum_scale = hi
    if unit is not None:
        ax.major_unit = unit
    ax.has_major_gridlines = grid
    if grid:
        ax.major_gridlines.format.line.color.rgb = rgb(LINE2)
        ax.major_gridlines.format.line.width = pt(1)
    ax.has_minor_gridlines = False
    ax.major_tick_mark = XL_TICK_MARK.NONE
    ax.minor_tick_mark = XL_TICK_MARK.NONE
    if line_visible:
        ax.format.line.color.rgb = rgb("#B4BFCB")
        ax.format.line.width = pt(1)
    else:
        ax.format.line.fill.background()
    if fmt:
        ax.tick_labels.number_format = fmt
        ax.tick_labels.number_format_is_linked = False
    ax.tick_labels.font.size = pt(size)
    ax.tick_labels.font.name = font
    ax.tick_labels.font.color.rgb = rgb(color)
    if not visible:
        ax.tick_label_position = XL_TICK_LABEL_POSITION.NONE


def _cross_between(chart, val="midCat"):
    for el in chart._chartSpace.iter(qn("c:crossBetween")):
        el.set("val", val)


def _series_line(series, color, width=2.5, dash=None, cap_flat=False):
    series.format.line.color.rgb = rgb(color)
    series.format.line.width = pt(width)
    if dash:
        series.format.line.dash_style = dash
    if cap_flat:
        series.format.line._get_or_add_ln().set("cap", "flat")
    series.smooth = False


def _no_line(series):
    series.format.line.fill.background()


def _marker(series, style, size, fill, border=WHITE, border_w=1.5, alpha=None):
    series.marker.style = style
    series.marker.size = size
    series.marker.format.fill.solid()
    series.marker.format.fill.fore_color.rgb = rgb(fill)
    if alpha is not None:
        _set_alpha(series.marker.format.fill._xPr.find(qn("a:solidFill")), alpha)
    if border:
        series.marker.format.line.color.rgb = rgb(border)
        series.marker.format.line.width = pt(border_w)
    else:
        series.marker.format.line.fill.background()


def _no_marker(series):
    series.marker.style = XL_MARKER_STYLE.NONE


class Mapper:
    """Datenkoordinaten -> Folienkoordinaten für eine feste Plotfläche (Folien-px)."""

    def __init__(self, rect, x0, x1, y0, y1):
        self.left, self.top, self.w, self.h = rect
        self.x0, self.x1, self.y0, self.y1 = x0, x1, y0, y1

    def x(self, v):
        return self.left + (v - self.x0) / (self.x1 - self.x0) * self.w

    def y(self, v):
        return self.top + (1 - (v - self.y0) / (self.y1 - self.y0)) * self.h


def xy_data(series_spec):
    data = XyChartData()
    for name, points in series_spec:
        ser = data.add_series(name)
        for px_, py_ in points:
            ser.add_data_point(px_, py_)
    return data


def _merge_line_into(target_chart, line_chart):
    """Linienserien in die Plotfläche eines Säulen-/Flächendiagramms übernehmen (gemeinsame Achsen)."""
    plot_area = target_chart._chartSpace.chart.plotArea
    host = next(el for el in plot_area if el.tag in (qn("c:areaChart"), qn("c:barChart")))
    ax_ids = [el.get("val") for el in host.findall(qn("c:axId"))]
    line_el = copy.deepcopy(line_chart._chartSpace.chart.plotArea.find(qn("c:lineChart")))
    for i, el in enumerate(line_el.findall(qn("c:axId"))):
        el.set("val", ax_ids[i])
    offset = len(host.findall(qn("c:ser")))
    for ser in line_el.findall(qn("c:ser")):
        for tag in ("c:idx", "c:order"):
            el = ser.find(qn(tag))
            el.set("val", str(int(el.get("val")) + offset))
    host.addnext(line_el)


def _combo(s, host_type, host_data, line_data, rect, pad):
    """Säulen- oder Flächendiagramm mit zusätzlicher Linie auf denselben Achsen."""
    gf, host = place_chart(s, host_type, host_data, rect, pad)
    tmp = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, E(rect[0]), E(rect[1]), E(rect[2]), E(rect[3]),
                             line_data)
    _merge_line_into(host, tmp.chart)
    rid = tmp._element.xpath(".//c:chart/@r:id")[0]
    s.shapes._spTree.remove(tmp._element)
    s.part.drop_rel(rid)
    return gf, host


def _diamond(s, cx, cy, size, color, name=None):
    d = s.shapes.add_shape(MSO_SHAPE.DIAMOND, E(cx - size / 2), E(cy - size / 2), E(size), E(size))
    if name:
        d.name = name
    d.fill.solid()
    d.fill.fore_color.rgb = rgb(color)
    d.line.color.rgb = rgb(WHITE)
    d.line.width = pt(2)
    _no_shadow(d)
    return d


def chart_card_title(s, el, x, y, w, h=20):
    """p.ch-t: fetter Titel, Zusatz in Klammern normal und grau."""
    text(s, x, y, w, h, rich(el, [("span", {"bold": False, "color": SOFT})]), 16, INK, True, lh=20)


# --------------------------------------------------------------------------- Folie 1: ML Canvas


def canvas_card(s, el, x, y, w, h):
    ml = "ml" in classes(el)
    if ml:
        card(s, x, y, w, h, border=TEAL6, lw=2, name=f"Canvas {plain(q(el, 'span.cv-nr'))}")
    else:
        card(s, x, y, w, h, fill=MIST, border=LINE2, name=f"Canvas {plain(q(el, 'span.cv-nr'))}")
    b = 2 if ml else 1
    cx, cy = x + b + 16, y + b + 14
    disc(s, cx, cy, disc_of(q(el, "span.disc")))
    nr = plain(q(el, "span.cv-nr"))
    text(s, cx + 44, cy + 9, 24, 18, nr, 14, TEAL6 if ml else SOFT, True, MONO, lh=18, wrap=False)
    nx = cx + 44 + measure(nr, 14, True, MONO) + 8
    text(s, nx, cy + 9, 160, 18, plain(q(el, "span.cv-name")), 14, MUTED, True, lh=18, tracking=.06,
         upper=True, wrap=False)
    for own in qa(el, "span.cv-own"):
        text(s, x + w - b - 16 - 90, cy + 9.5, 90, 17, plain(own), 13, SOFT, align="r", lh=17, wrap=False)
    tw = w - 2 * b - 32
    text(s, cx, cy + 46, tw, 23, plain(q(el, "h4")), 18, NAVY7 if ml else INK, True, lh=23)
    text(s, cx, cy + 73, tw, 40, plain(q(el, "p")), 15, MUTED, lh=20)


def chevron(s, cx, cy):
    """CSS-Chevron: 12-px-Quadrat mit 3-px-Rand oben/rechts, um 45° gedreht (Linienmitte 4,5 px innen)."""
    r = 4.5 * math.sqrt(2)
    return poly(s, [(cx, cy - r), (cx + r, cy), (cx, cy + r)], stroke=TEAL3, width=3, miter=True, name="Chevron")


class _UnitMap(SvgMap):
    """1:1-Abbildung für SVGs ohne viewBox (der Canvas-Pfad liegt direkt in Folienkoordinaten)."""

    def __init__(self):
        self.ox, self.oy, self.k = 0.0, 0.0, 1.0


def slide_canvas(prs, spec, num, facts):
    s, root = page(prs, spec)
    header(s, root)
    path_svg = q(root, "svg.cv-path")
    svg_draw(s, path_svg, SvgMap(path_svg, 0, 0, W) if path_svg.get("viewbox") else _UnitMap())
    for art in qa(root, "article.cv"):
        canvas_card(s, art, _style_px(art, "left"), _style_px(art, "top"), _style_px(art, "width"),
                    _style_px(art, "height"))
    for ch in qa(root, "span.chev"):
        chevron(s, _style_px(ch, "left") + 6, _style_px(ch, "top") + 6)
    back = q(root, "span.cv-back")
    disc(s, 1172, 522, disc_of(q(back, "span.disc")), name="Rücksprung zu 01")
    text(s, 1168, 558, 40, 18, plain(q(back, "b")), 14, TEAL6, True, MONO, "c", lh=18, wrap=False)
    footer(s, root, num)


# --------------------------------------------------------------------------- Folie 2: Fehlerkosten


def slide_fehlerkosten(prs, spec, num, facts):
    s, root = page(prs, spec, MIST)
    header(s, root)
    xs, ws = [72, 433.2, 892.8], [315.2, 413.6, 315.2]
    for el, x, w in zip(qa(root, "div.ucard"), xs, ws):
        wide = "wide" in classes(el)
        card(s, x, 143, w, 89.2, fill="#E3F3F7" if wide else WHITE, border="#B2DEE7" if wide else LINE)
        disc(s, x + 17, 170.6, disc_of(q(el, "span.disc")))
        label(s, x + 63, 156, w - 80, plain(q(el, "span.uk")), SOFT)
        text(s, x + 63, 176, w - 78, 44, rich(q(el, "strong"), [(".mono", {"font": MONO, "color": TEAL6})]), 16,
             INK, True, lh=21.6)
    for svg, x in zip(qa(root, "span.uarrow > svg"), (397.2, 856.8)):
        glyph(s, x, 174.6, svg)
    for el, x in zip(qa(root, "div.mhead"), (254, 738)):
        disc(s, x + 4, 256.2, disc_of(q(el, "span.disc")))
        text(s, x + 42, 250.2, 420, 20, plain(q(el, "b")), 16, INK, True, lh=20, wrap=False)
        text(s, x + 42, 272.2, 420, 18, plain(q(el, "div > span")), 14, SOFT, lh=18, wrap=False)
    rows = [(302.2, 101.8), (414.1, 122.8)]
    for el, (y, h) in zip(qa(root, "div.mrow"), rows):
        top = y + (h - 59.4) / 2
        text(s, 76, top, 164, 23, plain(q(el, "b")), 18, INK, True, lh=23)
        text(s, 76, top + 23, 164, 37, plain(q(el, "span")), 14, SOFT, lh=18.2)
    for i, el in enumerate(qa(root, "div.cost")):
        x, (y, h) = (254, 738)[i % 2], rows[i // 2]
        card(s, x, y, 470, h)
        disc(s, x + 17, y + 15, disc_of(q(el, "span.disc")))
        tw = 470 - 67 - 17
        text(s, x + 67, y + 15, tw, 22, plain(q(el, "strong")), 17, INK, True, lh=22)
        body, lever = qa(el, "p")
        text(s, x + 67, y + 39, tw, 21, plain(body), 15.5, MUTED, lh=20.9)
        lines = len(wrap_lines(plain(lever), tw - 21, 15.5))
        glyph(s, x + 67, y + 65.9 + (lines * 20.9 - 15) / 2, q(lever, "svg"))
        text(s, x + 88, y + 65.9, tw - 21, 21 * lines, plain(lever), 15.5, TEAL6, lh=20.9)
    text(s, 72, 552.8, 1136, 18, plain(q(root, "p.fnote")), 14, SOFT, lh=18)
    footer(s, root, num)


# --------------------------------------------------------------------------- Folie 3: Roter Faden (Split)


def ring_arc(s, cx, cy, r, clip, color, width):
    """Sichtbarer Teil eines Kreisrings innerhalb von ``clip`` (overflow:hidden im CSS)."""
    x0, y0, x1, y1 = clip

    def point(a):
        return cx + r * math.cos(a), cy + r * math.sin(a)

    def inside(a):
        px, py = point(a)
        return x0 <= px <= x1 and y0 <= py <= y1

    def edge(a_in, a_out):
        for _ in range(40):
            mid = (a_in + a_out) / 2
            a_in, a_out = (mid, a_out) if inside(mid) else (a_in, mid)
        return a_in

    n = 720
    step = 2 * math.pi / n
    flags = [inside(i * step) for i in range(n)]
    start = next(i for i in range(n) if flags[i] and not flags[i - 1])
    pts = [point(edge(start * step, (start - 1) * step))]
    i = start
    while flags[i % n] and i < start + n:
        pts.append(point(i * step))
        i += 1
    pts.append(point(edge((i - 1) * step, i * step)))
    pts = [(min(max(px, x0), x1), min(max(py, y0), y1)) for px, py in pts]
    return poly(s, pts, stroke=color, width=width, name="Ring")


def slide_faden(prs, spec, num, facts):
    s, root = page(prs, spec)
    box(s, 0, 0, 512, FOOT_Y, fill=NAVY, name="Split links")
    # Ringe aus dem CSS (.r1/.r2): Mittelpunkt und Radius der Linienmitte, abgeschnitten am Panel
    ring_arc(s, 502, 20, 209, (0, 0, 512, FOOT_Y), "rgba(108,192,210,.2)", 2)
    ring_arc(s, 492, 30, 139, (0, 0, 512, FOOT_Y), "rgba(108,192,210,.12)", 2)
    text(s, 56, 40, 408, 86.4, plain(q(root, "span.big-nr")), 96, mix(TEAL3, NAVY, .5), True, MONO, lh=86.4,
         wrap=False)
    text(s, 56, 136.4, 420, 42, plain(q(root, "div.split-l h2")), 38, WHITE, True, lh=41.8, tracking=-0.02,
         name="Titel")
    hook = q(root, "div.hook")
    box(s, 56, 337, 408, 284, fill=GLASS, line=GLASS_LINE, radius=12, name="Hook")
    top = q(hook, "div.hook-top")
    disc(s, 75, 354, disc_of(q(top, "span.disc")))
    text(s, 119, 360, 330, 22, plain(top), 17, TEAL3, font=MONO, lh=22, wrap=False)
    text(s, 75, 400, 370, 18, plain(q(hook, "span.hook-lbl")), 14, "#9CC3E0", lh=18)
    end = zw_from(s, q(hook, "span.zw"), 75, 424)
    text(s, end + 10, 447, 70, 29, plain(kid(q(hook, "div.hook-val"), "span")), 22, PALE, lh=29, wrap=False)
    text(s, 75, 492, 380, 22, rich(q(hook, "p.hook-exp"), [("b", {"bold": True, "color": WHITE}),
                                                             (".mono", {"font": MONO})]), 17, PALE, lh=22)
    for el, y in zip(qa(hook, "div.hook-then > span"), (526, 569)):
        new = "new" in classes(el)
        box(s, 75, y, 370, 35, fill=TEAL6 if new else "rgba(255,255,255,.05)", radius=8)
        glyph(s, 87, y + 8.5, q(el, "svg"))
        text(s, 115, y + 7.75, 325, 19.5, plain(el), 15, WHITE if new else PALE, lh=19.5, wrap=False)
    lq = q(root, "div.split-r div.lq")
    disc(s, 560, 36, disc_of(q(lq, "span.disc")), name="Leitfrage-Scheibe")
    text(s, 604, 38.5, 620, 27, plain(q(lq, "p")), 20, TEAL6, lh=27, wrap=False, name="Leitfrage")
    box(s, 579, 128, 2, 341, fill=LINE, name="Faden")
    for el, y in zip(qa(root, "div.grp"), (86, 307)):
        main = (el.text or "").strip()
        color = TEAL6 if "late" in classes(el) else NAVY7
        text(s, 616, y, 400, 18, main, 13, color, True, lh=18, tracking=.08, upper=True, wrap=False)
        mx = 616 + measure(main.upper(), 13, True, tracking=.08) + 6
        text(s, mx, y + 1, 300, 17, plain(q(el, "span")), 13, SOFT, lh=17, wrap=False)
    for el, y in zip(qa(root, "div.fstep"), (110, 166, 222, 331, 387, 443)):
        disc(s, 560, y + 6, disc_of(q(el, "span.disc")))
        nr = plain(q(el, "strong span.mono"))
        text(s, 616, y + 4.8, 30, 22, nr, 17, SOFT, True, MONO, lh=22, wrap=False)
        text(s, 616 + measure(nr, 17, True, MONO) + 4 + measure(" ", 17, True), y + 4.8, 560, 22,
             plain(q(el, "strong"))[len(nr):].strip(), 17, INK, True, lh=22, wrap=False)
        text(s, 616, y + 27.8, 600, 20, plain(q(el, "p")), 15, MUTED, lh=19.5)
    cut = q(root, "div.fcut")
    glyph(s, 616, 280.5, q(cut, "svg"))
    cut_text = plain(cut)
    text(s, 640, 280, 200, 17, cut_text, 13, SOFT, True, lh=17, tracking=.08, upper=True, wrap=False)
    line(s, 640 + measure(cut_text.upper(), 13, True, tracking=.08) + 8, 288.5, 1224, 288.5, "#B4BFCB", 2,
         MSO_LINE_DASH_STYLE.DASH)
    story = q(root, "p.story")
    glyph(s, 560, 622, q(story, "svg"))
    text(s, 584, 621, 640, 18, plain(story), 14, SOFT, lh=18, wrap=False)
    footer(s, root, num)


# --------------------------------------------------------------------------- Folie 4: Validierung


def slide_validierung(prs, spec, num, facts):
    s, root = page(prs, spec)
    header(s, root)
    card(s, 72, 143, 1136, 281.5)
    svg = chart_svg(q(root, "div.chart-card"))
    svg_draw(s, svg, SvgMap(svg, 93, 160, 1094))
    for el, x in zip(qa(root, "div.row-3 div.icard"), (72, 456, 840)):
        icard(s, el, x, 440.6, 368, 108.9, pad=(11, 14), psize=14.5)
    lo = q(root, "div.leftout")
    box(s, 72, 565.5, 1136, 46, fill="#F7F9FB", line=LINE2, radius=10)
    disc(s, 87, 573.5, disc_of(q(lo, "span.disc")))
    key = plain(q(lo, "span.lo-k"))
    label(s, 125, 580, 200, key, SOFT)
    x = 125 + measure(key.upper(), 13, True, tracking=.08) + 12
    for pill in qa(lo, "span.lo"):
        b = plain(q(pill, "b"))
        rest = plain(pill)[len(b):].strip()
        pw = 24 + measure(b, 14, True) + 4 + measure(" " + rest, 14) + 2
        box(s, x, 574.5, pw, 28, fill=WHITE, line=LINE2, radius=14)
        text(s, x + 12, 574.5, pw - 12, 28, [[(b, {"bold": True, "color": INK}), ("  " + rest, {})]], 14, MUTED,
             anchor="m", lh=18, wrap=False)
        x += pw + 8
    footer(s, root, num)


# --------------------------------------------------------------------------- Folie 5: Zielgröße


VLS_EXAMPLE = [([10_000, 40_000], 40_000), ([200, 200], 250)]  # wie chart_vls_example im HTML-Generator


def slide_vls(prs, spec, num, facts):
    s, root = page(prs, spec, MIST)
    header(s, root)
    for el, y, h in zip(qa(root, "div.fcard"), (143, 257.6), (102.6, 79)):
        card(s, 72, y, 618.9, h)
        disc(s, 91, y + 15, disc_of(q(el, "span.disc")))
        label(s, 139, y + 15, 400, plain(q(el, "span.fk")))
        fexpr(s, q(el, "div.fexpr"), 139, y + 38, 20)
    icard(s, q(root, "div.icard"), 72, 348.6, 618.9, 94)
    chip_note(s, q(root, "div.chip-note"), 72, 454.6, 618.9, 44)
    card(s, 712.9, 143, 495.1, 343.4)
    svg = chart_svg(q(root, "div.chart-card"))
    m = SvgMap(svg, 733.9, 160, 453.1)
    svg_draw(s, svg, m, skip=lambda el: el.tag == "path")  # Balken: natives Diagramm
    for p, (values, vmax) in enumerate(VLS_EXAMPLE):
        data = CategoryChartData()
        data.categories = ["Zähler A", "Zähler B"]
        data.add_series(["Energiemenge (kWh)", "Vollaststunden (VLS-h)"][p], values)
        # Balken bei x0+14 und x0+114 (Breite 78): Kategorien 100 breit, Lücke 22/78
        _, ch = place_chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, data, m.rect(240 * p + 3, 70, 200, 180),
                            (2, 2, 2, 2), name=f"Diagramm VLS-Beispiel {p + 1}")
        _axis(ch.value_axis, 0, vmax, grid=False, visible=False)
        _axis(ch.category_axis, grid=False, visible=False)
        plot = ch.plots[0]
        plot.gap_width, plot.overlap = 28, 0
        for j, point in enumerate(plot.series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = rgb([C["navy300"], C["ist"]][j])
    ev = q(root, "div.evidence")
    box(s, 712.9, 498.4, 495.1, 112, fill=NAVY, radius=10)
    text(s, 732.9, 524.4, 175, 60, plain(q(ev, "span.ev-v")), 46, WHITE, True, MONO, lh=60, tracking=-0.02,
         wrap=False)
    text(s, 911, 512.4, 285, 84, rich(q(ev, "p"), [(".mono", {"font": MONO, "color": WHITE})]), 15, PALE, lh=21)
    footer(s, root, num)


# --------------------------------------------------------------------------- Folie 6: Modellwahl


def model_chart(s, m: SvgMap, rows, bar_h, folds=True):
    """Balken als dicke XY-Linien (lineare x-Achse wie im SVG), Fold-Rauten und ±1-Std-Striche als Formen."""
    n = len(rows)
    rect = m.rect(200, 60, 234, 36 * n)
    spec = [(row["name"], [(0, n - i), (row["value"], n - i)]) for i, row in enumerate(rows)]
    _, chart = place_chart(s, XL_CHART_TYPE.XY_SCATTER_LINES, xy_data(spec), rect, (10, 8, 14, 32),
                           name="Diagramm Modellvergleich")
    _axis(chart.value_axis, 0.5, n + 0.5, 1, grid=False, visible=False)
    _axis(chart.category_axis, 0, 28_000, 7_000, fmt='#,##0,', grid=True)
    for i, series in enumerate(chart.plots[0].series):
        _series_line(series, C["prognose"] if i == 0 else C["grey300"], bar_h * m.k, cap_flat=True)
        _no_marker(series)
    mp = Mapper(rect, 0, 28_000, 0.5, n + 0.5)
    for i, row in enumerate(rows):
        cy = mp.y(n - i)
        if row.get("std"):
            lo, hi = mp.x(row["value"] - row["std"]), mp.x(row["value"] + row["std"])
            line(s, lo, cy, hi, cy, MUTED, 1.5)
            line(s, lo, cy - 6 * m.k, lo, cy + 6 * m.k, MUTED, 1.5)
            line(s, hi, cy - 6 * m.k, hi, cy + 6 * m.k, MUTED, 1.5)
        for k, v in enumerate(row.get("folds", []) if folds else []):
            _diamond(s, mp.x(v), cy, 15 * m.k, FOLD_COLORS[k], f"{FOLD_LABELS[k]} {row['name']} {round(v)}")


def slide_modell(prs, spec, num, facts):
    s, root = page(prs, spec)
    header(s, root)
    cv = sorted(facts["cv"], key=lambda r: r["mean"])
    bench = sorted(facts["bench"], key=lambda r: r["rmse"])
    rows = [
        ([{"name": r["name"], "value": r["mean"], "folds": r["folds"], "std": r["std"]} for r in cv], 24),
        ([{"name": r["name"], "value": r["rmse"]} for r in bench], 26),
    ]
    for i, (el, (data, bar_h)) in enumerate(zip(qa(root, "div.two-charts div.chart-card"), rows)):
        x = 72 + i * 578
        card(s, x, 143, 558, 328.2)
        chart_card_title(s, q(el, "p.ch-t"), x + 17, 156, 524)
        svg = chart_svg(el)
        m = SvgMap(svg, x + 17, 184, 524)
        model_chart(s, m, data, bar_h)
        svg_draw(s, svg, m, plot=(200, 54, 434, 208), skip=is_tick)
        note = q(el, "p.ch-note")
        if "mono" in classes(note):
            text(s, x + 17, 427.1, 530, 36, plain(note), 13, NAVY7, font=MONO, lh=17.55)
        else:
            text(s, x + 17, 427.1, 530, 19, plain(note), 14, MUTED, lh=18.9)
    tiles = qa(root, "div.tiles > div")
    t0, t1, t2 = tiles
    box(s, 72, 487.2, 380.1, 128.8, fill=NAVY, line=NAVY9, radius=10)
    zw_from(s, q(t0, "span.zw"), 89, 500.2)
    text(s, 89, 554.6, 350, 40, rich(q(t0, "p"), [(".mono", {"font": MONO, "color": WHITE})]), 14.5, PALE,
         lh=19.6)
    card(s, 468.1, 487.2, 362, 128.8)
    disc(s, 485.1, 534.6, disc_of(q(t1, "span.disc")))
    label(s, 533.1, 508.7, 280, plain(q(t1, "span.tk")), SOFT)
    text(s, 533.1, 528.7, 284, 45, rich(q(t1, "p.tv"), [("b", {"bold": True, "color": NAVY7})]), 16, INK, lh=22.4)
    text(s, 533.1, 576.5, 284, 18, plain(q(t1, "p.ts")), 13.5, SOFT, lh=18)
    card(s, 846, 487.2, 362, 128.8)
    disc(s, 863, 534.6, disc_of(q(t2, "span.disc")))
    label(s, 911, 500.2, 280, plain(q(t2, "span.tk")), SOFT)
    for j, li in enumerate(qa(t2, "ul.top3 > li")):
        y = 521.2 + 20.3 * j
        text(s, 911, y, 200, 20.3, plain(q(li, "span")), 15, INK, lh=20.3, wrap=False)
        text(s, 1030, y, 149.8, 20.3, plain(q(li, "b")), 15, TEAL6, True, MONO, "r", lh=20.3, wrap=False)
    text(s, 911, 585, 280, 18, plain(q(t2, "p.ts")), 13.5, SOFT, lh=18)
    footer(s, root, num)


# --------------------------------------------------------------------------- Folie 7: Prüfhinweis


def slide_schwelle(prs, spec, num, facts):
    s, root = page(prs, spec, MIST)
    header(s, root)
    thr = facts["threshold"]
    box(s, 106, 173, 2, 246, fill=LINE, name="Kette")
    for el, (y, h) in zip(qa(root, "div.clink"), [(143, 94.4), (247.5, 97), (354.4, 97)]):
        card(s, 72, y, 484.3, h)
        disc(s, 89, y + 13, disc_of(q(el, "span.disc")))
        label(s, 139, y + 13, 400, plain(q(el, "span.fk")))
        fh = fexpr(s, q(el, "div.fexpr"), 139, y + 36, 18)
        for p in qa(el, "p"):
            text(s, 139, y + 36 + fh + 5, 410, 20, rich(p, [("b", {"bold": True, "color": NAVY7})]), 14.5, MUTED,
                 lh=19.6, wrap=False)
    x = 72
    stairs = q(root, "div.stairs")
    for el in stairs:
        if el.tag == "svg":
            glyph(s, x, 463.4 + 7.5, el)
            x += 16 + 6
            continue
        on = "on" in classes(el)
        t = plain(el)
        pw = measure(t, 14.5, on) + 26
        box(s, x, 463.4, pw, 31, fill="#A86505" if on else WHITE, line="#A86505" if on else LINE, radius=15.5)
        text(s, x, 463.4, pw, 31, t, 14.5, WHITE if on else MUTED, on, align="c", anchor="m", lh=19, wrap=False)
        x += pw + 6
    chart_el = q(root, "div.chart-card")
    card(s, 578.3, 143, 629.7, 416.7)
    chart_card_title(s, q(chart_el, "p.ch-t"), 599.3, 160, 600)
    svg = chart_svg(chart_el)
    m = SvgMap(svg, 599.3, 188, 587.7)
    pts = facts["ranked"]
    base = [(x_, y_) for x_, y_ in pts if y_ <= thr]
    tail = [(x_, y_) for x_, y_ in pts if y_ > thr]
    _, chart = place_chart(s, XL_CHART_TYPE.XY_SCATTER_LINES, xy_data([
        ("bis zur Schwelle", base), ("oberhalb der Schwelle", [base[-1], *tail]),
        ("Schwelle q99", [(0, thr), (100, thr)]), ("99 %", [(99, 0), (99, 500)]),
    ]), m.rect(58, 28, 518, 252), (52, 8, 10, 30), name="Diagramm Kalibrierung")
    _axis(chart.value_axis, 0, 500, 100, fmt="0")
    _axis(chart.category_axis, 0, 100, 25, fmt='[<100]0" %";""', grid=False)
    sr = list(chart.plots[0].series)
    _series_line(sr[0], C["residuum"], 2.5)
    _no_marker(sr[0])
    _series_line(sr[1], C["anomalie"], 2.5)
    _marker(sr[1], XL_MARKER_STYLE.CIRCLE, 7, C["anomalie"], border_w=1.5)
    _series_line(sr[2], C["schwelle"], 2, MSO_LINE_DASH_STYLE.DASH)
    _no_marker(sr[2])
    _series_line(sr[3], C["grey400"], 1, MSO_LINE_DASH_STYLE.SQUARE_DOT)
    _no_marker(sr[3])
    svg_draw(s, svg, m, plot=(58, 28, 576, 280), skip=is_tick)
    text(s, 599.3, 521.8, 600, 19, plain(q(chart_el, "p.ch-note")), 14, MUTED, lh=18.9)
    merk(s, q(root, "div.merk"), 72, 587, 1136, 52)
    footer(s, root, num)


# --------------------------------------------------------------------------- Folie 8: Kalibrierung (dunkel)


def slide_kalibrierung(prs, spec, num, facts):
    s, root = page(prs, spec, NAVY)
    header(s, root, dark=True)
    for el, y in zip(qa(root, "div.cal-step"), (143, 259, 375)):
        dcard(s, 72, y, 543.4, 104)
        disc(s, 91, y + 15, disc_of(q(el, "span.disc")))
        text(s, 145, y + 15, 455, 23, plain(q(el, "strong")), 18, WHITE, True, lh=23)
        text(s, 145, y + 44, 455, 44, plain(q(el, "p")), 15.5, PALE, lh=21.7)
    p1, p2 = qa(root, "div.dpanel")
    dcard(s, 637.4, 143, 570.6, 242.3)
    label(s, 656.4, 156, 532, plain(q(p1, "span.dnr")), TEAL3)
    svg = chart_svg(p1)
    svg_draw(s, svg, SvgMap(svg, 656.4, 181, 532.6))
    line(s, 656.4, 295.6, 1189, 295.6, GLASS_LINE, 1)
    for col, x in zip(qa(p1, "div.cal-stats > div"), (656.4, 928.7)):
        text(s, x, 304.6, 262, 26, plain(q(col, "b")), 20, WHITE, True, MONO, lh=26, wrap=False)
        text(s, x, 332.6, 262, 40, plain(q(col, "span")), 14, "#9CC3E0", lh=18.9)
    dcard(s, 637.4, 397.4, 570.6, 118.4)
    label(s, 656.4, 410.4, 532, plain(q(p2, "span.dnr")), TEAL3)
    text(s, 656.4, 435.4, 540, 70, ["• " + plain(li) for li in qa(p2, "li")], 15.5, PALE, lh=22.475)
    merk(s, q(root, "div.merk"), 72, 566.4, 1136, 72.6, dark=True)
    footer(s, root, num, dark=True)


# --------------------------------------------------------------------------- Folie 9: Band 2025


def slide_band(prs, spec, num, facts):
    s, root = page(prs, spec)
    header(s, root)
    card(s, 72, 143, 520, 496.4)
    svg = chart_svg(q(root, "div.chart-card"))
    m = SvgMap(svg, 87, 156, 490)
    band = q(svg, "polygon")  # Toleranzband als Fläche hinter dem Diagramm
    raw = [float(v) for v in re.findall(r"[-\d.]+", band.get("points"))]
    poly(s, [(m.x(x_), m.y(y_)) for x_, y_ in zip(raw[::2], raw[1::2])], fill=band.get("fill"), closed=True,
         name="Toleranzband")
    vmax = 800
    rect = m.rect(62, 84, 424, 346)
    normal = [(x_, y_) for x_, y_ in facts["scatter_normal"]
              if x_ is not None and y_ is not None and 0 <= x_ <= vmax and 0 <= y_ <= vmax]
    alerts_in = [(x_, y_) for x_, y_ in facts["scatter_alerts"] if 0 <= x_ <= vmax and 0 <= y_ <= vmax]
    edge_y = vmax - 5 / 346 * vmax  # Dreieck am oberen Rand wie im SVG
    alerts_out = [(min(max(x_, 0), vmax), edge_y) for x_, y_ in facts["scatter_alerts"] if not (0 <= y_ <= vmax)]
    _, chart = place_chart(s, XL_CHART_TYPE.XY_SCATTER, xy_data([
        ("kein Hinweis", normal), ("Prüfhinweis", alerts_in), (f"über {vmax} VLS-h", alerts_out),
        ("Ist = Prognose", [(0, 0), (vmax, vmax)]),
    ]), rect, (46, 8, 12, 30), name="Diagramm Ist gegen Prognose")
    _axis(chart.value_axis, 0, vmax, 200, fmt="0")
    _axis(chart.category_axis, 0, vmax, 200, fmt="0", grid=False)
    sr = list(chart.plots[0].series)
    for ser in sr[:3]:
        _no_line(ser)
    _marker(sr[0], XL_MARKER_STYLE.CIRCLE, 3, C["navy300"], border=None, alpha=.55)
    _marker(sr[1], XL_MARKER_STYLE.CIRCLE, 7, C["anomalie"], border=WHITE, border_w=.75)
    _marker(sr[2], XL_MARKER_STYLE.TRIANGLE, 9, C["anomalie"], border=None)
    _series_line(sr[3], C["ist"], 1.5)
    _no_marker(sr[3])
    if os.environ.get("PPTX_QA_RENDER"):  # nur für die Sichtprüfung, nie ausgeliefert
        chart.plots[0]._element.find(qn("c:scatterStyle")).set("val", "marker")
        for ser, col in zip(sr[:3], [C["navy300"], C["anomalie"], C["anomalie"]]):
            ser.format.line.color.rgb = rgb(col)
    svg_draw(s, svg, m, plot=(62, 84, 486, 430), skip=is_tick)
    hero = q(root, "div.hero")
    box(s, 614, 143, 594, 99.6, fill=NAVY, radius=10)
    zw_from(s, q(hero, "span.zw"), 634, 160.3)
    text(s, 760.6, 159, 440, 23, plain(q(hero, "strong")), 18, WHITE, True, lh=23)
    text(s, 760.6, 186, 440, 41, rich(q(hero, "p")), 14.5, PALE, lh=20.3)
    wl = q(root, "div.wl-box")
    box(s, 614, 254.6, 594, 210.6, fill="#F7F9FB", line=LINE, radius=10)
    label(s, 631, 267.6, 560, plain(q(wl, "span.flabel")))
    for j, el in enumerate(qa(wl, "div.wl")):
        on = "on" in classes(el)
        x = 631 + 142.5 * j
        box(s, x, 292.6, 132.5, 80, fill=NAVY if on else WHITE, line=NAVY9 if on else LINE, radius=10)
        text(s, x + 13, 300.6, 110, 18, plain(q(el, "span.wl-q")), 14, PALE if on else MUTED, font=MONO, lh=18)
        text(s, x + 13, 318.6, 110, 28, plain(q(el, "span.wl-v")), 22, WHITE if on else INK, True, MONO, lh=28)
        text(s, x + 13, 346.6, 110, 18, plain(q(el, "span.wl-u")), 13.5, PALE if on else SOFT, lh=18)
    line(s, 631, 381.1, 1191, 381.1, "#B4BFCB", 1, MSO_LINE_DASH_STYLE.DASH)
    left, right = qa(wl, "div.tradeoff > span")
    text(s, 631, 387.6, 280, 18, plain(left), 13.5, SOFT, lh=18, wrap=False)
    text(s, 911, 387.6, 280, 18, plain(right), 13.5, SOFT, align="r", lh=18, wrap=False)
    text(s, 631, 411.6, 565, 41, plain(kid(wl, "p")), 14.5, MUTED, lh=20.3)
    chip_note(s, q(root, "div.chip-note"), 614, 477.2, 594, 54.5)
    footer(s, root, num)


# --------------------------------------------------------------------------- Folie 10: Fallbeispiel


def slide_fall(prs, spec, num, facts):
    s, root = page(prs, spec)
    header(s, root)
    thr, kw = facts["threshold"], facts["case_kw"]
    band = thr * kw
    ist, prog, fac = facts["case_ist"], facts["case_prognose"], facts["case_factor"]
    card(s, 72, 143, 618.9, 494)
    ref = q(root, "div.backref")
    disc(s, 93, 160, disc_of(q(ref, "span.disc")))
    text(s, 129, 163.5, 545, 19, rich(kid(ref, "span"), [("b", {"bold": True}), (".mono", {"font": MONO})]),
         15, MUTED, lh=19, wrap=False)
    svg = chart_svg(q(root, "div.chart-card"))
    m = SvgMap(svg, 93, 192, 576.9)
    lower = [max(0.0, p - band) for p in prog]
    width = [p + band - lo for p, lo in zip(prog, lower)]
    area_data = CategoryChartData()
    area_data.categories = MONTHS
    area_data.add_series("Basis", lower)
    area_data.add_series(f"Toleranz ± {de(band)} kWh", width)
    line_data = CategoryChartData()
    line_data.categories = MONTHS
    line_data.add_series("Prognose", prog)
    line_data.add_series("Ist", ist)
    rect = m.rect(84, 66, 468, 298)
    _, area = _combo(s, XL_CHART_TYPE.AREA_STACKED, area_data, line_data, rect, (66, 8, 16, 32))
    _axis(area.value_axis, 0, 26_000, 5_000, fmt="#,##0")
    _axis(area.category_axis, grid=False, font=SANS)
    _cross_between(area, "midCat")
    plots = list(area.plots)
    base_ser, band_ser = plots[0].series
    base_ser.format.fill.background()
    base_ser.format.line.fill.background()
    _fill(band_ser.format.fill, C["band"])
    band_ser.format.line.fill.background()
    prog_ser, ist_ser = plots[1].series
    _series_line(prog_ser, C["prognose"], 2.5, MSO_LINE_DASH_STYLE.DASH)
    _no_marker(prog_ser)
    _series_line(ist_ser, C["ist"], 2.5)
    _marker(ist_ser, XL_MARKER_STYLE.CIRCLE, 7, C["ist"], border_w=1.5)
    for i, f in enumerate(fac):
        if abs(f) >= 1:
            point = ist_ser.points[i]
            point.marker.style = XL_MARKER_STYLE.CIRCLE
            point.marker.size = 11
            point.marker.format.fill.solid()
            point.marker.format.fill.fore_color.rgb = rgb(C["anomalie"])
            point.marker.format.line.color.rgb = rgb(WHITE)
    svg_draw(s, svg, m, plot=(70, 66, 566, 364), skip=is_tick)
    calc = q(root, "div.calc")
    card(s, 712.9, 143, 495.1, 253.8)
    label(s, 733.9, 158, 453, plain(q(calc, "span.flabel")))
    for el, (y, h) in zip(qa(calc, "div.calc-row"), [(183, 36), (219, 37), (256, 36), (292, 36)]):
        text(s, 733.9, y + 7, 250, 20, plain(q(el, "span")), 16, MUTED, lh=20, wrap=False)
        text(s, 887, y + 7, 300, 21, plain(q(el, "b")), 16, INK, True, MONO, "r", lh=21, wrap=False)
        if "sum" in classes(el):
            line(s, 733.9, y + h - 1, 1187, y + h - 1, NAVY7, 2)
        else:
            line(s, 733.9, y + h - .5, 1187, y + h - .5, LINE2, 1)
    res = q(calc, "div.calc-res")
    text(s, 733.9, 349.9, 120, 22, plain(kid(res, "span")), 17, INK, True, lh=22)
    zw_from(s, q(res, "span.zw"), 0, 340, right=1187)
    icard(s, q(root, "div.icard"), 712.9, 408.8, 495.1, 94)
    footer(s, root, num)


# --------------------------------------------------------------------------- Folie 11: Dashboard (dunkel)


def _highlight(s, hl, ix, iy, iw, ih):
    """Markierung über einem Screenshot (Prozentangaben aus dem HTML)."""
    style = hl.get("style")
    left, top, width, height = (float(re.search(rf"{p}:([\d.]+)%", style).group(1)) / 100
                                for p in ("left", "top", "width", "height"))
    x, y, w, h = ix + iw * left, iy + ih * top, iw * width, ih * height
    box(s, x - 2, y - 2, w + 4, h + 4, line="rgba(0,128,160,.18)", lw=4, radius=8)  # box-shadow 4 px
    box(s, x + 1.5, y + 1.5, w - 3, h - 3, line=TEAL, lw=3, radius=5, name="Markierung")
    tag = plain(q(hl, "em"))
    tw = measure(tag, 13, True) + 16
    tx = x if "left" in classes(hl) else x + w - tw
    box(s, tx, y - 21, tw, 21.6, fill=TEAL6, radius=5)
    text(s, tx, y - 21, tw, 21.6, tag, 13, WHITE, True, align="c", anchor="m", lh=15.6, wrap=False)


def slide_dashboard(prs, spec, num, facts):
    s, root = page(prs, spec, NAVY)
    header(s, root, dark=True)
    for fig, (x, y, w) in zip(qa(root, "figure.shot-card"), [(72, 143, 646), (72, 402.2, 646), (738, 143, 470)]):
        img = q(fig, "img")
        data = base64.b64decode(img.get("src").split(",", 1)[1])
        iw, ih = Image.open(io.BytesIO(data)).size
        h_img = (w - 2) * ih / iw
        box(s, x, y, w, h_img + 38, fill="rgba(255,255,255,.06)", radius=10)
        cap = q(fig, "figcaption")
        disc(s, x + 13, y + 6, disc_of(q(cap, "span.disc")))
        text(s, x + 49, y + 1, w - 60, 36, plain(cap), 14.5, WHITE, True, anchor="m", lh=18.9, wrap=False)
        picture(s, data, x + 1, y + 37, w=w - 2, name=f"Screenshot {img.get('alt')}"[:80])
        box(s, x, y, w, h_img + 38, line=GLASS_LINE, radius=10)
        for hl in qa(fig, "span.hl"):
            _highlight(s, hl, x + 1, y + 37, w - 2, h_img)
    learn = q(root, "div.learn")
    dcard(s, 72, 556, 1136, 44)
    x = 87
    rule = q(learn, "span.rule")
    for el, tail in [(None, rule.text)] + [(ch, ch.tail) for ch in rule]:
        if el is not None and el.tag == "code":
            t = plain(el)
            cw = measure(t, 13.5, font=MONO) + 10
            box(s, x, 568.5, cw, 20, fill="rgba(255,255,255,.12)", radius=4)
            text(s, x, 568.5, cw, 20, t, 13.5, WHITE, font=MONO, align="c", anchor="m", lh=17, wrap=False)
            x += cw
        elif el is not None:
            t = plain(el)
            text(s, x, 568.5, measure(t, 14.5, True) + 10, 19, t, 14.5, WHITE, True, lh=19, wrap=False)
            x += measure(t, 14.5, True)
        if tail and tail.strip():
            t = f" {tail.strip()} "
            text(s, x, 568.5, measure(t, 14.5) + 6, 19, t, 14.5, PALE, lh=19, wrap=False)
            x += measure(t, 14.5)
    x += 12
    arrow_w = .838 * 14.5 + 2  # Pfeil aus dem DejaVu-Ersatz im HTML
    for step in qa(learn, "span.lstep"):
        text(s, x, 568.5, arrow_w + 6, 19, "→", 14.5, TEAL3, lh=19, wrap=False)
        disc(s, x + arrow_w + 8, 565, disc_of(q(step, "span.disc")))
        t = plain(step)
        tx = x + arrow_w + 8 + 26 + 8
        text(s, tx, 568.5, measure(t, 14.5) + 10, 19, t, 14.5, WHITE, lh=19, wrap=False)
        x = tx + measure(t, 14.5) + 12
    closing = q(root, "div.closing")
    glyph(s, 72, 615.5, q(closing, "p > svg"))
    text(s, 100, 613.5, 780, 22, plain(q(closing, "p")), 17, WHITE, lh=22, wrap=False)
    hand = q(closing, "span.handover")
    t = plain(hand)
    hw = 16 + 18 + 8 + measure(t, 15, True) + 16
    hx = 1208 - hw
    box(s, hx, 608, hw, 33, fill=TEAL3, radius=16.5, name="Übergabe")
    glyph(s, hx + 16, 615.5, q(hand, "svg"))
    text(s, hx + 42, 608, hw - 50, 33, t, 15, NAVY9, True, anchor="m", lh=19.5, wrap=False)
    footer(s, root, num, dark=True)


# --------------------------------------------------------------------------- Backups


def slide_b1_canvas(prs, spec, num, facts):
    s, root = page(prs, spec, NAVY)
    header(s, root, dark=True)
    for k, el in enumerate(qa(root, "div.a1")):
        x, y = 72 + (k % 2) * 428.5, 143 + (k // 2) * 91
        ml = "ml" in classes(el)
        box(s, x, y, 421.5, 84, fill="rgba(255,255,255,.06)", line=TEAL3 if ml else "rgba(255,255,255,.14)",
            radius=10, name=f"A1 {k + 1:02d}")
        disc(s, x + 13, y + 8, disc_of(q(el, "span.disc")))
        key = q(el, "span.a1-k")
        nr = plain(q(key, "b"))
        text(s, x + 51, y + 8, 360, 17, [[(nr + " ", {"font": MONO, "color": TEAL3}),
                                            (plain(key)[len(nr):].strip(), {"upper": True, "tracking": .06})]],
             13, "#E4EFF7", True, lh=17, wrap=False)
        text(s, x + 51, y + 27, 357.5, 37, plain(q(el, "p")), 14, PALE, lh=18.5)
    coh = q(root, "div.coh")
    dcard(s, 940, 143, 268, 368)
    disc(s, 959, 160, disc_of(q(coh, "span.disc")))
    label(s, 959, 202, 230, plain(q(coh, "span.dnr")), TEAL3)
    y = 229
    for li in qa(coh, "li"):
        text(s, 959, y, 230, 19, plain(q(li, "b")), 15, TEAL3, True, MONO, lh=19, wrap=False)
        body = plain(q(li, "span"))
        lines = len(wrap_lines(body, 230, 15))
        text(s, 959, y + 20, 232, 20 * lines, body, 15, "#E4EFF7", lh=20)
        y += 20 + 20 * lines + 14
    footer(s, root, num, dark=True)


def _cell_borders(cell, color=None):
    """Nur eine feine waagerechte Linie unten (oder keine), keine senkrechten Linien."""
    tc_pr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tc_pr.findall(qn(tag)):
            tc_pr.remove(old)
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = OxmlElement(tag)
        if tag == "a:lnB" and color:
            ln.set("w", str(Pt(0.75)))
            fill = OxmlElement("a:solidFill")
            clr = OxmlElement("a:srgbClr")
            clr.set("val", color.lstrip("#"))
            fill.append(clr)
            ln.append(fill)
        else:
            ln.set("w", "0")
            ln.append(OxmlElement("a:noFill"))
        tc_pr.append(ln)


def slide_b2_hyper(prs, spec, num, facts):
    s, root = page(prs, spec)
    header(s, root)
    tbl = q(root, "table.tbl")
    rows = [tr for tr in qa(tbl, "tr")]
    widths = [105.9, 127, 100.2, 111.3, 173.4]
    heights = [58.5] + [39.5] + [40] * (len(rows) - 2)
    shape = s.shapes.add_table(len(rows), 5, E(72.5), E(143.5), E(sum(widths)), E(sum(heights)))
    shape.name = "Tabelle Hyperparameter"
    table = shape.table
    table.first_row = False
    table.horz_banding = False
    for j, wv in enumerate(widths):
        table.columns[j].width = E(wv)
    for i, (tr, hv) in enumerate(zip(rows, heights)):
        table.rows[i].height = E(hv)
        best = "best" in classes(tr)
        for j, td in enumerate(tr):
            cell = table.cell(i, j)
            num_col = "num" in classes(td)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if num_col else PP_ALIGN.LEFT
            p.line_spacing = pt(19 if i == 0 else 20.8)
            r = p.add_run()
            r.text = plain(td)
            r.font.name = MONO if num_col else SANS
            r.font.size = pt(14.5 if i == 0 else 16)
            r.font.bold = best
            r.font.italic = False
            r.font.color.rgb = rgb(WHITE if i == 0 else (TEAL7 if best else INK))
            _cell_borders(cell, None if i == 0 else LINE2)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(NAVY9 if i == 0 else ("#E3F3F7" if best else WHITE))
            if i == 0:  # Kopf wie im HTML umbrechen (Innenabstand 12 px je Seite)
                cell.margin_left = cell.margin_right = E(12)
            else:  # Daten: Abstand nur auf der Leseseite, damit „unbegrenzt“ nicht umbricht
                cell.margin_left, cell.margin_right = (E(2), E(12)) if num_col else (E(12), E(2))
            cell.margin_top = cell.margin_bottom = E(0)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    box(s, 72, 143, 618.9, sum(heights) + 1, line=LINE, name="Tabellenrahmen")
    for el, y in zip(qa(root, "div.icard"), (143, 228, 313)):
        icard(s, el, 712.9, y, 495.1, 73)
    text(s, 712.9, 398, 495.1, 42, plain(q(root, "p.caveat")), 15, SOFT, lh=21)
    footer(s, root, num)


def slide_b3_treiber(prs, spec, num, facts):
    s, root = page(prs, spec)
    header(s, root)
    el = q(root, "div.chart-card")
    card(s, 72, 143, 618.9, 411.3)
    chart_card_title(s, q(el, "p.ch-t"), 93, 160, 580, 40)
    m = SvgMap(chart_svg(el), 93, 208, 576.9)
    rows = BERICHT["importance"]
    data = CategoryChartData()
    data.categories = [name for name, _ in rows]
    data.add_series("RMSE-Anstieg (kWh, Bericht Tab. D1)", [value for _, value in rows])
    _, ch = place_chart(s, XL_CHART_TYPE.BAR_CLUSTERED, data, m.rect(200, 10, 258, 42.857 * len(rows)),
                        (192, 2, 116, 2), name="Diagramm Permutation Importance")
    _axis(ch.value_axis, 0, 7_500, grid=False, visible=False)
    _axis(ch.category_axis, grid=False, font=SANS, size=16, color=INK)
    ch.category_axis.reverse_order = True
    plot = ch.plots[0]
    plot.gap_width = 79
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = '"+"#,##0" kWh"', False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.name, dl.font.size, dl.font.bold = MONO, pt(15), False
    dl.font.color.rgb = rgb(INK)
    for j, point in enumerate(plot.series[0].points):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = rgb(C["ist"] if j == 0 else C["navy300"])
    for card_el, y in zip(qa(root, "div.icard"), (143, 249, 355)):
        icard(s, card_el, 712.9, y, 495.1, 94)
    text(s, 712.9, 461, 495.1, 42, plain(q(root, "p.caveat")), 15, SOFT, lh=21)
    footer(s, root, num)


def slide_b4_lags(prs, spec, num, facts):
    s, root = page(prs, spec)
    header(s, root)
    el = q(root, "div.chart-card")
    card(s, 72, 143, 618.9, 401.3)
    chart_card_title(s, q(el, "p.ch-t"), 93, 160, 580)
    svg = chart_svg(el)
    m = SvgMap(svg, 93, 188, 576.9)
    names = [("Ist-VLS des Monats", "Ist des Monats", C["ist"], None),
             ("Vormonat als Merkmal", "Vormonat", C["grey500"], MSO_LINE_DASH_STYLE.DASH),
             ("bis zu 3 Vormonate", "Ø bis zu 3 Vormonate", C["navy300"], MSO_LINE_DASH_STYLE.ROUND_DOT)]
    data = CategoryChartData()
    data.categories = MONTHS
    for key, legend, _, _ in names:
        data.add_series(legend, [None if (v is None or v != v) else v for v in facts["lag"][key]])
    _, ch = place_chart(s, XL_CHART_TYPE.LINE, data, m.rect(60, 44, 492, 260), (44, 6, 10, 30),
                        name="Diagramm Lag-Features")
    _axis(ch.value_axis, 50, 130, 20, fmt="0")
    _axis(ch.category_axis, grid=False, font=SANS)
    _cross_between(ch, "midCat")
    for ser, (_, _, col, dash) in zip(ch.plots[0].series, names):
        _series_line(ser, col, 2.5, dash)
        _no_marker(ser)
    svg_draw(s, svg, m, plot=(50, 44, 562, 304), skip=is_tick)
    for card_el, (y, h) in zip(qa(root, "div.icard"), [(143, 94), (249, 73), (334, 94)]):
        icard(s, card_el, 712.9, y, 495.1, h)
    text(s, 712.9, 440, 495.1, 42, plain(q(root, "p.caveat")), 15, SOFT, lh=21)
    footer(s, root, num)


def slide_b5_grenzen(prs, spec, num, facts):
    s, root = page(prs, spec, NAVY)
    header(s, root, dark=True)
    for k, el in enumerate(qa(root, "div.dcard")):
        x = 72 + 288 * k
        dcard(s, x, 143, 272, 411.4)
        disc(s, x + 19, 162, disc_of(q(el, "span.disc")))
        label(s, x + 19, 208, 234, plain(q(el, "span.dnr")), TEAL3)
        head = plain(q(el, "h4"))
        lines = len(wrap_lines(head, 240, 19, True))
        text(s, x + 19, 233, 240, 23.75 * lines, head, 19, WHITE, True, lh=23.75)
        text(s, x + 19, 233 + 23.75 * lines + 16, 240, 250, ["• " + plain(li) for li in qa(el, "li")], 15, PALE,
             lh=21.75, gap=4)
    merk(s, q(root, "div.merk"), 72, 566.4, 1136, 72.6, dark=True)
    footer(s, root, num, dark=True)


def slide_b6_aufwand(prs, spec, num, facts):
    s, root = page(prs, spec)
    header(s, root)
    el = q(root, "div.chart-card")
    card(s, 72, 143, 618.9, 391.3)
    chart_card_title(s, q(el, "p.ch-t"), 93, 160, 400)
    svg = chart_svg(el)
    m = SvgMap(svg, 93, 188, 576.9)
    mean = facts["n_alerts"] / 12
    bars = CategoryChartData()
    bars.categories = MONTHS
    bars.add_series("Prüfhinweise 2025 (q99)", facts["alerts_by_month"])
    avg = CategoryChartData()
    avg.categories = MONTHS
    avg.add_series(f"Ø {de(mean, 1)} je Monat", [mean] * 12)
    gf, ch = _combo(s, XL_CHART_TYPE.COLUMN_CLUSTERED, bars, avg, m.rect(46, 30, 522, 264), (40, 8, 10, 30))
    gf.name = "Diagramm Prüfhinweise je Monat"
    _axis(ch.value_axis, 0, 18, 6, fmt="0")
    _axis(ch.category_axis, grid=False, font=SANS)
    bar_plot, line_plot = list(ch.plots)
    bar_plot.gap_width, bar_plot.overlap = 38, 0
    bar_plot.has_data_labels = True
    dl = bar_plot.data_labels
    dl.number_format, dl.number_format_is_linked = "0", False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.name, dl.font.size, dl.font.bold = MONO, pt(15), True
    dl.font.color.rgb = rgb(INK)
    bar_plot.series[0].format.fill.solid()
    bar_plot.series[0].format.fill.fore_color.rgb = rgb(C["residuum"])
    avg_ser = line_plot.series[0]
    _series_line(avg_ser, C["schwelle"], 2, MSO_LINE_DASH_STYLE.DASH)
    _no_marker(avg_ser)
    # Ø-Beschriftung aus dem SVG, aber in die Titelzeile gesetzt: im HTML überdeckt sie die Dez-Säule
    avg_label = next(t for t in svg.iter("text") if (t.get("fill") or "").upper() == C["amber700"])
    t = (avg_label.text or "").strip()
    tw = measure(t, 16, True)
    line(s, 669.9 - tw - 34, 170, 669.9 - tw - 8, 170, C["schwelle"], 2, MSO_LINE_DASH_STYLE.DASH)
    text(s, 669.9 - tw - 4, 160, tw + 4, 20, t, 16, C["amber700"], True, align="r", lh=20, wrap=False,
         name="Ø-Linie Beschriftung")
    for card_el, y in zip(qa(root, "div.icard"), (143, 249, 355)):
        icard(s, card_el, 712.9, y, 495.1, 94)
    footer(s, root, num)


# --------------------------------------------------------------------------- Zusammenbau


def normalize_axis_ids(prs):
    """Achsen-IDs positiv und reproduzierbar setzen.

    python-pptx vergibt zufällige, teils negative IDs; das Schema verlangt
    xsd:unsignedInt, und strenge Leser (z. B. PPTX-Importer) brechen daran ab.
    """
    for slide in prs.slides:
        for n, shape in enumerate(s for s in slide.shapes if s.has_chart):
            space = shape.chart._chartSpace
            mapping = {}
            for el in space.iter(qn("c:axId"), qn("c:crossAx")):
                old = el.get("val")
                mapping.setdefault(old, str(500_000 + 1_000 * slide.slide_id + 10 * n + len(mapping)))
                el.set("val", mapping[old])


BUILDERS = [slide_canvas, slide_fehlerkosten, slide_faden, slide_validierung, slide_vls, slide_modell,
            slide_schwelle, slide_kalibrierung, slide_band, slide_fall, slide_dashboard,
            slide_b1_canvas, slide_b2_hyper, slide_b3_treiber, slide_b4_lags, slide_b5_grenzen, slide_b6_aufwand]


def _assemble(specs, numbers, facts):
    prs = Presentation()
    prs.slide_width, prs.slide_height = E(W), E(H)
    prs.core_properties.title = "ML Canvas, Methodik und Modell, Prüffall"
    prs.core_properties.author = "Kiko · Gruppe 6"
    prs.core_properties.language = "de-DE"
    for builder, spec, num in zip(BUILDERS, specs, numbers):
        builder(prs, spec, num, facts)
    normalize_axis_ids(prs)
    return prs


def build(path: Path = OUT) -> Path:
    facts = BUILD["extract_facts"](BUILD["Notebook"].load())
    specs = BUILD["build_slides"](facts)
    numbers = BUILD["slide_numbers"](specs)
    assert len(specs) == len(BUILDERS), "HTML- und PPTX-Foliensatz laufen auseinander"
    _MISSING.clear()
    prs = _assemble(specs, numbers, facts)
    if _MISSING:  # fehlende Icons einmal gesammelt rendern, dann neu zusammenbauen
        render_icons(_MISSING)
        _MISSING.clear()
        prs = _assemble(specs, numbers, facts)
        if _MISSING:
            raise RuntimeError(f"Icons fehlen weiterhin: {sorted(_MISSING)}")
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path


if __name__ == "__main__":
    out = build()
    print(f"{out.relative_to(ROOT)} ({out.stat().st_size // 1024} KB)")
